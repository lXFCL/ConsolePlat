from __future__ import annotations

import csv
import re
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import urlsplit, urlunsplit
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook


RESULT_HEADERS = (
    "排名",
    "商品标题",
    "销量",
    "销量数值",
    "价格",
    "商品链接",
    "搜索关键词",
    "采集时间",
    "主图链接",
    "本地主图",
    "状态",
    "失败原因",
)


@dataclass
class SelectionCandidate:
    title: str
    sales_text: str
    sales_count: int | None
    price: str
    product_url: str
    image_url: str
    keyword: str
    rank: int = 0
    status: str = "待下载"
    error: str = ""
    local_image_path: str = ""
    collected_at: str = ""


@dataclass(frozen=True)
class SelectionBatch:
    batch_dir: Path
    image_dir: Path
    csv_path: Path
    xlsx_path: Path
    failure_path: Path


@dataclass(frozen=True)
class ExportResult:
    batch_dir: Path
    csv_path: Path
    xlsx_path: Path
    failure_path: Path


def parse_sales_count(text: str) -> int | None:
    clean = str(text or "").replace(",", "").strip()
    if not clean:
        return None
    match = re.search(r"(\d+(?:\.\d+)?)\s*(万|k|K)?", clean)
    if match is None:
        return None
    value = float(match.group(1))
    unit = match.group(2)
    if unit == "万":
        value *= 10_000
    elif unit and unit.lower() == "k":
        value *= 1_000
    return int(value)


def canonical_product_url(url: str) -> str:
    parts = urlsplit(str(url or "").strip())
    if not parts.scheme or not parts.netloc:
        return str(url or "").strip()
    return urlunsplit((parts.scheme, parts.netloc.lower(), parts.path, "", ""))


def rank_candidates(candidates: Iterable[SelectionCandidate], limit: int = 30) -> list[SelectionCandidate]:
    unique: dict[str, SelectionCandidate] = {}
    for candidate in candidates:
        if candidate.sales_count is None:
            continue
        key = canonical_product_url(candidate.product_url)
        previous = unique.get(key)
        if previous is None or int(candidate.sales_count) > int(previous.sales_count or 0):
            unique[key] = candidate
    ordered = sorted(
        unique.values(),
        key=lambda candidate: (-int(candidate.sales_count or 0), candidate.title.casefold()),
    )[: max(0, int(limit))]
    return [replace(candidate, rank=index) for index, candidate in enumerate(ordered, start=1)]


def create_batch(
    program_data_dir: str | Path,
    keywords: list[str],
    created_at: str | None = None,
) -> SelectionBatch:
    root_text = str(program_data_dir or "").strip()
    if not root_text:
        raise ValueError("请先在设置中配置程序数据目录")
    root = Path(root_text).expanduser()
    stamp = _parse_created_at(created_at).strftime("%Y%m%d_%H%M%S")
    keyword = _safe_name(next((item.strip() for item in keywords if item.strip()), "选品"))
    batch_dir = root / "AI选品" / f"{stamp}_{keyword}"
    return SelectionBatch(
        batch_dir=batch_dir,
        image_dir=batch_dir / "主图",
        csv_path=batch_dir / "选品结果.csv",
        xlsx_path=batch_dir / "选品结果.xlsx",
        failure_path=batch_dir / "失败清单.csv",
    )


def export_batch(
    batch: SelectionBatch,
    candidates: list[SelectionCandidate],
    failures: list[str] | None = None,
) -> ExportResult:
    batch.image_dir.mkdir(parents=True, exist_ok=True)
    rows = [_candidate_row(candidate) for candidate in candidates]
    with batch.csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(RESULT_HEADERS)
        writer.writerows(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "选品结果"
    sheet.append(RESULT_HEADERS)
    for row in rows:
        sheet.append(row)
    for column in sheet.columns:
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(batch.xlsx_path)
    _write_failures(batch.failure_path, failures or [])
    return ExportResult(batch.batch_dir, batch.csv_path, batch.xlsx_path, batch.failure_path)


def load_batches(program_data_dir: str | Path) -> list[SelectionBatch]:
    root_text = str(program_data_dir or "").strip()
    if not root_text:
        return []
    parent = Path(root_text).expanduser() / "AI选品"
    if not parent.is_dir():
        return []
    batches = [
        SelectionBatch(
            batch_dir=directory,
            image_dir=directory / "主图",
            csv_path=directory / "选品结果.csv",
            xlsx_path=directory / "选品结果.xlsx",
            failure_path=directory / "失败清单.csv",
        )
        for directory in parent.iterdir()
        if directory.is_dir() and (directory / "选品结果.csv").is_file()
    ]
    return sorted(batches, key=lambda batch: batch.batch_dir.name, reverse=True)


def load_candidates(batch: SelectionBatch) -> list[SelectionCandidate]:
    if not batch.csv_path.is_file():
        return []
    with batch.csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    candidates: list[SelectionCandidate] = []
    for row in rows:
        sales_count = parse_sales_count(row.get("销量数值") or "")
        candidates.append(
            SelectionCandidate(
                title=row.get("商品标题") or "",
                sales_text=row.get("销量") or "",
                sales_count=sales_count,
                price=row.get("价格") or "",
                product_url=row.get("商品链接") or "",
                image_url=row.get("主图链接") or "",
                keyword=row.get("搜索关键词") or "",
                rank=int(row.get("排名") or 0),
                status=row.get("状态") or "",
                error=row.get("失败原因") or "",
                local_image_path=row.get("本地主图") or "",
                collected_at=row.get("采集时间") or "",
            )
        )
    return candidates


def download_image(url: str, target: str | Path, timeout: int = 20) -> None:
    request = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(request, timeout=timeout) as response:  # noqa: S310 - image URL is read from the public page.
        Path(target).write_bytes(response.read())


def download_candidate_images(candidates: Iterable[SelectionCandidate], image_dir: str | Path) -> None:
    target_dir = Path(image_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    for index, candidate in enumerate(candidates, start=1):
        if not candidate.image_url:
            candidate.status = "主图缺失"
            candidate.error = "搜索结果没有可下载的主图"
            continue
        target = target_dir / f"{candidate.rank or index:02d}_{_safe_name(candidate.title)[:60]}.jpg"
        try:
            download_image(candidate.image_url, target)
        except Exception as exc:  # noqa: BLE001 - failures must not stop the batch.
            candidate.status = "主图下载失败"
            candidate.error = str(exc)
            continue
        candidate.local_image_path = str(target)
        candidate.status = "已下载"
        candidate.error = ""


def read_xlsx_headers(path: str | Path) -> tuple[str, ...]:
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook.active
        return tuple(str(cell.value or "") for cell in next(sheet.iter_rows(max_row=1)))
    finally:
        workbook.close()


def _candidate_row(candidate: SelectionCandidate) -> list[str | int]:
    return [
        candidate.rank,
        candidate.title,
        candidate.sales_text,
        candidate.sales_count or "",
        candidate.price,
        candidate.product_url,
        candidate.keyword,
        candidate.collected_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        candidate.image_url,
        candidate.local_image_path,
        candidate.status,
        candidate.error,
    ]


def _write_failures(path: Path, failures: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("失败原因",))
        for failure in failures:
            writer.writerow((failure,))


def _parse_created_at(value: str | None) -> datetime:
    if not value:
        return datetime.now()
    return datetime.fromisoformat(value)


def _safe_name(value: str) -> str:
    clean = re.sub(r'[\\/:*?"<>|\r\n]+', "_", str(value or "").strip())
    return clean.strip(" ._") or "选品"
