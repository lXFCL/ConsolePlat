from __future__ import annotations

import hashlib
import colorsys
import random
import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from consoleplat.config import resolve_project_dir
from consoleplat.paths import project_root
from consoleplat.services.putaway_sync_service import PutawaySyncSummary, sync_putaway_assets


POSAI_ROOT = resolve_project_dir("posaiimg") or (project_root() / "modules" / "PosAiImg")
if str(POSAI_ROOT) not in sys.path:
    sys.path.append(str(POSAI_ROOT))

try:
    from tools.tshirt_print_tool import (  # type: ignore
        Placement,
        composite_one,
        crop_to_alpha,
        remove_near_white_background,
        wave_displace,
    )
except ModuleNotFoundError:  # pragma: no cover
    from tshirt_print_tool import (  # type: ignore
        Placement,
        composite_one,
        crop_to_alpha,
        remove_near_white_background,
        wave_displace,
    )


STORE_NAMES = {"BO": "YUHAOBO", "SZW": "YUHOOBO"}
PRODUCT_CATEGORY = "T恤"
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
DEFAULT_MODEL_DIR = POSAI_ROOT / "模特图/干净"
DEFAULT_WHITE_MODEL_NAMES = ["主图5-白.jpg", "主图6-白.jpg"]
DEFAULT_BLACK_MODEL_NAMES = ["主图1-黑.jpg", "主图2-黑.jpg"]
WHITE_COLOR_NAME = "白"
BLACK_COLOR_NAME = "黑"
HEADER_ROW = ["店铺名称", "产品分类", "产品标题", "产品序列号", "颜色"]
_PRINT_METRIC_CACHE: dict[tuple[str, int], tuple[float, float]] = {}


@dataclass(frozen=True)
class AIEditFormalizeSummary:
    ok: bool
    renamed_outputs: list[str]
    product_outputs: list[str]
    xlsx_path: str
    putaway: PutawaySyncSummary | None
    color_assignments: dict[str, str] | None = None
    message: str = ""


@dataclass(frozen=True)
class MockupAssignment:
    model_path: Path
    color_name: str


DEFAULT_CHEST_SAFE_BOX = {
    "center_x": 0.50,
    "center_y": 0.43,
    "max_width_ratio": 0.30,
    "max_height_ratio": 0.22,
}


def _stable_seed_text(text: str) -> int:
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    return int(digest[:16], 16)


def safe_filename(text: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if char in invalid or ord(char) < 32 else char for char in text)
    return re.sub(r"_+", "_", cleaned).strip(" ._")


def choose_white_model(model_dir: Path) -> Path:
    for name in DEFAULT_WHITE_MODEL_NAMES:
        path = model_dir / name
        if path.exists():
            return path
    for path in sorted(model_dir.iterdir()):
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES and "白" in path.stem:
            return path
    raise FileNotFoundError(f"未找到白色模特图：{model_dir}")


def choose_black_model(model_dir: Path) -> Path:
    for name in DEFAULT_BLACK_MODEL_NAMES:
        path = model_dir / name
        if path.exists():
            return path
    for path in sorted(model_dir.iterdir()):
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES and "黑" in path.stem:
            return path
    raise FileNotFoundError(f"未找到黑色模特图：{model_dir}")


def _subject_brightness(print_path: Path) -> float:
    brightness, _saturation = _subject_metrics(print_path)
    return brightness


def _subject_metrics(print_path: Path) -> tuple[float, float]:
    cache_key = (str(print_path.resolve()), int(print_path.stat().st_mtime_ns))
    cached = _PRINT_METRIC_CACHE.get(cache_key)
    if cached is not None:
        return cached
    image = Image.open(print_path).convert("RGBA")
    pixels = image.load()
    brightness_total = 0.0
    saturation_total = 0.0
    alpha_total = 0.0
    for y in range(image.height):
        for x in range(image.width):
            red, green, blue, alpha = pixels[x, y]
            if alpha <= 0:
                continue
            brightness = red * 0.299 + green * 0.587 + blue * 0.114
            _hue, saturation, _value = colorsys.rgb_to_hsv(red / 255.0, green / 255.0, blue / 255.0)
            weight = alpha / 255.0
            brightness_total += brightness * weight
            saturation_total += saturation * weight
            alpha_total += weight
    if alpha_total <= 0:
        metrics = (255.0, 0.0)
    else:
        metrics = (brightness_total / alpha_total, saturation_total / alpha_total)
    _PRINT_METRIC_CACHE[cache_key] = metrics
    return metrics


def _subject_saturation(print_path: Path) -> float:
    _brightness, saturation = _subject_metrics(print_path)
    return saturation


def is_grayscale_print(print_path: Path, *, saturation_threshold: float) -> bool:
    return _subject_saturation(print_path) < max(0.0, min(1.0, saturation_threshold))


def choose_mockup_model_for_print(print_path: Path) -> str:
    return "white" if _subject_brightness(print_path) < 128 else "black"


def _collect_model_pools(model_dir: Path) -> tuple[list[Path], list[Path]]:
    black_pool: list[Path] = []
    white_pool: list[Path] = []
    for path in sorted(model_dir.iterdir()):
        if not path.is_file() or path.suffix.lower() not in IMAGE_SUFFIXES:
            continue
        stem = path.stem
        if "黑" in stem:
            black_pool.append(path)
        elif "白" in stem:
            white_pool.append(path)
    if not black_pool:
        black_pool = [choose_black_model(model_dir)]
    if not white_pool:
        white_pool = [choose_white_model(model_dir)]
    return black_pool, white_pool


def choose_mockup_models_for_batch(
    print_paths: list[Path],
    *,
    model_dir: Path,
    assignment_seed: str,
    saturation_threshold: float = 0.15,
) -> dict[Path, MockupAssignment]:
    if not print_paths:
        return {}
    black_pool, white_pool = _collect_model_pools(model_dir)
    ordered_paths = sorted(print_paths, key=lambda path: path.name.lower())
    rng = random.Random(_stable_seed_text(assignment_seed))
    grayscale_paths: list[Path] = []
    colorful_paths: list[Path] = []
    for path in ordered_paths:
        if is_grayscale_print(path, saturation_threshold=saturation_threshold):
            grayscale_paths.append(path)
        else:
            colorful_paths.append(path)
    shuffled_indexes = list(range(len(colorful_paths)))
    rng.shuffle(shuffled_indexes)

    black_target = len(colorful_paths) // 2
    white_target = len(colorful_paths) - black_target
    color_slots = [BLACK_COLOR_NAME] * black_target + [WHITE_COLOR_NAME] * white_target
    rng.shuffle(color_slots)

    black_cycle = 0
    white_cycle = 0
    assignments: dict[Path, MockupAssignment] = {}
    for print_path in grayscale_paths:
        model_path = white_pool[white_cycle % len(white_pool)]
        white_cycle += 1
        assignments[print_path] = MockupAssignment(model_path=model_path, color_name=WHITE_COLOR_NAME)
    for order_index, shuffled_index in enumerate(shuffled_indexes):
        print_path = colorful_paths[shuffled_index]
        color_name = color_slots[order_index]
        if color_name == BLACK_COLOR_NAME:
            model_path = black_pool[black_cycle % len(black_pool)]
            black_cycle += 1
        else:
            model_path = white_pool[white_cycle % len(white_pool)]
            white_cycle += 1
        assignments[print_path] = MockupAssignment(model_path=model_path, color_name=color_name)
    return assignments


def _assignment_to_color_name(assignment: str) -> str:
    return WHITE_COLOR_NAME if assignment == "white" else BLACK_COLOR_NAME


def fit_print_within_safe_box(
    *,
    print_img: Image.Image,
    base_size: tuple[int, int],
    width_ratio: float,
    height_ratio: float,
    remove_white_bg: bool,
    wave_strength: float,
    rotation: float,
    opacity: float,
) -> Image.Image:
    if remove_white_bg:
        print_img = remove_near_white_background(print_img)
    print_img = crop_to_alpha(print_img)

    max_width = max(1, int(base_size[0] * width_ratio))
    max_height = max(1, int(base_size[1] * height_ratio))
    ratio = min(max_width / max(1, print_img.width), max_height / max(1, print_img.height))
    target_w = max(1, int(print_img.width * ratio))
    target_h = max(1, int(print_img.height * ratio))
    print_img = print_img.resize((target_w, target_h), Image.Resampling.LANCZOS)

    if wave_strength:
        print_img = wave_displace(print_img, wave_strength)

    if rotation:
        print_img = print_img.rotate(
            rotation,
            expand=True,
            resample=Image.Resampling.BICUBIC,
            fillcolor=(0, 0, 0, 0),
        )

    if opacity < 1:
        r, g, b, a = print_img.split()
        alpha = max(0, min(1, opacity))
        a = a.point(lambda value: round(value * alpha))
        print_img = Image.merge("RGBA", (r, g, b, a))
    return print_img


def rename_split_outputs(
    *,
    split_paths: list[Path],
    final_transparent_dir: Path,
    prefix: str,
    start_number: int,
) -> list[Path]:
    final_transparent_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for index, source in enumerate(split_paths):
        sku = f"{prefix}-{start_number + index}"
        target = final_transparent_dir / f"{sku}.png"
        if source.resolve() != target.resolve():
            target.unlink(missing_ok=True)
            shutil.copy2(source, target)
        outputs.append(target)
    return outputs


def build_product_images(
    *,
    print_paths: list[Path],
    final_product_dir: Path,
    product_title: str,
    model_dir: Path,
    saturation_threshold: float = 0.15,
) -> tuple[list[Path], dict[str, str]]:
    if not model_dir.exists() or not model_dir.is_dir():
        raise FileNotFoundError(f"模特底图目录不存在: {model_dir}")
    final_product_dir.mkdir(parents=True, exist_ok=True)
    placement = Placement(
        center_x=DEFAULT_CHEST_SAFE_BOX["center_x"],
        center_y=DEFAULT_CHEST_SAFE_BOX["center_y"],
        width=DEFAULT_CHEST_SAFE_BOX["max_width_ratio"],
        opacity=0.92,
        rotation=0.0,
        shadow_strength=0.32,
        wave_strength=0.012,
        remove_white_bg=False,
    )
    assignments = choose_mockup_models_for_batch(
        print_paths,
        model_dir=model_dir,
        assignment_seed=f"{print_paths[0].stem}-{len(print_paths)}",
        saturation_threshold=saturation_threshold,
    )
    outputs: list[Path] = []
    color_assignments: dict[str, str] = {}
    for print_path in print_paths:
        sku = print_path.stem
        output = final_product_dir / safe_filename(f"{sku}_{product_title}.png")
        output.unlink(missing_ok=True)
        assignment = assignments[print_path]
        color_assignments[sku] = assignment.color_name
        base = Image.open(assignment.model_path).convert("RGBA")
        design = Image.open(print_path).convert("RGBA")
        design = fit_print_within_safe_box(
            print_img=design,
            base_size=base.size,
            width_ratio=DEFAULT_CHEST_SAFE_BOX["max_width_ratio"],
            height_ratio=DEFAULT_CHEST_SAFE_BOX["max_height_ratio"],
            remove_white_bg=placement.remove_white_bg,
            wave_strength=placement.wave_strength,
            rotation=placement.rotation,
            opacity=placement.opacity,
        )
        x = int(base.width * placement.center_x - design.width / 2)
        y = int(base.height * placement.center_y - design.height / 2)
        result = composite_one.__globals__["cloth_blend"](base, design, x, y, placement)
        output.parent.mkdir(parents=True, exist_ok=True)
        result.save(output)
        outputs.append(output)
    return outputs, color_assignments


def _build_product_images_balanced(
    *,
    print_paths: list[Path],
    final_product_dir: Path,
    product_title: str,
    model_dir: Path,
    saturation_threshold: float = 0.15,
) -> tuple[list[Path], dict[str, str]]:
    if not model_dir.exists() or not model_dir.is_dir():
        raise FileNotFoundError(f"模特底图目录不存在: {model_dir}")
    final_product_dir.mkdir(parents=True, exist_ok=True)
    placement = Placement(
        center_x=DEFAULT_CHEST_SAFE_BOX["center_x"],
        center_y=DEFAULT_CHEST_SAFE_BOX["center_y"],
        width=DEFAULT_CHEST_SAFE_BOX["max_width_ratio"],
        opacity=0.92,
        rotation=0.0,
        shadow_strength=0.32,
        wave_strength=0.012,
        remove_white_bg=False,
    )
    assignments = choose_mockup_models_for_batch(
        print_paths,
        model_dir=model_dir,
        assignment_seed=f"{print_paths[0].stem}-{len(print_paths)}",
        saturation_threshold=saturation_threshold,
    )
    cloth_blend = composite_one.__globals__["cloth_blend"]
    outputs: list[Path] = []
    color_assignments: dict[str, str] = {}
    for print_path in print_paths:
        sku = print_path.stem
        output = final_product_dir / safe_filename(f"{sku}_{product_title}.png")
        output.unlink(missing_ok=True)
        assignment = assignments[print_path]
        color_assignments[sku] = assignment.color_name
        with Image.open(assignment.model_path).convert("RGBA") as base_image:
            base = base_image.copy()
        with Image.open(print_path).convert("RGBA") as design_image:
            design = fit_print_within_safe_box(
                print_img=design_image.copy(),
                base_size=base.size,
                width_ratio=DEFAULT_CHEST_SAFE_BOX["max_width_ratio"],
                height_ratio=DEFAULT_CHEST_SAFE_BOX["max_height_ratio"],
                remove_white_bg=placement.remove_white_bg,
                wave_strength=placement.wave_strength,
                rotation=placement.rotation,
                opacity=placement.opacity,
            )
        x = int(base.width * placement.center_x - design.width / 2)
        y = int(base.height * placement.center_y - design.height / 2)
        result = cloth_blend(base, design, x, y, placement)
        output.parent.mkdir(parents=True, exist_ok=True)
        result.save(output)
        outputs.append(output)
    return outputs, color_assignments


build_product_images = _build_product_images_balanced


def _load_existing_color_assignments(xlsx_path: Path) -> dict[str, str]:
    if not xlsx_path.exists():
        return {}
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        assignments: dict[str, str] = {}
        row = 2
        while True:
            sku = ws[f"D{row}"].value
            if not sku:
                break
            color = ws[f"E{row}"].value
            if color:
                assignments[str(sku)] = str(color)
            row += 1
        return assignments
    finally:
        wb.close()


def _collect_color_assignments_from_transparent_dir(
    final_transparent_dir: Path,
    *,
    saturation_threshold: float = 0.15,
) -> dict[str, str]:
    if not final_transparent_dir.exists():
        return {}
    assignments: dict[str, str] = {}
    for path in sorted(final_transparent_dir.glob("*.png")):
        if not path.is_file():
            continue
        assignments[path.stem] = WHITE_COLOR_NAME if is_grayscale_print(
            path,
            saturation_threshold=saturation_threshold,
        ) else BLACK_COLOR_NAME
    return assignments


def backfill_xlsx_colors_from_transparent_dir(
    *,
    xlsx_path: Path,
    final_transparent_dir: Path,
    saturation_threshold: float = 0.15,
) -> int:
    if not xlsx_path.exists() or not final_transparent_dir.exists():
        return 0
    inferred_colors = _collect_color_assignments_from_transparent_dir(
        final_transparent_dir,
        saturation_threshold=saturation_threshold,
    )
    if not inferred_colors:
        return 0

    wb = load_workbook(xlsx_path)
    try:
        ws = wb.active
        updated = 0
        row = 2
        while True:
            sku = ws[f"D{row}"].value
            if not sku:
                break
            color_cell = ws[f"E{row}"]
            if not color_cell.value:
                inferred = inferred_colors.get(str(sku))
                if inferred:
                    color_cell.value = inferred
                    updated += 1
            row += 1
        if updated:
            wb.save(xlsx_path)
        return updated
    finally:
        wb.close()


def write_xlsx(
    *,
    xlsx_path: Path,
    prefix: str,
    start_number: int,
    count: int,
    product_title: str,
    color_assignments: dict[str, str] | None = None,
) -> Path:
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    merged_color_assignments = _load_existing_color_assignments(xlsx_path)
    if color_assignments:
        merged_color_assignments.update(color_assignments)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADER_ROW)
    store_name = STORE_NAMES[prefix]
    for number in range(start_number, start_number + count):
        sku = f"{prefix}-{number}"
        ws.append([store_name, PRODUCT_CATEGORY, product_title, sku, merged_color_assignments.get(sku, "")])
    for col, width in {"A": 16, "B": 12, "C": 90, "D": 18, "E": 10}.items():
        ws.column_dimensions[col].width = width
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


def formalize_ai_edit_outputs(
    *,
    split_paths: list[Path],
    final_transparent_dir: Path,
    final_product_dir: Path,
    xlsx_path: Path,
    putaway_data_dir: Path,
    model_dir: Path,
    prefix: str,
    start_number: int,
    product_title: str,
    xlsx_batch_start_number: int | None = None,
    xlsx_batch_count: int | None = None,
    saturation_threshold: float = 0.15,
) -> AIEditFormalizeSummary:
    if not model_dir.exists() or not model_dir.is_dir():
        raise FileNotFoundError(f"模特底图目录不存在: {model_dir}")
    renamed_outputs = rename_split_outputs(
        split_paths=split_paths,
        final_transparent_dir=final_transparent_dir,
        prefix=prefix,
        start_number=start_number,
    )
    product_outputs, color_assignments = build_product_images(
        print_paths=renamed_outputs,
        final_product_dir=final_product_dir,
        product_title=product_title,
        model_dir=model_dir,
        saturation_threshold=saturation_threshold,
    )
    write_xlsx(
        xlsx_path=xlsx_path,
        prefix=prefix,
        start_number=int(xlsx_batch_start_number or start_number),
        count=max(1, int(xlsx_batch_count or len(renamed_outputs))),
        product_title=product_title,
        color_assignments=color_assignments,
    )
    putaway = sync_putaway_assets(
        source_images_dir=final_product_dir,
        source_xlsx_path=xlsx_path,
        target_data_dir=putaway_data_dir,
        replace_image_names={path.name for path in product_outputs},
    )
    ok = bool(renamed_outputs) and bool(product_outputs) and xlsx_path.exists() and putaway.ok
    return AIEditFormalizeSummary(
        ok=ok,
        renamed_outputs=[str(path) for path in renamed_outputs],
        product_outputs=[str(path) for path in product_outputs],
        xlsx_path=str(xlsx_path),
        putaway=putaway,
        color_assignments=color_assignments,
        message=(
            f"正式模式后处理完成：透明底 {len(renamed_outputs)} 张，产品图 {len(product_outputs)} 张，XLSX {xlsx_path.name}"
            if ok
            else "正式模式后处理失败"
        ),
    )
