from __future__ import annotations

from collections import defaultdict
from copy import copy
from pathlib import Path
from typing import Callable, Iterable

import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter

from image_utils import download_image, fit_dimensions, prepare_excel_image, to_high_res_url
from models import TemuSkuRecord, WriteSummary


Logger = Callable[[str], None]

SIZE_COLUMNS = {
    "black": {"S": "C", "M": "D", "L": "E", "XL": "F", "XXL": "G"},
    "white": {"S": "N", "M": "O", "L": "P", "XL": "Q", "XXL": "R"},
}

BASE_COLUMNS = {
    "black": {"sku": "B", "total": "H", "front": "I", "sum_start": "C", "sum_end": "G"},
    "white": {"sku": "M", "total": "S", "front": "T", "sum_start": "N", "sum_end": "R"},
}

COLOR_ALIASES = {
    "黑": "black",
    "黑色": "black",
    "白": "white",
    "白色": "white",
}


def normalize_color(color: str) -> str:
    text = (color or "").strip()
    return COLOR_ALIASES.get(text, "")


def clear_target_area(ws, start_row: int = 5) -> None:
    for row in range(start_row, ws.max_row + 1):
        for col in ["B", "C", "D", "E", "F", "G", "I", "J", "M", "N", "O", "P", "Q", "R", "T", "U"]:
            ws[f"{col}{row}"].value = None

    # Remove old images in the two front/back image areas.
    kept_images = []
    for img in getattr(ws, "_images", []):
        anchor = getattr(img, "anchor", None)
        marker = getattr(anchor, "_from", None)
        col_idx = getattr(marker, "col", None)
        row_idx = getattr(marker, "row", None)
        if col_idx is None or row_idx is None:
            kept_images.append(img)
            continue
        cell_col = get_column_letter(col_idx + 1)
        cell_row = row_idx + 1
        if cell_row >= start_row and cell_col in {"I", "J", "T", "U"}:
            continue
        kept_images.append(img)
    ws._images = kept_images


def ensure_row(ws, row: int) -> None:
    if row <= ws.max_row:
        return
    template_row = max(5, ws.max_row)
    ws.append([])
    ws.row_dimensions[row].height = ws.row_dimensions[template_row].height
    for col in range(1, ws.max_column + 1):
        src = ws.cell(template_row, col)
        dst = ws.cell(row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.border:
            dst.border = copy(src.border)
        if src.fill:
            dst.fill = copy(src.fill)


def _group_records(records: Iterable[TemuSkuRecord]) -> tuple[dict[tuple[str, str], dict], list[str], int]:
    groups: dict[tuple[str, str], dict] = {}
    warnings: list[str] = []
    skipped = 0

    for record in records:
        color_key = normalize_color(record.color)
        if not color_key:
            skipped += 1
            warnings.append(f"跳过无法识别颜色的数据：{record.sku_code or record.sku_attr}")
            continue
        if record.size not in SIZE_COLUMNS[color_key]:
            skipped += 1
            warnings.append(f"跳过无法识别尺码的数据：{record.sku_code or record.sku_attr}")
            continue
        if not record.product_sku:
            skipped += 1
            warnings.append(f"跳过缺少商品货号的数据：{record.sku_code or record.sku_attr}")
            continue

        key = (color_key, record.product_sku)
        group = groups.setdefault(
            key,
            {
                "color": color_key,
                "product_sku": record.product_sku,
                "sizes": defaultdict(int),
                "image_url": "",
            },
        )
        group["sizes"][record.size] += int(record.quantity or 0)
        if not group["image_url"] and record.image_url:
            group["image_url"] = to_high_res_url(record.image_url)

    return groups, warnings, skipped


def write_purchase_sheet(
    records: Iterable[TemuSkuRecord],
    template_path: str | Path,
    output_path: str | Path,
    image_dir: str | Path,
    logger: Logger | None = None,
) -> WriteSummary:
    log = logger or (lambda _message: None)
    template_path = Path(template_path)
    output_path = Path(output_path)
    image_dir = Path(image_dir)

    records = list(records)
    groups, warnings, skipped = _group_records(records)

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    clear_target_area(ws)

    next_rows = {"black": 5, "white": 5}
    written_rows = {"black": 0, "white": 0}

    for (color_key, product_sku), group in sorted(groups.items(), key=lambda item: (item[0][0], item[0][1])):
        row = next_rows[color_key]
        ensure_row(ws, row)
        base = BASE_COLUMNS[color_key]

        ws[f"{base['sku']}{row}"] = product_sku
        for size, qty in group["sizes"].items():
            ws[f"{SIZE_COLUMNS[color_key][size]}{row}"] = qty

        ws[f"{base['total']}{row}"] = f"=SUM({base['sum_start']}{row}:{base['sum_end']}{row})"

        image_url = group.get("image_url") or ""
        if image_url:
            try:
                source = download_image(image_url, image_dir / "originals")
                prepared = prepare_excel_image(source, image_dir / "excel")
                excel_img = ExcelImage(str(prepared))
                excel_img.width, excel_img.height = fit_dimensions(prepared)
                excel_img.anchor = f"{base['front']}{row}"
                ws.add_image(excel_img)
            except Exception as exc:  # noqa: BLE001 - keep batch write running.
                warnings.append(f"{product_sku} 图片写入失败：{exc}")
                log(f"{product_sku} 图片写入失败：{exc}")

        next_rows[color_key] += 1
        written_rows[color_key] += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return WriteSummary(
        output_path=str(output_path),
        total_records=len(records),
        black_rows=written_rows["black"],
        white_rows=written_rows["white"],
        skipped_records=skipped,
        warnings=warnings,
    )
