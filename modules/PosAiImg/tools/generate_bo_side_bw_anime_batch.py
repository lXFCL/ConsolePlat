from __future__ import annotations

import argparse
import random
import re
import sys
import time
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageOps

sys.path.insert(0, str(Path(__file__).resolve().parent))

from comfy_print_batch import copy_outputs, make_workflow, queue_prompt, start_comfyui, stop_comfyui, wait_prompt, wait_server  # noqa: E402
from tshirt_print_tool import Placement, composite_one, list_images  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
MODEL_DIR = ROOT / "模特图-干净"
TEMPLATE_DIR = ROOT / "衣物对应的xlsx" / "简约200"


def batch_name(start: int, count: int, batch_date: str | None = None) -> str:
    end = start + count - 1
    stamp = batch_date or date.today().isoformat()
    return f"侧脸黑白灰二次元头像_BO-{start}-BO-{end}_{stamp}"


def configure_batch_paths(start: int, count: int, batch_date: str | None = None) -> None:
    global BATCH_NAME, PROMPT_FILE, PRINT_DIR, MOCKUP_DIR, OUTPUT_XLSX
    BATCH_NAME = batch_name(start, count, batch_date)
    PROMPT_FILE = ROOT / "生成提示词" / f"{BATCH_NAME}.txt"
    PRINT_DIR = ROOT / "印花图_透明底" / BATCH_NAME
    MOCKUP_DIR = ROOT / "批量贴图结果" / f"{BATCH_NAME}_随机主图{count}"
    OUTPUT_XLSX = TEMPLATE_DIR / f"{BATCH_NAME}.xlsx"


configure_batch_paths(556, 50, "2026-06-11")

STORE_NAME = "YUHAOBO"
CATEGORY = "T恤"
HEADERS = ("店铺名称", "产品分类", "产品标题", "产品序列号", "颜色")
SELLING_POINTS = (
    "圆领短袖 透气微弹针织上衣 日常休闲百搭",
    "圆领短袖 透气舒适针织上衣 户外休闲日常百搭",
    "圆领短袖 轻薄透气针织上衣 休闲通勤日常百搭",
    "圆领短袖 柔软透气针织上衣 夏季日常百搭",
)


@dataclass(frozen=True)
class DesignSpec:
    title_word: str
    prompt: str


def sku_for(start: int, index: int) -> str:
    return f"BO-{start + index}"


def safe_filename(text: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if char in invalid or ord(char) < 32 else char for char in text)
    return re.sub(r"_+", "_", cleaned).strip(" ._")


def build_specs(count: int) -> list[DesignSpec]:
    hairs = [
        ("短发少女头像", "short black bob hair girl"),
        ("低马尾头像", "low ponytail girl"),
        ("齐刘海头像", "blunt bangs girl with tied hair"),
        ("层次短发头像", "layered short hair androgynous face"),
        ("长直发头像", "long straight dark hair behind neck"),
        ("编发头像", "single braid behind ear"),
        ("银灰短发头像", "silver gray short hair boyish face"),
        ("微卷短发头像", "slightly wavy short bob hair"),
        ("高马尾头像", "high ponytail girl"),
        ("耳侧碎发头像", "loose side strands around ear"),
    ]
    moods = [
        ("冷感", "cold aloof expression"),
        ("沉静", "quiet neutral expression"),
        ("清冷", "high-cold distant expression"),
        ("低调", "low key calm expression"),
        ("克制", "restrained expressionless mouth"),
    ]
    line_styles = [
        "bold manga ink with flat gray screenprint shadows",
        "clean manga line art with sparse gray hatching",
        "matte black ink outline with soft gray tone blocks",
        "fine monochrome manga lines with thick white outer sticker stroke",
        "high contrast black and gray screen print portrait decal",
    ]

    specs: list[DesignSpec] = []
    for index in range(count):
        hair_title, hair_prompt = hairs[index % len(hairs)]
        mood_title, mood_prompt = moods[(index // len(hairs)) % len(moods)]
        direction = "left" if index % 2 == 0 else "right"
        profile_detail = (
            "protruding nose silhouette, small closed mouth in profile, one visible ear"
            if direction == "left"
            else "clear nose bridge silhouette, thin neutral mouth in profile, one visible ear"
        )
        line_style = line_styles[index % len(line_styles)]
        title_word = f"{mood_title}{hair_title}"
        prompt = (
            "original monochrome anime manga avatar print, absolute pure 90 degree side profile view "
            f"facing {direction}, profile silhouette portrait, character looks away from viewer, "
            "only one side-view eye shape, no eye contact, no second eye, no second eyebrow, "
            f"{profile_detail}, {mood_prompt}, {hair_prompt}, black white gray only, no color, "
            f"{line_style}, clean black outline with white outer stroke, centered isolated head and neck only, "
            "generous empty margin, pure white background, transparent decal preparation, no front view, "
            "no three-quarter view, no symmetrical face, no shoulders, no body, no text, no symbols, "
            "no brand, no logo, no mockup, no shirt, no clothing, no background scene, "
            "no square panel, no rectangular background, no gray box, no framed portrait card"
        )
        specs.append(DesignSpec(title_word, prompt))
    return specs


def write_prompt_file(items: list[tuple[str, DesignSpec]]) -> None:
    PROMPT_FILE.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "侧脸黑白灰二次元头像印花 50 款，参考已筛选的精选组风格。",
        "要求：黑白灰、透明底准备、侧脸头像、无文字、无品牌、无版权角色，黑白 T 恤上都要清晰。",
        "",
    ]
    lines.extend(f"{sku}: {spec.prompt}" for sku, spec in items)
    PROMPT_FILE.write_text("\n".join(lines) + "\n", encoding="utf-8")


def generate_prints(items: list[tuple[str, DesignSpec]], args: argparse.Namespace) -> list[Path]:
    PRINT_DIR.mkdir(parents=True, exist_ok=True)
    if not args.regenerate_existing:
        existing_paths = [PRINT_DIR / f"{sku}.png" for sku, _ in items]
        if existing_paths and all(path.exists() for path in existing_paths):
            for sku, _ in items:
                print(f"{sku}: reuse {PRINT_DIR / f'{sku}.png'}")
            return existing_paths

    started = None
    if args.auto_start_comfyui:
        started = start_comfyui(args.comfy_timeout)
    else:
        wait_server()

    try:
        output_paths: list[Path] = []
        for index, (sku, spec) in enumerate(items):
            final_path = PRINT_DIR / f"{sku}.png"
            if final_path.exists() and not args.regenerate_existing:
                output_paths.append(final_path)
                print(f"{sku}: reuse {final_path}")
                continue

            workflow = make_workflow(spec.prompt, args.seed + index * 9973, args.steps, args.width, args.height)
            prompt_id = queue_prompt(workflow)
            print(f"{sku}: queued {spec.title_word}")
            history = wait_prompt(prompt_id, timeout_seconds=args.prompt_timeout)
            copied = copy_outputs(
                history,
                PRINT_DIR,
                sku,
                args.bg_threshold,
                args.bg_color_distance,
                keep_raw_comfy_output=False,
            )
            if len(copied) != 1:
                raise RuntimeError(f"{sku}: expected one generated output, got {len(copied)}")
            copied[0].replace(final_path)
            output_paths.append(final_path)
            print(f"{sku}: saved {final_path}")
        return output_paths
    finally:
        if not args.keep_comfyui:
            stop_comfyui(started)


def shirt_color(model_path: Path) -> tuple[str, str]:
    if "白" in model_path.stem:
        return "白色", "白"
    if "黑" in model_path.stem:
        return "黑色", "黑"

    img = ImageOps.exif_transpose(Image.open(model_path)).convert("RGB")
    w, h = img.size
    crop = img.crop((int(w * 0.32), int(h * 0.22), int(w * 0.68), int(h * 0.62)))
    arr = np.asarray(crop, dtype=np.float32)
    luma = 0.2126 * arr[..., 0] + 0.7152 * arr[..., 1] + 0.0722 * arr[..., 2]
    return ("白色", "白") if float(np.median(luma)) >= 150 else ("黑色", "黑")


def title_for(color_word: str, spec: DesignSpec, index: int) -> str:
    suffix = SELLING_POINTS[index % len(SELLING_POINTS)]
    return f"夏季{color_word}二次元黑白灰侧脸{spec.title_word}印花T恤 {suffix}"


def make_mockups(items: list[tuple[str, DesignSpec]], seed: int) -> list[Path]:
    models = list_images(MODEL_DIR)
    if not models:
        raise FileNotFoundError(f"No model images found: {MODEL_DIR}")

    rng = random.Random(seed)
    MOCKUP_DIR.mkdir(parents=True, exist_ok=True)
    placement = Placement(
        center_x=0.50,
        center_y=0.42,
        width=0.28,
        opacity=0.96,
        rotation=0.0,
        shadow_strength=0.28,
        wave_strength=0.006,
        remove_white_bg=False,
    )

    output_paths: list[Path] = []
    for index, (sku, spec) in enumerate(items):
        print_path = PRINT_DIR / f"{sku}.png"
        if not print_path.exists():
            raise FileNotFoundError(f"Missing print: {print_path}")
        model_path = rng.choice(models)
        color_word, _ = shirt_color(model_path)
        title = title_for(color_word, spec, index)
        output_path = MOCKUP_DIR / safe_filename(f"{sku}_{title}.png")
        composite_one(model_path, print_path, output_path, placement)
        output_paths.append(output_path)
        print(f"{sku}: mockup {output_path.name}")
    return output_paths


def make_overview(paths: list[Path], output_path: Path) -> None:
    tile_w, tile_h = 230, 310
    cols = 10
    rows = (len(paths) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tile_w, rows * tile_h), (235, 235, 235))
    for index, path in enumerate(paths):
        img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
        img.thumbnail((210, 245), Image.Resampling.LANCZOS)
        tile = Image.new("RGB", (tile_w, tile_h), "white")
        tile.paste(img, ((tile_w - img.width) // 2, 8))
        draw = ImageDraw.Draw(tile)
        draw.text((8, 260), path.stem[:30], fill=(0, 0, 0))
        sheet.paste(tile, ((index % cols) * tile_w, (index // cols) * tile_h))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output_path, quality=92)


def rows_from_product_filenames(mockup_dir: Path) -> list[tuple[str, str, str, str, str]]:
    pattern = re.compile(r"^(BO-\d+)_(.+)\.png$", re.IGNORECASE)
    rows: list[tuple[str, str, str, str, str]] = []
    for path in list_images(mockup_dir):
        if path.name.startswith("_"):
            continue
        match = pattern.match(path.name)
        if not match:
            continue
        sku, title = match.groups()
        color = "白" if "白色" in title else "黑" if "黑色" in title else ""
        rows.append((STORE_NAME, CATEGORY, title, sku, color))
    rows.sort(key=lambda row: int(row[3].split("-")[1]))
    return rows


def find_template_xlsx() -> Path:
    candidates = sorted(path for path in TEMPLATE_DIR.glob("*.xlsx") if not path.name.startswith("~$"))
    if not candidates:
        raise FileNotFoundError(f"No xlsx template found in {TEMPLATE_DIR}")
    return candidates[0]


def write_xlsx_from_filenames(expected_count: int) -> None:
    rows = rows_from_product_filenames(MOCKUP_DIR)
    if len(rows) != expected_count:
        raise RuntimeError(f"Expected {expected_count} product rows, got {len(rows)}")

    template = find_template_xlsx()
    workbook = load_workbook(template)
    sheet = workbook.worksheets[0]
    for col, header in enumerate(HEADERS, start=1):
        sheet.cell(row=1, column=col, value=header)
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append(row)
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_XLSX)


def validate_outputs(start: int, count: int) -> None:
    expected = [f"BO-{start + i}" for i in range(count)]
    prints = sorted(path for path in PRINT_DIR.glob("BO-*.png") if path.is_file())
    products = sorted(path for path in MOCKUP_DIR.glob("BO-*.png") if path.is_file())
    rows = rows_from_product_filenames(MOCKUP_DIR)
    print_skus = [path.stem for path in prints]
    product_skus = [path.name.split("_", 1)[0] for path in products]
    row_skus = [row[3] for row in rows]
    if print_skus != expected:
        raise RuntimeError(f"Print SKU mismatch: {print_skus[:3]} ... {print_skus[-3:]}")
    if product_skus != expected:
        raise RuntimeError(f"Product SKU mismatch: {product_skus[:3]} ... {product_skus[-3:]}")
    if row_skus != expected:
        raise RuntimeError(f"XLSX row SKU mismatch: {row_skus[:3]} ... {row_skus[-3:]}")

    workbook = load_workbook(OUTPUT_XLSX, read_only=True)
    sheet = workbook.worksheets[0]
    headers = tuple(sheet.cell(row=1, column=col).value for col in range(1, 6))
    if headers != HEADERS:
        raise RuntimeError(f"Header mismatch: {headers}")
    if sheet.max_row != count + 1:
        raise RuntimeError(f"XLSX row count mismatch: {sheet.max_row}")
    workbook.close()

    print(f"Validated prints={len(prints)}, products={len(products)}, xlsx_rows={count + 1}")
    print(f"First SKU={expected[0]}, last SKU={expected[-1]}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate 50 BO side-profile black-white-gray anime prints, mockups, and xlsx.")
    parser.add_argument("--start", type=int, default=556)
    parser.add_argument("--count", type=int, default=50)
    parser.add_argument("--seed", type=int, default=2026061103)
    parser.add_argument("--steps", type=int, default=24)
    parser.add_argument("--width", type=int, default=1024)
    parser.add_argument("--height", type=int, default=1024)
    parser.add_argument("--bg-threshold", type=int, default=244)
    parser.add_argument("--bg-color-distance", type=float, default=42.0)
    parser.add_argument("--auto-start-comfyui", action="store_true")
    parser.add_argument("--comfy-timeout", type=int, default=240)
    parser.add_argument("--prompt-timeout", type=int, default=900)
    parser.add_argument("--keep-comfyui", action="store_true")
    parser.add_argument("--regenerate-existing", action="store_true")
    parser.add_argument("--only-skus", default="", help="Comma-separated SKUs to regenerate, for example BO-584,BO-599.")
    parser.add_argument("--date", default=None, help="Batch date stamp, defaults to today (YYYY-MM-DD).")
    parser.add_argument("--prints-only", action="store_true", help="Only generate print PNGs; skip mockups, xlsx, and validation.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    configure_batch_paths(args.start, args.count, args.date)

    specs = build_specs(args.count)
    items = [(sku_for(args.start, index), specs[index]) for index in range(args.count)]
    write_prompt_file(items)
    generate_items = items
    if args.only_skus.strip():
        wanted = {value.strip().upper() for value in args.only_skus.split(",") if value.strip()}
        generate_items = [(sku, spec) for sku, spec in items if sku.upper() in wanted]
        missing = wanted - {sku.upper() for sku, _ in generate_items}
        if missing:
            raise ValueError(f"Unknown SKU(s): {sorted(missing)}")

    started_at = time.time()
    generate_prints(generate_items, args)
    if args.prints_only:
        print(f"Generated print-only batch: {PRINT_DIR}")
        print(f"Elapsed seconds: {time.time() - started_at:.1f}")
        return 0
    mockups = make_mockups(items, args.seed)
    make_overview(mockups, MOCKUP_DIR / "_overview.jpg")
    write_xlsx_from_filenames(args.count)
    validate_outputs(args.start, args.count)
    print(f"Prompt file: {PROMPT_FILE}")
    print(f"Print dir: {PRINT_DIR}")
    print(f"Product dir: {MOCKUP_DIR}")
    print(f"XLSX: {OUTPUT_XLSX}")
    print(f"Elapsed seconds: {time.time() - started_at:.1f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
