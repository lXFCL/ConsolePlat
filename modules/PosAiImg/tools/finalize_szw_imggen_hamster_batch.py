from __future__ import annotations

import argparse
import math
import re
import random
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageOps

try:
    from tshirt_print_tool import Placement, composite_one, list_images
except ModuleNotFoundError:
    from tools.tshirt_print_tool import Placement, composite_one, list_images


ROOT = Path(__file__).resolve().parents[1]
MODEL_DIR = ROOT / "模特图-干净"
STORE_NAME = "YUHOOBO"
PRODUCT_CATEGORY = "T恤"
PRODUCT_TITLE = "男士休闲，女士时尚，男女百搭，情侣装，夏季速干2D印花T恤 休闲圆领罗纹短袖 透气针织上衣 户外跑步健身沙滩度假日常百搭T恤，宽松，高级，舒适"
HEADERS = ["店铺名称", "产品分类", "产品标题", "产品序列号", "颜色"]
STYLE_NAME = "韩系仓鼠短句印花"


@dataclass(frozen=True)
class BatchPaths:
    print_dir: Path
    product_dir: Path
    product_test_dir: Path
    xlsx_path: Path


def parse_sku(value: str) -> tuple[str, int]:
    match = re.fullmatch(r"([A-Za-z]+)-(\d+)", value.strip())
    if not match:
        raise ValueError(f"Invalid SKU: {value}")
    return match.group(1).upper(), int(match.group(2))


def safe_filename(text: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if char in invalid or ord(char) < 32 else char for char in text)
    return re.sub(r"_+", "_", cleaned).strip(" ._")


def batch_name(prefix: str, start: int, count: int, stamp: str) -> str:
    return f"{STYLE_NAME}_{prefix}-{start}-{prefix}-{start + count - 1}_{stamp}"


def default_print_dir(prefix: str, start: int, count: int, stamp: str) -> Path:
    year = stamp.split("-")[0]
    month = f"{int(stamp.split('-')[1])}月"
    return ROOT / "图库" / prefix / year / month / batch_name(prefix, start, count, stamp) / "最终透明底"


def make_paths(prefix: str, start: int, count: int, stamp: str, print_dir: Path | None) -> BatchPaths:
    year = stamp.split("-")[0]
    month = f"{int(stamp.split('-')[1])}月"
    name = batch_name(prefix, start, count, stamp)
    product_batch = ROOT / "批量贴图结果" / prefix / year / month / name
    return BatchPaths(
        print_dir=print_dir or default_print_dir(prefix, start, count, stamp),
        product_dir=product_batch / "最终产品图",
        product_test_dir=product_batch / "测试",
        xlsx_path=ROOT / "衣物对应的xlsx" / prefix / f"{name}.xlsx",
    )


def sku_number(path: Path, prefix: str) -> int:
    match = re.fullmatch(rf"{re.escape(prefix)}-(\d+)\.png", path.name, re.IGNORECASE)
    if not match:
        raise ValueError(f"Unexpected print filename: {path.name}")
    return int(match.group(1))


def collect_prints(print_dir: Path, prefix: str, start: int, count: int) -> list[Path]:
    prints = sorted(
        [p for p in list_images(print_dir) if re.fullmatch(rf"{re.escape(prefix)}-\d+\.png", p.name, re.IGNORECASE)],
        key=lambda p: sku_number(p, prefix),
    )
    expected = list(range(start, start + count))
    got = [sku_number(p, prefix) for p in prints]
    if got != expected:
        raise RuntimeError(f"Print SKU range mismatch. expected={expected[:3]}...{expected[-3:]}, got_count={len(got)}, got={got[:5]}...{got[-5:] if got else []}")
    return prints


def white_model() -> Path:
    candidates = [MODEL_DIR / "主图5-白.jpg", MODEL_DIR / "主图6-白.jpg"]
    for path in candidates:
        if path.exists():
            return path
    white = [p for p in list_images(MODEL_DIR) if "白" in p.stem]
    if white:
        return white[0]
    raise FileNotFoundError(f"No white model image found in {MODEL_DIR}")


def black_model() -> Path:
    candidates = [MODEL_DIR / "主图1-黑.jpg", MODEL_DIR / "主图2-黑.png", MODEL_DIR / "主图3-黑.jpg", MODEL_DIR / "主图4-黑.jpg"]
    for path in candidates:
        if path.exists():
            return path
    black = [p for p in list_images(MODEL_DIR) if "黑" in p.stem]
    if black:
        return black[0]
    raise FileNotFoundError(f"No black model image found in {MODEL_DIR}")


def color_plan(prefix: str, start: int, count: int, black_count: int, seed: int) -> dict[str, str]:
    if not 0 <= black_count <= count:
        raise ValueError(f"black_count must be between 0 and {count}, got {black_count}")
    skus = [f"{prefix}-{number}" for number in range(start, start + count)]
    rng = random.Random(seed)
    black_skus = set(rng.sample(skus, black_count))
    return {sku: ("黑" if sku in black_skus else "白") for sku in skus}


def write_color_selection(paths: BatchPaths, plan: dict[str, str]) -> None:
    paths.product_test_dir.mkdir(parents=True, exist_ok=True)
    lines = ["sku\tcolor", *[f"{sku}\t{color}" for sku, color in plan.items()]]
    (paths.product_test_dir / "tshirt_color_selection.tsv").write_text("\n".join(lines) + "\n", encoding="utf-8")


def make_products(paths: BatchPaths, prefix: str, start: int, count: int, plan: dict[str, str]) -> list[Path]:
    prints = collect_prints(paths.print_dir, prefix, start, count)
    models = {"白": white_model(), "黑": black_model()}
    placement = Placement(
        center_x=0.50,
        center_y=0.42,
        width=0.28,
        opacity=0.96,
        rotation=0.0,
        shadow_strength=0.28,
        wave_strength=0.006,
        remove_white_bg=False,
    )
    paths.product_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for print_path in prints:
        sku = print_path.stem.upper()
        model = models[plan[sku]]
        output = paths.product_dir / safe_filename(f"{sku}_{PRODUCT_TITLE}.png")
        composite_one(model, print_path, output, placement)
        outputs.append(output)
    make_overview(outputs, paths.product_test_dir / "_overview.jpg")
    write_color_selection(paths, plan)
    return outputs


def make_overview(images: list[Path], output: Path, cols: int = 8) -> None:
    tile_w, tile_h = 250, 330
    rows = max(1, math.ceil(len(images) / cols))
    sheet = Image.new("RGB", (cols * tile_w, rows * tile_h), (238, 238, 238))
    try:
        label_font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 12)
    except OSError:
        label_font = ImageFont.load_default()
    for idx, path in enumerate(images):
        img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
        img.thumbnail((226, 268), Image.Resampling.LANCZOS)
        tile = Image.new("RGB", (tile_w, tile_h), "white")
        tile.paste(img, ((tile_w - img.width) // 2, 6))
        ImageDraw.Draw(tile).text((8, 284), path.stem[:34], fill=(0, 0, 0), font=label_font)
        sheet.paste(tile, ((idx % cols) * tile_w, (idx // cols) * tile_h))
    output.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output, quality=92)


def product_skus(product_dir: Path, prefix: str) -> list[int]:
    values: list[int] = []
    for path in list_images(product_dir):
        match = re.match(rf"^{re.escape(prefix)}-(\d+)_", path.name, re.IGNORECASE)
        if match:
            values.append(int(match.group(1)))
    return sorted(values)


def write_xlsx(paths: BatchPaths, prefix: str, start: int, count: int, plan: dict[str, str]) -> None:
    paths.xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADERS)
    for number in range(start, start + count):
        sku = f"{prefix}-{number}"
        ws.append([STORE_NAME, PRODUCT_CATEGORY, PRODUCT_TITLE, sku, plan[sku]])
    for col, width in {"A": 16, "B": 12, "C": 90, "D": 18, "E": 10}.items():
        ws.column_dimensions[col].width = width
    wb.save(paths.xlsx_path)


def validate_xlsx(xlsx_path: Path, prefix: str, start: int, count: int, expected_color_counts: dict[str, int] | None = None) -> dict:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0]) if rows else []
    data = rows[1:]
    summary = {
        "exists": xlsx_path.exists(),
        "rows_including_header": len(rows),
        "header": header,
        "first_sku": data[0][3] if data else "",
        "last_sku": data[-1][3] if data else "",
        "store_names": sorted({row[0] for row in data}),
        "titles": sorted({row[2] for row in data}),
        "colors": sorted({row[4] for row in data}),
    }
    expected_skus = [f"{prefix}-{n}" for n in range(start, start + count)]
    got_skus = [row[3] for row in data]
    if header != HEADERS or len(rows) != count + 1 or got_skus != expected_skus:
        raise RuntimeError(f"XLSX validation failed: {summary}")
    if summary["store_names"] != [STORE_NAME] or summary["titles"] != [PRODUCT_TITLE]:
        raise RuntimeError(f"XLSX content validation failed: {summary}")
    if expected_color_counts is not None:
        got_counts = {color: sum(1 for row in data if row[4] == color) for color in sorted({row[4] for row in data})}
        if got_counts != expected_color_counts:
            raise RuntimeError(f"XLSX color count validation failed: got={got_counts}, expected={expected_color_counts}")
    return summary


def validate_outputs(paths: BatchPaths, prefix: str, start: int, count: int, black_count: int) -> dict:
    expected = list(range(start, start + count))
    prints = collect_prints(paths.print_dir, prefix, start, count)
    products = [p for p in list_images(paths.product_dir) if p.name.startswith(f"{prefix}-")]
    product_nums = product_skus(paths.product_dir, prefix)
    bad_title_files = [p.name for p in products if not p.name.endswith(f"_{PRODUCT_TITLE}.png")]
    transparency = []
    for path in prints:
        alpha = Image.open(path).convert("RGBA").getchannel("A")
        transparency.append(alpha.getextrema()[0] == 0)
    expected_color_counts = {"白": count - black_count, "黑": black_count}
    xlsx_summary = validate_xlsx(paths.xlsx_path, prefix, start, count, expected_color_counts)
    selection_path = paths.product_test_dir / "tshirt_color_selection.tsv"
    selection_counts: dict[str, int] = {}
    if selection_path.exists():
        for line in selection_path.read_text(encoding="utf-8").splitlines()[1:]:
            if not line.strip():
                continue
            _, color = line.split("\t", 1)
            selection_counts[color] = selection_counts.get(color, 0) + 1
    summary = {
        "print_count": len(prints),
        "product_count": len(products),
        "print_skus": [sku_number(p, prefix) for p in prints],
        "product_skus": product_nums,
        "transparent_pngs_ok": all(transparency),
        "bad_title_file_count": len(bad_title_files),
        "overview_exists": (paths.product_test_dir / "_overview.jpg").exists(),
        "selection_counts": selection_counts,
        "xlsx": xlsx_summary,
    }
    if len(products) != count or product_nums != expected:
        raise RuntimeError(f"Product validation failed: {summary}")
    if not all(transparency) or bad_title_files:
        raise RuntimeError(f"Print/title validation failed: {summary}")
    if selection_counts != expected_color_counts:
        raise RuntimeError(f"Color selection validation failed: {summary}")
    if not summary["overview_exists"]:
        raise RuntimeError(f"Overview missing: {summary}")
    return summary


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Make product mockups and xlsx for the imgGen SZW hamster batch.")
    parser.add_argument("--start-sku", default="SZW-3013")
    parser.add_argument("--count", type=int, default=100)
    parser.add_argument("--stamp", default="2026-06-18")
    parser.add_argument("--print-dir", type=Path)
    parser.add_argument("--black-count", type=int, default=70)
    parser.add_argument("--color-seed", type=int, default=2026061901)
    parser.add_argument("--validate-only", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    prefix, start = parse_sku(args.start_sku)
    paths = make_paths(prefix, start, args.count, args.stamp, args.print_dir)
    products: list[Path] = []
    plan = color_plan(prefix, start, args.count, args.black_count, args.color_seed)
    if not args.validate_only:
        products = make_products(paths, prefix, start, args.count, plan)
        write_xlsx(paths, prefix, start, args.count, plan)
    validation = validate_outputs(paths, prefix, start, args.count, args.black_count)
    print(f"print_dir={paths.print_dir}")
    print(f"product_dir={paths.product_dir}")
    print(f"overview={paths.product_test_dir / '_overview.jpg'}")
    print(f"xlsx={paths.xlsx_path}")
    print(f"product_count={len(products) if products else validation['product_count']}")
    print(f"validation={validation}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
