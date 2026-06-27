from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Callable
from urllib.parse import quote

from openpyxl import load_workbook

from browser_session import BrowserSession
from app_paths import resource_path, runtime_root
from models import GoodsListResult, InventorySettingUploadResult, SpuIdentifierUploadResult


Logger = Callable[[str], None]

GOODS_LIST_URL = "https://agentseller.temu.com/goods/list"
SELLER_LOGIN_URL = "https://seller.kuajingmaihuo.com/settle/seller-login"
TEMPLATE_GROUP_URL = "https://agentseller.temu.com/sample/clothing-set"
COMPLIANCE_INFO_URL = "https://agentseller.temu.com/govern/information-supplementation"
COMPLIANT_LIVE_PHOTOS_URL = "https://agentseller.temu.com/govern/compliant-live-photos"
COMPLIANT_LIVE_PHOTOS_BATCH_URL = "https://agentseller.temu.com/govern/compliant-live-photos-batch"
JIT_PRODUCT_SELECT_URL = "https://agentseller.temu.com/newon/product-select"
STOCK_SALE_MANAGE_URL = "https://agentseller.temu.com/stock/fully-mgt/sale-manage/main"
TARGET_PAGE_SIZE = 50
JIT_TARGET_PAGE_SIZE = 10
JIT_PRODUCT_INFO_CONFIRM_PAGE_SIZE = 100

TEXT_GOODS_LIST = "商品列表"
TEXT_GOODS_INFO = "商品信息"
TEXT_ALL = "全部"
TEXT_TOTAL = "共有"
TEXT_MORE = "更多"
TEXT_BATCH_DELETE = "批量删除"
TEXT_BATCH_COPY_ID = "批量复制ID"
TEXT_SKC_ID = "SKC ID"
TEXT_SPU_ID = "SPU ID"
TEXT_SKU_ID = "SKU ID"
TEXT_TEMPLATE_GROUP = "套版组管理"
TEXT_ADD = "添加"
TEXT_ADD_SAME_MATERIAL = "添加同面料套版"
TEXT_SEARCH = "搜索"
TEXT_QUERY = "查询"
TEXT_SUBMIT = "提交"
TEXT_COMPLIANCE_INFO = "商品合规信息"
TEXT_BATCH_UPLOAD_COMPLIANCE = "批量上传合规信息"
TEXT_BATCH_UPLOAD_IDENTIFIER = "批量上传商品识别码"
TEXT_BATCH_IMPORT_INVENTORY = "批量导入库存"
TEXT_BATCH_IMPORT_INVENTORY_SETTING = "批量导入-库存设置"
TEXT_COMPLIANT_LIVE_PHOTOS = "商品实拍图"
TEXT_BATCH_UPLOAD_PACKAGE_PHOTOS = "具有相同包装的商品批量上传或更新"
TEXT_UPLOAD_LIVE_PHOTO = "上传实拍图"
TEXT_HISTORY = "历史记录"
TEXT_PRODUCT_CATEGORY = "商品所属类目"
TEXT_FRONT_VIEW = "正视图"
TEXT_SIDE_VIEW = "侧视图"
TEXT_CONFIRM_SUBMIT = "确认提交"
TEXT_NEWON_PRODUCT_SELECT = "上新生命周期管理"
TEXT_BATCH_ADJUST_JIT = "批量调整JIT"
TEXT_BATCH_OPEN_JIT = "批量开通JIT"
TEXT_OPEN_JIT = "开通JIT"
TEXT_PENDING_PRODUCT_INFO_CONFIRM = "商品信息待确认"
TEXT_BATCH_CONFIRM_PRODUCT_INFO = "批量确认商品信息"
TEXT_STOCK_SALE_MANAGE = "销售管理"
TEXT_EXPECTED_ARRIVAL_AREA_SETTING = "期望到货区域设置"
TEXT_MODIFY_EXPECTED_ARRIVAL_AREA = "修改期望到货区域"
TEXT_ALL_JIT_CUSTOM_PRODUCTS = "全部JIT/定制品"
TEXT_SELECTED_JIT_CUSTOM_PRODUCTS = "已选JIT/定制品"
TEXT_EXPECTED_ARRIVAL_AREA = "期望到货区域"
TEXT_YIWU = "义乌"
TEXT_HISTORY_NEAREST_ARRIVAL_AREA = "按照历史发货地就近推荐"
TEXT_UPDATE_SELECTED_PACKAGE_PHOTOS = "更新所选全部商品的包装实拍图"
TEXT_UPLOAD_FILE = "上传文件"
TEXT_START_IMPORT = "开始导入"
TEXT_IMPORT_CONFIRMED = "我已确认数据无误，确认导入"
TEXT_COMPLIANCE_MODAL_TITLE = "批量上传合规信息"
TEXT_IDENTIFIER_UPLOAD_MODAL_TITLE = "上传文件"
TEXT_COMPLIANCE_TYPE = "合规信息类型"
TEXT_CALIFORNIA_65 = "加州 65 号提案"
TEXT_MANUFACTURER_ATTRIBUTE = "制造商属性"
TEXT_PRODUCTION_SHELF_LIFE = "生产/保质期"
TEXT_WARNING_SAFETY_SUPPLEMENT = "警告或安全信息（补充）"
TEXT_PACKAGING_MATERIAL_INFO = "包装材料信息收集"
TEXT_TURKEY_RESPONSIBLE_PERSON = "土耳其负责人"
TEXT_MANUFACTURER_INFO = "制造商信息"
TEXT_EU_RESPONSIBLE_PERSON = "欧盟负责人"
TEXT_STATUS = "状态"
TEXT_PENDING_UPLOAD = "待上传"
TEXT_WARNING_TYPE = "警示类型"
TEXT_NO_WARNING = "No Warning Applicable/无需警示"
TEXT_AGREEMENT_SIGN = "协议签署"
TEXT_MANUFACTURER_IMPORTER_INFO = "制造商/进口商信息"
TEXT_PRODUCTION_SHELF_LIFE_INFO = "制造日期/保质期"
TEXT_NOT_APPLICABLE_PRODUCT = "该项目不适用该产品"
TEXT_TURKEY_RESPONSIBLE_PERSON_VALUE = "LİNOSAHİN DISTRIBUTION İHRACAT İTHALAT VE DİŞ TİCARET LTD ŞTİ"
TEXT_MANUFACTURER_INFO_VALUE = "Shangrao jiegi Clothing Co. Ltd."
TEXT_EU_RESPONSIBLE_PERSON_VALUE = "TOP PLUS SOLUTION LTD"
TEXT_CONFIRM_UPLOAD = "确认上传"
TEXT_CONFIRM = "确认"
TEXT_JIT_CONFIRM_OPTIONS = (TEXT_CONFIRM, "确定")
TEXT_CANCEL = "取消"
TEXT_I_KNOW = "我知道了"
COMPLIANCE_TEMPLATE_PATH = resource_path("合规相关资料") / "合规.xlsx"
INVENTORY_SETTING_TEMPLATE_PATH = resource_path("合规相关资料") / "库存设置.xlsx"
GENERATED_COMPLIANCE_DIR = runtime_root() / "合规相关资料" / "generated"
COMPLIANT_FRONT_IMAGE_PATH = resource_path("合规相关资料") / "商品正视图.jpg"
COMPLIANT_SIDE_IMAGE_PATH = resource_path("合规相关资料") / "商品侧视图.jpg"


def page_size_option_labels(size: int) -> list[str]:
    return [f"{size} 条/页", f"{size}条/页", str(size)]


def compliance_modal_container_selector() -> str:
    return (
        ".rocket-drawer-content-wrapper,.rocket-drawer-content,.rocket-drawer-body,.rocket-drawer,"
        "[role=\"dialog\"],[aria-modal=\"true\"],.rocket-modal,.rocket-modal-content"
    )


def jit_confirm_modal_selector() -> str:
    return (
        "[role=\"dialog\"],[aria-modal=\"true\"],.rocket-modal,.rocket-dialog,.rocket-modal-content,"
        ".MDL_outerWrapper_5-120-1,.MDL_container_5-120-1,.MDL_innerWrapper_5-120-1,"
        ".MDL_inner_5-120-1,.MDL_body_5-120-1,[class*=\"MDL_outerWrapper\"],"
        "[class*=\"MDL_container\"],[class*=\"MDL_innerWrapper\"],[class*=\"MDL_body\"]"
    )


class TemuGoodsList:
    def __init__(
        self,
        user_data_dir: str | Path,
        cdp_endpoint: str = "http://127.0.0.1:9222",
        logger: Logger | None = None,
    ) -> None:
        self.logger = logger or (lambda _message: None)
        self.session = BrowserSession(user_data_dir=user_data_dir, cdp_endpoint=cdp_endpoint, logger=self.logger)
        self.login_phone = ""
        self.login_password = ""

    def log(self, message: str) -> None:
        self.logger(message)

    def set_login_credentials(self, phone: str, password: str) -> None:
        self.login_phone = phone.strip()
        self.login_password = password

    def open_goods_list(self):
        self.session.connect_or_launch()
        page = self.session.page_for("agentseller.temu.com")
        if "agentseller.temu.com/goods/list" not in page.url:
            self.log("正在打开商品列表。")
            page.goto(GOODS_LIST_URL, wait_until="domcontentloaded")
        else:
            self.log("当前已在商品列表页。")
        page.wait_for_load_state("domcontentloaded", timeout=15000)
        self._wait_for_page_ready(page)
        page = self._ensure_authenticated_after_navigation(page, GOODS_LIST_URL)
        return page

    def prepare_and_collect_skc_ids(
        self,
        copy_to_clipboard: bool = True,
        target_page_size: int = TARGET_PAGE_SIZE,
    ) -> GoodsListResult:
        if target_page_size <= 0:
            raise RuntimeError("分页数量必须大于 0。")

        page = self.open_goods_list()
        self._ensure_goods_page(page)
        self._grant_clipboard_permissions(page)

        total_count = self._read_total_count(page)
        self.log(f"商品总数：{total_count if total_count >= 0 else '未知'}")

        page_size = self._read_page_size(page)
        if page_size != target_page_size:
            self.log(f"当前每页 {page_size or '未知'} 条，准备切换到 {target_page_size} 条。")
            self._set_page_size(page, target_page_size)
            page_size = self._read_page_size(page)
        else:
            self.log(f"当前已经是每页 {target_page_size} 条。")

        self._select_all_current_page(page)
        selected = self._is_header_checkbox_checked(page)
        if not selected:
            raise RuntimeError("已尝试全选，但未检测到全选状态。")
        self.log("已全选当前页商品。")

        self._copy_skc_ids_from_batch_menu(page)
        clipboard_text = self._read_clipboard_text(page)
        skc_ids = self._parse_copied_ids(clipboard_text)
        if not skc_ids:
            raise RuntimeError("已执行批量复制，但剪贴板中没有解析到 SKC ID。")
        self.log(f"已从批量复制结果解析 SKC ID：{len(skc_ids)} 个。")
        if len(skc_ids) != target_page_size:
            self.log(f"提示：当前解析到 {len(skc_ids)} 个，不等于 {target_page_size}，请检查页面选中数量。")

        return GoodsListResult(
            total_count=total_count,
            page_size=page_size,
            selected=selected,
            skc_ids=skc_ids,
            copied_to_clipboard=True,
            raw_skc_text=clipboard_text,
        )

    def prepare_and_collect_spu_ids(
        self,
        target_page_size: int = TARGET_PAGE_SIZE,
    ) -> GoodsListResult:
        if target_page_size <= 0:
            raise RuntimeError("分页数量必须大于 0。")

        page = self.open_goods_list()
        self._ensure_goods_page(page)
        self._grant_clipboard_permissions(page)

        total_count = self._read_total_count(page)
        self.log(f"商品总数：{total_count if total_count >= 0 else '未知'}")

        page_size = self._read_page_size(page)
        if page_size != target_page_size:
            self.log(f"当前每页 {page_size or '未知'} 条，准备切换到 {target_page_size} 条。")
            self._set_page_size(page, target_page_size)
            page_size = self._read_page_size(page)
        else:
            self.log(f"当前已经是每页 {target_page_size} 条。")

        self._select_all_current_page(page)
        selected = self._is_header_checkbox_checked(page)
        if not selected:
            raise RuntimeError("已尝试全选，但未检测到全选状态。")
        self.log("已全选当前页商品。")

        clipboard_text = ""
        spu_ids: list[str] = []
        last_error = ""
        for attempt in range(1, 4):
            try:
                self._copy_ids_from_batch_menu(page, TEXT_SPU_ID)
                clipboard_text = self._read_clipboard_text(page)
                spu_ids = self._parse_copied_ids(clipboard_text)
                if spu_ids:
                    break
                last_error = "剪贴板中没有解析到 SPU ID"
            except Exception as exc:  # noqa: BLE001
                last_error = str(exc)
            self.log(f"第 {attempt} 次复制 SPU ID 未成功，准备重试：{last_error}")
            page.wait_for_timeout(1000)
        if not spu_ids:
            raise RuntimeError(f"已重试复制 SPU ID，但仍未成功：{last_error}")
        self.log(f"已从批量复制结果解析 SPU ID：{len(spu_ids)} 个。")
        if len(spu_ids) != target_page_size:
            self.log(f"提示：当前解析到 {len(spu_ids)} 个，不等于 {target_page_size}，请检查页面选中数量。")

        return GoodsListResult(
            total_count=total_count,
            page_size=page_size,
            selected=selected,
            skc_ids=spu_ids,
            copied_to_clipboard=True,
            raw_skc_text=clipboard_text,
        )

    def prepare_and_collect_sku_ids(
        self,
        target_page_size: int = TARGET_PAGE_SIZE,
    ) -> GoodsListResult:
        if target_page_size <= 0:
            raise RuntimeError("分页数量必须大于 0。")

        page = self.open_goods_list()
        self._ensure_goods_page(page)
        self._grant_clipboard_permissions(page)

        total_count = self._read_total_count(page)
        self.log(f"商品总数：{total_count if total_count >= 0 else '未知'}")

        page_size = self._read_page_size(page)
        if page_size != target_page_size:
            self.log(f"当前每页 {page_size or '未知'} 条，准备切换到 {target_page_size} 条。")
            self._set_page_size(page, target_page_size)
            page_size = self._read_page_size(page)
        else:
            self.log(f"当前已经是每页 {target_page_size} 条。")

        self._select_all_current_page(page)
        selected = self._is_header_checkbox_checked(page)
        if not selected:
            raise RuntimeError("已尝试全选，但未检测到全选状态。")
        self.log("已全选当前页商品。")

        clipboard_text = ""
        sku_ids: list[str] = []
        last_error = ""
        for attempt in range(1, 4):
            try:
                self._copy_ids_from_batch_menu(page, TEXT_SKU_ID)
                clipboard_text = self._read_clipboard_text(page)
                sku_ids = self._parse_copied_ids(clipboard_text)
                if sku_ids:
                    break
                last_error = "剪贴板中没有解析到 SKU ID"
            except Exception as exc:  # noqa: BLE001
                last_error = str(exc)
            self.log(f"第 {attempt} 次复制 SKU ID 未成功，准备重试：{last_error}")
            page.wait_for_timeout(1000)
        if not sku_ids:
            raise RuntimeError(f"已重试复制 SKU ID，但仍未成功：{last_error}")
        self.log(f"已从批量复制结果解析 SKU ID：{len(sku_ids)} 个。")
        if len(sku_ids) != target_page_size:
            self.log(f"提示：当前解析到 {len(sku_ids)} 个，不等于 {target_page_size}，请检查页面选中数量。")

        return GoodsListResult(
            total_count=total_count,
            page_size=page_size,
            selected=selected,
            skc_ids=sku_ids,
            copied_to_clipboard=True,
            raw_skc_text=clipboard_text,
        )

    def upload_product_identifier_file(self, target_page_size: int = TARGET_PAGE_SIZE) -> SpuIdentifierUploadResult:
        result = self.prepare_and_collect_spu_ids(target_page_size=target_page_size)
        upload_file = self.generate_identifier_upload_file(result.skc_ids)
        page = self.open_compliance_info()
        self._ensure_compliance_info_page(page)
        self._open_identifier_upload_dialog(page)
        self._upload_identifier_file_to_dialog(page, upload_file)
        self._wait_for_identifier_file_parsed(page)
        self._submit_identifier_import(page)
        self.log(f"已导入商品识别码，SPU 数量：{result.skc_count} 个。")
        return SpuIdentifierUploadResult(
            spu_ids=result.skc_ids,
            upload_file=str(upload_file),
            imported=True,
        )

    def upload_inventory_setting_file(self, target_page_size: int = TARGET_PAGE_SIZE) -> InventorySettingUploadResult:
        result = self.prepare_and_collect_sku_ids(target_page_size=target_page_size)
        upload_file = self.generate_inventory_setting_upload_file(result.skc_ids)
        page = self.open_goods_list()
        self._ensure_goods_page(page)
        self._open_inventory_setting_upload_dialog(page)
        self._upload_identifier_file_to_dialog(page, upload_file)
        self._wait_for_inventory_setting_edit_modal(page)
        self._save_inventory_setting_import(page)
        self.log(f"库存设置文件已上传并保存，SKU 数量：{result.skc_count} 个。")
        return InventorySettingUploadResult(
            sku_ids=result.skc_ids,
            upload_file=str(upload_file),
            saved=True,
        )

    def upload_compliant_live_photos(self, submit: bool = True) -> int:
        self._ensure_compliant_live_photo_files()
        page = self.open_compliant_live_photos()
        self._ensure_compliant_live_photos_batch_page(page)
        self._select_first_live_photo_history(page)
        self._set_live_photo_page_size(page, 50)
        selected_count = self._select_all_live_photo_products(page)
        self._upload_live_photo_image(page, TEXT_FRONT_VIEW, COMPLIANT_FRONT_IMAGE_PATH)
        self._upload_live_photo_image(page, TEXT_SIDE_VIEW, COMPLIANT_SIDE_IMAGE_PATH)
        self._select_live_photo_update_all_option(page)
        if submit:
            self._submit_live_photo_upload(page)
            self.log(f"已提交商品合规图，选中商品：{selected_count} 个。")
        else:
            self.log(f"商品合规图调试流程已停在提交前，选中商品：{selected_count} 个。")
        return selected_count

    def batch_open_jit_management(self, submit: bool = True, target_page_size: int = JIT_TARGET_PAGE_SIZE) -> int:
        page = self.open_jit_product_select()
        self._ensure_jit_product_select_page(page)
        if self._has_jit_confirm_modal(page):
            modal_count = self._read_jit_confirm_modal_count(page)
            if submit:
                self._confirm_all_jit_open_modals(page)
                self.log(f"检测到JIT开通确认弹窗已打开，已直接确认，商品数量：{modal_count or '未知'}。")
            else:
                self.log(f"检测到JIT开通确认弹窗已打开，调试流程停在确认前，商品数量：{modal_count or '未知'}。")
            return modal_count
        self._confirm_pending_jit_product_info_if_needed(page)
        self._set_jit_product_page_size(page, target_page_size)
        selected_count = self._select_all_jit_products(page)
        self._open_batch_adjust_jit_menu(page)
        self._click_batch_open_jit(page)
        modal_count = self._wait_for_jit_confirm_modal(page)
        result_count = modal_count or selected_count
        if submit:
            self._confirm_all_jit_open_modals(page)
            self.log(f"已确认批量开通JIT管理，选中商品：{result_count} 个。")
        else:
            self.log(f"批量开通JIT管理调试流程已停在确认前，选中商品：{result_count} 个。")
        return result_count

    def batch_set_expected_arrival_area(self, submit: bool = True, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        page = self.open_stock_sale_manage()
        self._ensure_stock_sale_manage_page(page)
        if self._has_expected_arrival_confirm_modal(page):
            modal_count = self._read_expected_arrival_confirm_count(page)
            if submit:
                self._confirm_expected_arrival_area(page)
                self.log(f"检测到期望到货区域二次确认弹窗已打开，已直接确认，商品数量：{modal_count or '未知'}。")
            else:
                self.log(f"检测到期望到货区域二次确认弹窗已打开，调试流程停在确认前，商品数量：{modal_count or '未知'}。")
            return modal_count
        self._open_expected_arrival_area_drawer(page)
        self._set_expected_arrival_page_size(page, target_page_size)
        self._filter_expected_arrival_current_area(page, TEXT_HISTORY_NEAREST_ARRIVAL_AREA)
        selected_count = self._select_all_expected_arrival_products(page)
        self._select_expected_arrival_area(page, TEXT_YIWU)
        self._click_expected_arrival_drawer_confirm(page)
        modal_count = self._wait_for_expected_arrival_confirm_modal(page)
        result_count = modal_count or selected_count
        if submit:
            self._confirm_expected_arrival_area(page)
            self.log(f"已确认批量设置期望到货区域为{TEXT_YIWU}，选中商品：{result_count} 个。")
        else:
            self.log(f"期望到货区域调试流程已停在二次确认前，选中商品：{result_count} 个。")
        return result_count

    def generate_identifier_upload_file(self, spu_ids: list[str]) -> Path:
        if not spu_ids:
            raise RuntimeError("没有可写入 Excel 的 SPU ID。")
        if not COMPLIANCE_TEMPLATE_PATH.exists():
            raise RuntimeError(f"没有找到合规模板：{COMPLIANCE_TEMPLATE_PATH}")

        GENERATED_COMPLIANCE_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = GENERATED_COMPLIANCE_DIR / f"商品识别码上传_{timestamp}.xlsx"

        workbook = load_workbook(COMPLIANCE_TEMPLATE_PATH)
        worksheet = workbook.worksheets[0]
        identifier_value = worksheet.cell(row=2, column=2).value or "202499619"
        max_row = max(worksheet.max_row, len(spu_ids) + 1)
        for row in range(2, max_row + 1):
            worksheet.cell(row=row, column=1).value = None
            worksheet.cell(row=row, column=2).value = None
            worksheet.cell(row=row, column=3).value = None

        for index, spu_id in enumerate(spu_ids, start=2):
            worksheet.cell(row=index, column=1).value = int(spu_id) if spu_id.isdigit() else spu_id
            worksheet.cell(row=index, column=2).value = identifier_value
            worksheet.cell(row=index, column=3).value = None

        workbook.save(output_path)
        self.log(f"已生成商品识别码上传文件：{output_path}")
        return output_path

    def generate_inventory_setting_upload_file(self, sku_ids: list[str]) -> Path:
        if not sku_ids:
            raise RuntimeError("没有可写入 Excel 的 SKU ID。")
        if not INVENTORY_SETTING_TEMPLATE_PATH.exists():
            raise RuntimeError(f"没有找到库存设置模板：{INVENTORY_SETTING_TEMPLATE_PATH}")

        GENERATED_COMPLIANCE_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = GENERATED_COMPLIANCE_DIR / f"库存设置上传_{timestamp}.xlsx"

        workbook = load_workbook(INVENTORY_SETTING_TEMPLATE_PATH)
        worksheet = workbook.worksheets[0]
        inventory_value = worksheet.cell(row=3, column=4).value or 999
        max_row = max(worksheet.max_row, len(sku_ids) + 2)
        for row in range(3, max_row + 1):
            for column in range(1, worksheet.max_column + 1):
                worksheet.cell(row=row, column=column).value = None

        for index, sku_id in enumerate(sku_ids, start=3):
            worksheet.cell(row=index, column=1).value = int(sku_id) if sku_id.isdigit() else sku_id
            worksheet.cell(row=index, column=4).value = inventory_value

        workbook.save(output_path)
        self.log(f"已生成库存设置上传文件：{output_path}")
        return output_path

    def apply_to_template_group(self, skc_text: str) -> int:
        clean_text = " ".join(re.findall(r"\d{5,}", skc_text))
        if not clean_text:
            raise RuntimeError("没有可填写的 SKC ID，请先执行“批量复制SKC ID”。")

        page = self.open_template_group()
        self._ensure_template_group_page(page)
        self._click_first_add_button(page)
        self._fill_same_material_skc(page, clean_text)
        self._click_modal_button(page, TEXT_SEARCH)
        self._wait_for_page_ready(page, timeout=7000)
        added_count = self._wait_for_modal_search_results(page)
        self.log(f"搜索完成，下方结果区检测到 {added_count} 个同面料套版结果。")
        self._click_modal_button(page, TEXT_SUBMIT)
        self._wait_for_page_ready(page, timeout=7000)
        page.wait_for_timeout(1200)
        self.log("已点击提交。")
        return added_count

    def upload_california_65_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        page = self.open_compliance_info()
        self._ensure_compliance_info_page(page)
        self._click_batch_upload_compliance(page)
        self._select_compliance_modal_dropdown(page, TEXT_COMPLIANCE_TYPE, TEXT_CALIFORNIA_65)
        self._select_compliance_modal_dropdown(page, TEXT_STATUS, TEXT_PENDING_UPLOAD)
        self._click_compliance_modal_button(page, TEXT_SEARCH)
        self._wait_for_page_ready(page, timeout=7000)
        if self._has_empty_compliance_products(page):
            self._skip_empty_compliance_upload(0, "加州65号提案")
            return 0
        self._set_compliance_modal_page_size(page, target_page_size)
        selected_count = self._select_all_compliance_products(page)
        if self._skip_empty_compliance_upload(selected_count, "加州65号提案"):
            return selected_count
        self._select_compliance_modal_dropdown(page, TEXT_WARNING_TYPE, TEXT_NO_WARNING, option_fallback="无需警示")
        self._check_compliance_confirmation(page)
        self._submit_compliance_upload(page)
        self.log(f"已确认上传合规信息，选中商品：{selected_count} 个。")
        return selected_count

    def upload_manufacturer_attribute_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        return self._upload_korea_disclosure_compliance_info(
            compliance_type=TEXT_MANUFACTURER_ATTRIBUTE,
            field_label=TEXT_MANUFACTURER_IMPORTER_INFO,
            log_label="制造商属性",
            target_page_size=target_page_size,
        )

    def upload_production_shelf_life_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        return self._upload_korea_disclosure_compliance_info(
            compliance_type=TEXT_PRODUCTION_SHELF_LIFE,
            field_label=TEXT_PRODUCTION_SHELF_LIFE_INFO,
            log_label="生产/保质期",
            target_page_size=target_page_size,
        )

    def upload_warning_safety_supplement_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        return self._upload_korea_disclosure_compliance_info(
            compliance_type=TEXT_WARNING_SAFETY_SUPPLEMENT,
            field_label=TEXT_WARNING_SAFETY_SUPPLEMENT,
            log_label="警告或安全信息",
            target_page_size=target_page_size,
        )

    def upload_packaging_material_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        page = self.open_compliance_info()
        self._ensure_compliance_info_page(page)
        self._click_batch_upload_compliance(page)
        self._select_compliance_modal_dropdown(page, TEXT_COMPLIANCE_TYPE, TEXT_PACKAGING_MATERIAL_INFO)
        self._select_compliance_modal_dropdown(page, TEXT_STATUS, TEXT_PENDING_UPLOAD)
        self._click_compliance_modal_button(page, TEXT_SEARCH)
        self._wait_for_page_ready(page, timeout=7000)
        if self._has_empty_compliance_products(page):
            self._skip_empty_compliance_upload(0, "包装材料信息收集")
            return 0
        self._set_compliance_modal_page_size(page, target_page_size)
        selected_count = self._select_all_compliance_products(page)
        if self._skip_empty_compliance_upload(selected_count, "包装材料信息收集"):
            return selected_count
        self._fill_packaging_material_info(page)
        self._check_compliance_confirmation(page)
        self._submit_compliance_upload(page)
        self.log(f"已确认上传包装材料信息收集，选中商品：{selected_count} 个。")
        return selected_count

    def upload_turkey_responsible_person_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        page = self.open_compliance_info()
        self._ensure_compliance_info_page(page)
        self._click_batch_upload_compliance(page)
        self._select_compliance_modal_dropdown(page, TEXT_COMPLIANCE_TYPE, TEXT_TURKEY_RESPONSIBLE_PERSON)
        self._select_compliance_modal_dropdown(page, TEXT_STATUS, TEXT_PENDING_UPLOAD)
        self._click_compliance_modal_button(page, TEXT_SEARCH)
        self._wait_for_page_ready(page, timeout=7000)
        if self._has_empty_compliance_products(page):
            self._skip_empty_compliance_upload(0, "土耳其负责人")
            return 0
        self._set_compliance_modal_page_size(page, target_page_size)
        selected_count = self._select_all_compliance_products(page)
        if self._skip_empty_compliance_upload(selected_count, "土耳其负责人"):
            return selected_count
        self._select_compliance_loose_label_dropdown(
            page,
            TEXT_TURKEY_RESPONSIBLE_PERSON,
            TEXT_TURKEY_RESPONSIBLE_PERSON_VALUE,
            search_text="LİNO",
        )
        self._check_compliance_confirmation(page)
        self._submit_compliance_upload(page)
        self.log(f"已确认上传土耳其负责人，选中商品：{selected_count} 个。")
        return selected_count

    def upload_manufacturer_info_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        page = self.open_compliance_info()
        self._ensure_compliance_info_page(page)
        self._click_batch_upload_compliance(page)
        self._select_compliance_modal_dropdown(page, TEXT_COMPLIANCE_TYPE, TEXT_MANUFACTURER_INFO)
        self._select_compliance_modal_dropdown(page, TEXT_STATUS, TEXT_PENDING_UPLOAD)
        self._click_compliance_modal_button(page, TEXT_SEARCH)
        self._wait_for_page_ready(page, timeout=7000)
        if self._has_empty_compliance_products(page):
            self._skip_empty_compliance_upload(0, "制造商信息")
            return 0
        self._set_compliance_modal_page_size(page, target_page_size)
        selected_count = self._select_all_compliance_products(page)
        if self._skip_empty_compliance_upload(selected_count, "制造商信息"):
            return selected_count
        self._select_compliance_loose_label_dropdown(
            page,
            TEXT_MANUFACTURER_INFO,
            TEXT_MANUFACTURER_INFO_VALUE,
            search_text="Shangrao",
        )
        self._check_compliance_confirmation(page)
        self._submit_compliance_upload(page)
        self.log(f"已确认上传制造商信息，选中商品：{selected_count} 个。")
        return selected_count

    def upload_eu_responsible_person_compliance_info(self, target_page_size: int = TARGET_PAGE_SIZE) -> int:
        page = self.open_compliance_info()
        self._ensure_compliance_info_page(page)
        self._click_batch_upload_compliance(page)
        self._select_compliance_modal_dropdown(page, TEXT_COMPLIANCE_TYPE, TEXT_EU_RESPONSIBLE_PERSON)
        self._select_compliance_modal_dropdown(page, TEXT_STATUS, TEXT_PENDING_UPLOAD)
        self._click_compliance_modal_button(page, TEXT_SEARCH)
        self._wait_for_page_ready(page, timeout=7000)
        if self._has_empty_compliance_products(page):
            self._skip_empty_compliance_upload(0, "欧盟负责人")
            return 0
        self._set_compliance_modal_page_size(page, target_page_size)
        selected_count = self._select_all_compliance_products(page)
        if self._skip_empty_compliance_upload(selected_count, "欧盟负责人"):
            return selected_count
        self._select_compliance_loose_label_dropdown(
            page,
            TEXT_EU_RESPONSIBLE_PERSON,
            TEXT_EU_RESPONSIBLE_PERSON_VALUE,
            search_text="TOP PLUS",
        )
        self._check_compliance_confirmation(page)
        self._submit_compliance_upload(page)
        self.log(f"已确认上传欧盟负责人，选中商品：{selected_count} 个。")
        return selected_count

    def _upload_korea_disclosure_compliance_info(
        self,
        compliance_type: str,
        field_label: str,
        log_label: str,
        target_page_size: int = TARGET_PAGE_SIZE,
    ) -> int:
        page = self.open_compliance_info()
        self._ensure_compliance_info_page(page)
        self._click_batch_upload_compliance(page)
        self._select_compliance_modal_dropdown(page, TEXT_COMPLIANCE_TYPE, compliance_type)
        self._select_compliance_modal_dropdown(page, TEXT_STATUS, TEXT_PENDING_UPLOAD)
        self._click_compliance_modal_button(page, TEXT_SEARCH)
        self._wait_for_page_ready(page, timeout=7000)
        if self._has_empty_compliance_products(page):
            self._skip_empty_compliance_upload(0, log_label)
            return 0
        self._set_compliance_modal_page_size(page, target_page_size)
        selected_count = self._select_all_compliance_products(page)
        if self._skip_empty_compliance_upload(selected_count, log_label):
            return selected_count
        self._check_optional_compliance_checkbox_by_text(page, "我认可")
        self._select_compliance_modal_dropdown(
            page,
            field_label,
            TEXT_NOT_APPLICABLE_PRODUCT,
        )
        self._check_compliance_confirmation(page)
        self._submit_compliance_upload(page)
        self.log(f"已确认上传{log_label}，选中商品：{selected_count} 个。")
        return selected_count

    def open_template_group(self):
        self.session.connect_or_launch()
        page = self.session.page_for("agentseller.temu.com")
        if "agentseller.temu.com/sample/clothing-set" not in page.url:
            self.log("正在打开套版组管理。")
            page.goto(TEMPLATE_GROUP_URL, wait_until="domcontentloaded")
        else:
            self.log("当前已在套版组管理页。")
        page.wait_for_load_state("domcontentloaded", timeout=15000)
        self._wait_for_page_ready(page)
        page = self._ensure_authenticated_after_navigation(page, TEMPLATE_GROUP_URL)
        return page

    def open_compliance_info(self):
        self.session.connect_or_launch()
        page = self.session.page_for("agentseller.temu.com")
        if "agentseller.temu.com/govern/information-supplementation" not in page.url:
            self.log("正在打开商品合规信息页面。")
            page.goto(COMPLIANCE_INFO_URL, wait_until="domcontentloaded")
        else:
            self.log("当前已在商品合规信息页面。")
        page.wait_for_load_state("domcontentloaded", timeout=15000)
        self._wait_for_page_ready(page)
        page = self._ensure_authenticated_after_navigation(page, COMPLIANCE_INFO_URL)
        self._dismiss_existing_compliance_success_modal(page)
        return page

    def open_compliant_live_photos(self):
        self.session.connect_or_launch()
        page = self.session.page_for("agentseller.temu.com")
        if "agentseller.temu.com/govern/compliant-live-photos-batch" not in page.url:
            self.log("正在打开商品合规图批量上传页面。")
            page.goto(COMPLIANT_LIVE_PHOTOS_BATCH_URL, wait_until="domcontentloaded")
        else:
            self.log("当前已在商品合规图批量上传页面。")
        page.wait_for_load_state("domcontentloaded", timeout=15000)
        self._wait_for_page_ready(page)
        page = self._ensure_authenticated_after_navigation(page, COMPLIANT_LIVE_PHOTOS_BATCH_URL)
        return page

    def open_jit_product_select(self):
        self.session.connect_or_launch()
        page = self.session.page_for("agentseller.temu.com")
        if "agentseller.temu.com/newon/product-select" not in page.url:
            self.log("正在打开JIT商品选择页面。")
            page.goto(JIT_PRODUCT_SELECT_URL, wait_until="domcontentloaded")
        else:
            self.log("当前已在JIT商品选择页面。")
        page.wait_for_load_state("domcontentloaded", timeout=15000)
        self._wait_for_page_ready(page, timeout=12000)
        page = self._ensure_authenticated_after_navigation(page, JIT_PRODUCT_SELECT_URL)
        return page

    def open_stock_sale_manage(self):
        self.session.connect_or_launch()
        page = self.session.page_for("agentseller.temu.com")
        if "agentseller.temu.com/stock/fully-mgt/sale-manage/main" not in page.url:
            self.log("正在打开销售管理库存页面。")
            page.goto(STOCK_SALE_MANAGE_URL, wait_until="domcontentloaded")
        else:
            self.log("当前已在销售管理库存页面。")
        page.wait_for_load_state("domcontentloaded", timeout=15000)
        self._wait_for_page_ready(page, timeout=12000)
        page = self._ensure_authenticated_after_navigation(page, STOCK_SALE_MANAGE_URL)
        return page

    def close(self) -> None:
        self.session.close()

    def dismiss_transient_site_popups(self) -> int:
        try:
            self.session.connect_or_launch()
            page = self.session.page_for("agentseller.temu.com")
        except Exception:  # noqa: BLE001 - popup cleanup must never block the main task.
            return 0
        closed_count = 0
        for _attempt in range(4):
            if self._has_jit_confirm_modal(page):
                break
            if self._has_compliance_upload_modal(page) or self._has_identifier_upload_dialog(page):
                break
            if self._has_inventory_setting_edit_modal(page):
                break
            if self._has_compliance_success_modal(page):
                if not self._close_compliance_success_modal(page, timeout=2500):
                    break
                closed_count += 1
                continue
            if not self._click_transient_site_popup_close(page):
                break
            closed_count += 1
        return closed_count

    def _wait_for_page_ready(self, page, timeout: int = 8000) -> None:
        try:
            page.wait_for_load_state("networkidle", timeout=timeout)
        except Exception:  # noqa: BLE001
            pass
        page.wait_for_timeout(500)

    def _restart_auth_flow_page(self, page, target_url: str):
        self.log("当前登录页状态不稳定，正在重新打开一个干净页面。")
        try:
            page.close()
        except Exception:  # noqa: BLE001
            pass
        new_page = self.session.page_for("agentseller.temu.com").context.new_page()
        new_page.goto(self._seller_login_url_for(target_url), wait_until="domcontentloaded")
        self._wait_for_page_ready(new_page, timeout=15000)
        return new_page

    def _ensure_authenticated_after_navigation(self, page, target_url: str):
        for _attempt in range(5):
            self._recover_blank_login_app_if_needed(page)
            if self._is_login_authorization_page(page):
                self._check_login_authorization(page)
                if not self._click_login_continue_button(page):
                    raise RuntimeError("没有找到或无法点击“确认授权并前往”按钮。")
                self.log("已点击确认授权并前往。")
                try:
                    page.wait_for_function(
                        "() => !String(location.href).includes('/settle/seller-login')",
                        timeout=20000,
                    )
                except Exception:  # noqa: BLE001
                    self.log("授权后页面跳转较慢，继续等待后续页面状态。")
                self._wait_for_page_ready(page, timeout=15000)
                if self._is_login_authorization_page(page) or self._is_login_page(page):
                    page = self._restart_auth_flow_page(page, target_url)
                    continue
                continue
            if self._is_login_page(page):
                self._login_with_saved_credentials(page)
                self._wait_for_page_ready(page, timeout=15000)
                if self._is_login_authorization_page(page) or self._is_login_page(page):
                    page = self._restart_auth_flow_page(page, target_url)
                    continue
                continue
            if self._enter_seller_center_if_needed(page, target_url):
                self._wait_for_page_ready(page, timeout=15000)
                continue
            if target_url and "agentseller.temu.com" in page.url and target_url not in page.url:
                try:
                    page.goto(target_url, wait_until="domcontentloaded")
                    self._wait_for_page_ready(page, timeout=12000)
                    continue
                except Exception:  # noqa: BLE001
                    pass
            break
        return page

    def _recover_blank_login_app_if_needed(self, page) -> None:
        if "seller.kuajingmaihuo.com/settle/seller-login" not in page.url:
            return
        try:
            is_blank_app = bool(
                page.evaluate(
                    """
                    () => {
                      const text = String(document.body.innerText || '').trim();
                      return !text || text.includes('You need to enable JavaScript to run this app.');
                    }
                    """
                )
            )
        except Exception:  # noqa: BLE001
            return
        if not is_blank_app:
            return
        self.log("登录页应用未完全加载，正在刷新登录页。")
        try:
            page.reload(wait_until="domcontentloaded", timeout=60000)
        except Exception:  # noqa: BLE001
            self.log("刷新登录页耗时较长，继续等待页面稳定。")
        self._wait_for_page_ready(page, timeout=15000)

    def _seller_login_url_for(self, target_url: str) -> str:
        return f"{SELLER_LOGIN_URL}?redirectUrl={quote(target_url, safe='')}"

    def _is_login_page(self, page) -> bool:
        url = page.url
        if "seller.kuajingmaihuo.com/settle/seller-login" in url:
            try:
                return bool(
                    page.evaluate(
                        """
                        () => {
                          return Boolean(document.querySelector('#usernameId,input[name="usernameId"],#passwordId,input[type="password"]'));
                        }
                        """
                    )
                )
            except Exception:  # noqa: BLE001
                return False
        try:
            return bool(
                page.evaluate(
                    """
                    () => {
                      const text = String(document.body.innerText || '').replace(/\\s+/g, ' ');
                      return text.includes('授权登录')
                        && text.includes('手机号')
                        && text.includes('密码');
                    }
                    """
                )
            )
        except Exception:  # noqa: BLE001
            return False

    def _is_login_authorization_page(self, page) -> bool:
        if "seller.kuajingmaihuo.com/settle/seller-login" not in page.url:
            return False
        try:
            return bool(
                page.evaluate(
                    """
                    () => {
                      const text = String(document.body.innerText || '').replace(/\\s+/g, ' ');
                      return text.includes('确认授权并前往')
                        || (text.includes('即将前往') && text.includes('店铺'));
                    }
                    """
                )
            )
        except Exception:  # noqa: BLE001
            return False

    def _login_with_saved_credentials(self, page) -> None:
        if not self.login_phone or not self.login_password:
            raise RuntimeError("检测到需要重新登录，但软件中没有设置登录账号或密码。")
        self.log("检测到登录页面，正在自动填写账号密码。")
        if not self._fill_login_input(page, "phone", self.login_phone):
            raise RuntimeError("没有找到登录手机号输入框。")
        if not self._fill_login_input(page, "password", self.login_password):
            raise RuntimeError("没有找到登录密码输入框。")
        self._check_login_authorization(page)
        if not self._click_login_authorize_button(page):
            raise RuntimeError("没有找到或无法点击登录按钮。")
        self.log("已点击登录按钮。")
        try:
            page.wait_for_function(
                """
                () => !String(location.href).includes('/settle/seller-login')
                """,
                timeout=20000,
            )
        except Exception:  # noqa: BLE001
            self.log("登录后页面跳转较慢，继续等待后续页面状态。")
        page.wait_for_timeout(1500)

    def _fill_login_input(self, page, field: str, value: str) -> bool:
        rect = None
        end_time = page.evaluate("Date.now()") + 8000
        while page.evaluate("Date.now()") < end_time:
            rect = page.evaluate(
                """
                ([field]) => {
                  const exactSelector = field === 'password' ? '#passwordId' : '#usernameId';
                  const exactInput = document.querySelector(exactSelector);
                  if (isUsable(exactInput)) return inputRect(exactInput);

                  const inputs = [...document.querySelectorAll('input')]
                    .filter(isUsable);
                  const candidates = inputs
                    .filter(input => {
                      const type = String(input.type || '').toLowerCase();
                      const placeholder = String(input.getAttribute('placeholder') || '');
                      const aria = String(input.getAttribute('aria-label') || '');
                      const name = String(input.getAttribute('name') || '');
                      const id = String(input.id || '');
                      const near = String(input.closest('label,div,section,form')?.innerText || '');
                      if (field === 'password') {
                        return id === 'passwordId'
                          || type === 'password'
                          || placeholder.includes('密码')
                          || aria.includes('密码')
                          || near.includes('密码');
                      }
                      return type !== 'password'
                        && type !== 'checkbox'
                        && (
                          id === 'usernameId'
                          || name === 'usernameId'
                          || placeholder.includes('手机号')
                          || placeholder.includes('手机号码')
                          || aria.includes('手机')
                          || near.includes('手机')
                        );
                    })
                    .sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
                  const input = candidates[0];
                  return input ? inputRect(input) : null;

                  function inputRect(input) {
                    input.scrollIntoView({ block: 'center', inline: 'center' });
                    const rect = input.getBoundingClientRect();
                    return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };
                  }
                  function isUsable(input) {
                    if (!input || input.disabled || input.readOnly || input.getAttribute('readonly') !== null) return false;
                    const rect = input.getBoundingClientRect();
                    const style = getComputedStyle(input);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                [field],
            )
            if rect:
                break
            page.wait_for_timeout(300)
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        page.keyboard.press("Control+A")
        page.keyboard.type(value, delay=25)
        page.wait_for_timeout(300)
        return True

    def _check_login_authorization(self, page) -> None:
        rect = page.evaluate(
            """
            () => {
              const targetText = '您授权您的账号ID和店铺名称在卖家中心各板块共享';
              const authLabel = [...document.querySelectorAll('label,[data-testid="beast-core-checkbox"]')]
                .filter(visible)
                .filter(el => ancestorText(el).includes(targetText) || ancestorText(el).includes('隐私政策'))
                .sort((a, b) => area(a) - area(b))[0];
              if (authLabel) {
                const input = authLabel.querySelector('input[type="checkbox"]');
                if (input?.checked === true || authLabel.getAttribute('data-checked') === 'true') return { checked: true };
                const box = authLabel.querySelector('.CBX_squareInputWrapper_5-116-1,[class*="CBX_squareInputWrapper"],[data-testid="beast-core-checkbox-checkIcon"]') || authLabel;
                box.scrollIntoView({ block: 'center', inline: 'center' });
                const rect = box.getBoundingClientRect();
                return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };
              }

              const candidates = [...document.querySelectorAll('input[type="checkbox"],[role="checkbox"],label,span,div')]
                .filter(visible)
                .filter(el => {
                  const text = ancestorText(el);
                  return text.includes(targetText) || text.includes('隐私政策');
                })
                .sort((a, b) => checkboxScore(a) - checkboxScore(b));
              const target = candidates[0];
              if (!target) return null;
              const input = target.matches?.('input[type="checkbox"]')
                ? target
                : target.querySelector?.('input[type="checkbox"]')
                  || target.closest('label')?.querySelector?.('input[type="checkbox"]');
              if (input?.checked === true) return { checked: true };
              const clickable = input || target.closest('label') || target;
              const box = input || clickable;
              box.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = box.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function ancestorText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 6; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                }
                return chunks.join(' ');
              }
              function checkboxScore(el) {
                let score = 20;
                if (el.matches?.('input[type="checkbox"]')) score -= 10;
                if (el.tagName === 'LABEL') score -= 5;
                return score;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
            }
            """
        )
        if not rect or rect.get("checked"):
            return
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(300)

    def _click_login_authorize_button(self, page) -> bool:
        for button_text in ("授权登录", "登录"):
            if self._click_visible_text_as_button(page, button_text) or self._click_global_button_if_visible(page, button_text, timeout=2000):
                return True
        return False

    def _click_login_continue_button(self, page) -> bool:
        for button_text in ("确认授权并前往", "授权登录", "登录"):
            if (
                self._click_visible_text_as_button(page, button_text)
                or self._click_global_button_if_visible(page, button_text, timeout=2000)
                or self._click_visible_text_contains_as_button(page, button_text)
            ):
                return True
        return False

    def _click_visible_text_contains_as_button(self, page, text: str) -> bool:
        rect = page.evaluate(
            """
            ([targetText]) => {
              const target = [...document.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent).includes(targetText))
                .map(el => el.closest('button,[role="button"],a') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .sort((a, b) => area(a) - area(b))[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''));
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [text],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(700)
        return True

    def _enter_seller_center_if_needed(self, page, target_url: str = GOODS_LIST_URL) -> bool:
        if "agentseller.temu.com/auth/authentication" not in page.url:
            return False
        rect = page.evaluate(
            """
            () => {
              const candidates = [...document.querySelectorAll('a,button,[role="button"],div,span')]
                .filter(visible)
                .filter(el => {
                  const text = norm(el.innerText || el.textContent);
                  return text.includes('中国地区')
                    && text.includes('商家中心')
                    && !text.includes('其他地区')
                    && !text.includes('敬请期待');
                })
                .map(el => el.closest('a,button,[role="button"],.authentication_regionItem__g1pZV,[class*="authentication_regionItem"]') || el)
                .filter(visible)
                .filter(el => !/disabled/.test(String(el.className || '')))
                .filter((el, index, list) => list.indexOf(el) === index)
                .sort((a, b) => area(a) - area(b));
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const goto = [...target.querySelectorAll('a,button,[role="button"],div,span')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === '商家中心')
                .filter(el => getComputedStyle(el).cursor === 'pointer' || /goto|link|button/i.test(String(el.className || '')))
                .sort((a, b) => area(a) - area(b))[0] || target;
              for (const eventType of ['pointerover', 'mouseover', 'pointermove', 'mousemove']) {
                const rect = goto.getBoundingClientRect();
                goto.dispatchEvent(new MouseEvent(eventType, {
                  bubbles: true,
                  cancelable: true,
                  view: window,
                  clientX: rect.left + rect.width / 2,
                  clientY: rect.top + rect.height / 2,
                }));
              }
              const rect = goto.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """
        )
        if not rect:
            return False
        self.log("检测到商家中心选择页，正在进入中国地区商家中心。")
        page.mouse.move(rect["x"], rect["y"])
        page.mouse.down()
        page.mouse.up()
        try:
            page.wait_for_function(
                "() => !String(location.href).includes('/auth/authentication')",
                timeout=20000,
            )
        except Exception:  # noqa: BLE001
            self.log("进入商家中心后页面跳转较慢，继续等待。")
        page.wait_for_timeout(1500)
        if "agentseller.temu.com/auth/authentication" in page.url:
            self.log("商家中心选择页未跳转，准备重新授权登录。")
            try:
                page.goto(self._seller_login_url_for(target_url), wait_until="domcontentloaded", timeout=60000)
            except Exception:  # noqa: BLE001
                self.log("重新打开卖家登录页耗时较长，继续等待页面稳定。")
            self._wait_for_page_ready(page, timeout=12000)
            if self._is_login_page(page):
                self._login_with_saved_credentials(page)
        return True

    def _ensure_goods_page(self, page) -> None:
        if "agentseller.temu.com/goods/list" not in page.url:
            raise RuntimeError(f"当前不在商品列表页：{page.url}")
        try:
            page.wait_for_function(
                """
                ([goodsListText, goodsInfoText, moreText]) => {
                  const text = document.body.innerText || '';
                  const buttons = [...document.querySelectorAll('button,[role="button"]')]
                    .map(el => String(el.innerText || el.textContent || '').replace(/\\s+/g, ' ').trim());
                  const hasGoodsList = text.includes(goodsListText);
                  const hasGoodsInfo = text.includes(goodsInfoText);
                  const hasTable = document.querySelectorAll('table').length > 0;
                  const hasMore = buttons.includes(moreText);
                  const hasSkuSignal = text.includes('SKC ID') || text.includes('SKU ID');
                  return hasGoodsList && (hasGoodsInfo || hasTable || hasMore || hasSkuSignal);
                }
                """,
                arg=[TEXT_GOODS_LIST, TEXT_GOODS_INFO, TEXT_MORE],
                timeout=8000,
            )
        except Exception as exc:  # noqa: BLE001
            signals = self._read_page_signals(page)
            if signals.get("urlLooksRight") and signals.get("bodyLength", 0) > 0:
                self.log(f"页面内容检测信号较弱，但当前 URL 已在商品列表，继续尝试执行。信号：{signals}")
                return
            raise RuntimeError(f"未检测到商品列表内容。当前页面信号：{signals}") from exc

    def _grant_clipboard_permissions(self, page) -> None:
        try:
            page.context.grant_permissions(
                ["clipboard-read", "clipboard-write"],
                origin="https://agentseller.temu.com",
            )
        except Exception:  # noqa: BLE001
            pass

    def _read_total_count(self, page) -> int:
        text = self._body_text(page)
        match = re.search(rf"{TEXT_ALL}\s*(\d+)", text)
        if match:
            return int(match.group(1))
        match = re.search(rf"{TEXT_TOTAL}\s*(\d+)\s*条", text)
        return int(match.group(1)) if match else -1

    def _read_page_size(self, page) -> int:
        return int(
            page.evaluate(
                """
                () => {
                  const visible = el => {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  };
                  const values = [...document.querySelectorAll('input')]
                    .filter(visible)
                    .filter(input => /^\\d+$/.test(input.value))
                    .map(input => ({ input, value: Number(input.value), score: score(input) }))
                    .filter(item => item.value > 0 && item.value <= 500)
                    .sort((a, b) => b.score - a.score);
                  return Number(values[0]?.value || 0);

                  function score(input) {
                    const text = String(input.closest('div,li,span')?.innerText || '');
                    let value = 0;
                    if (['20', '50', '100', '200', '500'].includes(input.value)) value += 20;
                    if (text.includes('每页')) value += 80;
                    if (text.includes('条')) value += 60;
                    if (text.includes('页') && !text.includes('每页')) value -= 40;
                    const rect = input.getBoundingClientRect();
                    value += Math.max(0, rect.top / 1000);
                    return value;
                  }
                }
                """
            )
            or 0
        )

    def _set_page_size(self, page, size: int) -> None:
        rect = self._locate_page_size_control(page)
        if not rect:
            raise RuntimeError("没有找到分页条数控件。")

        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(300)
        option_rect = self._locate_dropdown_option(page, str(size))
        if option_rect:
            page.mouse.click(option_rect["x"], option_rect["y"])
        else:
            self.log(f"分页下拉中没有找到 {size} 条选项，尝试直接填写分页数量。")
            self._fill_page_size_control(page, size)
        page.wait_for_timeout(1500)
        if self._read_page_size(page) != size:
            raise RuntimeError(f"分页条数切换后未检测到 {size}。")

    def _locate_page_size_control(self, page):
        return page.evaluate(
            """
            () => {
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              };
              const inputs = [...document.querySelectorAll('input')]
                .filter(visible)
                .filter(input => /^\\d+$/.test(input.value))
                .filter(input => Number(input.value) > 0 && Number(input.value) <= 500)
                .sort((a, b) => score(b) - score(a));
              const input = inputs[0];
              if (!input) return null;
              input.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = input.getBoundingClientRect();
              if (!visible(input)) return null;
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function score(input) {
                const text = String(input.closest('div,li,span')?.innerText || '');
                let value = 0;
                if (['20', '50', '100', '200', '500'].includes(input.value)) value += 20;
                if (text.includes('每页')) value += 80;
                if (text.includes('条')) value += 60;
                if (text.includes('页') && !text.includes('每页')) value -= 40;
                const rect = input.getBoundingClientRect();
                value += Math.max(0, rect.top / 1000);
                return value;
              }
            }
            """
        )

    def _fill_page_size_control(self, page, size: int) -> None:
        rect = self._locate_page_size_control(page)
        if not rect:
            raise RuntimeError("没有找到可填写的分页条数控件。")
        page.mouse.click(rect["x"], rect["y"])
        page.keyboard.press("Control+A")
        page.keyboard.type(str(size), delay=20)
        page.keyboard.press("Enter")

    def _locate_dropdown_option(self, page, text: str):
        return page.evaluate(
            """
            (target) => {
              const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              };
              const candidates = [...document.querySelectorAll('li,[role="option"],div,span')]
                .filter(el => norm(el.innerText || el.textContent) === target)
                .filter(visible);
              const item = candidates
                .sort((a, b) => area(a) - area(b))[0];
              if (!item) return null;
              item.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = item.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            text,
        )

    def _select_all_current_page(self, page) -> None:
        self._scroll_to_goods_toolbar(page)
        if self._is_batch_delete_enabled(page):
            return

        for attempt in range(3):
            target = self._locate_visible_header_checkbox(page)
            if not target:
                raise RuntimeError("没有找到可见的表头全选框。")
            self.log(
                "准备点击全选框："
                f"x={target['x']:.1f}, y={target['y']:.1f}, "
                f"候选序号={target.get('index')}, 已选={target.get('checked')}"
            )
            if target.get("checked") or self._is_batch_delete_enabled(page):
                return

            page.mouse.move(target["x"], target["y"])
            page.mouse.click(target["x"], target["y"])
            page.wait_for_timeout(700)
            if self._is_batch_delete_enabled(page) or self._is_header_checkbox_checked(page):
                return

            self._click_checkbox_by_dom_index(page, int(target.get("index", 0)))
            page.wait_for_timeout(700)
            if self._is_batch_delete_enabled(page) or self._is_header_checkbox_checked(page):
                return

            self.log(f"第 {attempt + 1} 次点击全选后仍未检测到选中，准备重新定位。")

    def _scroll_to_goods_toolbar(self, page) -> None:
        page.evaluate(
            """
            ([allText, moreText]) => {
              const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const moreButton = [...document.querySelectorAll('button,[role="button"]')]
                .find(el => norm(el.innerText || el.textContent) === moreText);
              const allTab = [...document.querySelectorAll('a,div,span')]
                .find(el => norm(el.innerText || el.textContent).startsWith(allText));
              const target = moreButton || allTab;
              if (target) target.scrollIntoView({ block: 'center', inline: 'nearest' });
              const table = [...document.querySelectorAll('table')]
                .find(item => (item.innerText || '').includes('SKU ID'))
                || [...document.querySelectorAll('table')][0];
              if (table) {
                let current = table;
                while (current && current !== document.body) {
                  if (current.scrollWidth > current.clientWidth + 8) current.scrollLeft = 0;
                  current = current.parentElement;
                }
              }
            }
            """,
            [TEXT_ALL, TEXT_MORE],
        )
        page.wait_for_timeout(300)

    def _locate_visible_header_checkbox(self, page):
        return page.evaluate(
            """
            () => {
              const candidates = [...document.querySelectorAll(
                'label[data-testid="beast-core-checkbox"],label,input[type="checkbox"],[role="checkbox"],[data-testid="beast-core-checkbox"]'
              )]
                .map((el, index) => ({ el, index }))
                .filter(item => visible(item.el))
                .map(item => {
                  const target = item.el.closest('label') || item.el;
                  const rect = target.getBoundingClientRect();
                  const text = nearbyText(target);
                  return {
                    index: item.index,
                    x: rect.left + rect.width / 2,
                    y: rect.top + rect.height / 2,
                    left: rect.left,
                    top: rect.top,
                    width: rect.width,
                    height: rect.height,
                    checked: isChecked(item.el) || isChecked(target),
                    text,
                    score: scoreCandidate(rect, text),
                  };
                })
                .filter(item => item.width > 0 && item.height > 0)
                .filter(item => item.y > 250 && item.y < window.innerHeight - 40)
                .sort((a, b) => b.score - a.score || a.y - b.y || a.x - b.x);
              return candidates[0] || null;

              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden'
                  && rect.bottom >= 0
                  && rect.right >= 0
                  && rect.top <= window.innerHeight
                  && rect.left <= window.innerWidth;
              }
              function isChecked(el) {
                return el.getAttribute('data-checked') === 'true'
                  || el.getAttribute('aria-checked') === 'true'
                  || el.checked === true
                  || el.querySelector?.('input[type="checkbox"]')?.checked === true;
              }
              function nearbyText(el) {
                const row = el.closest('tr');
                const parent = el.parentElement;
                return String((row || parent || el).innerText || '').replace(/\\s+/g, ' ').trim().slice(0, 120);
              }
              function scoreCandidate(rect, text) {
                let score = 0;
                if (text.includes('商品信息')) score += 1000;
                if (!text.includes('SPU ID') && !text.includes('SKC ID')) score += 200;
                if (rect.top < 620) score += 200;
                if (rect.left < 520) score += 120;
                if (rect.left > 80) score += 80;
                score -= Math.abs(rect.top - 475) * 0.2;
                score -= rect.left * 0.02;
                return score;
              }
            }
            """
        )

    def _click_checkbox_by_dom_index(self, page, index: int) -> None:
        page.evaluate(
            """
            (targetIndex) => {
              const items = [...document.querySelectorAll(
                'label[data-testid="beast-core-checkbox"],label,input[type="checkbox"],[role="checkbox"],[data-testid="beast-core-checkbox"]'
              )];
              const el = items[targetIndex];
              if (!el) return false;
              const target = el.closest('label') || el;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                const rect = target.getBoundingClientRect();
                target.dispatchEvent(new MouseEvent(type, {
                  bubbles: true,
                  cancelable: true,
                  view: window,
                  clientX: rect.left + rect.width / 2,
                  clientY: rect.top + rect.height / 2,
                }));
              }
              target.click?.();
              return true;
            }
            """,
            index,
        )

    def _is_header_checkbox_checked(self, page) -> bool:
        if self._is_batch_delete_enabled(page):
            return True
        return bool(
            page.evaluate(
                """
                ([goodsInfoText]) => {
                  const visible = el => {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  };
                  const table = [...document.querySelectorAll('table')]
                    .find(item => (item.innerText || '').includes(goodsInfoText))
                    || [...document.querySelectorAll('table')][0];
                  const header = table && (table.querySelector('thead') || table.querySelector('tr'));
                  const checkbox = [...(header || document).querySelectorAll('label,input[type="checkbox"],[role="checkbox"],[data-testid="beast-core-checkbox"]')]
                    .find(visible);
                  return Boolean(checkbox && (
                    checkbox.getAttribute('data-checked') === 'true'
                    || checkbox.getAttribute('aria-checked') === 'true'
                    || checkbox.checked === true
                    || checkbox.querySelector?.('input[type="checkbox"]')?.checked === true
                  ));
                }
                """,
                [TEXT_GOODS_INFO],
            )
        )

    def _is_batch_delete_enabled(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([batchDeleteText]) => {
                  const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
                  const button = [...document.querySelectorAll('button,[role="button"]')]
                    .find(el => norm(el.innerText || el.textContent) === batchDeleteText);
                  if (!button) return false;
                  return button.disabled !== true
                    && button.getAttribute('disabled') === null
                    && button.getAttribute('aria-disabled') !== 'true';
                }
                """,
                [TEXT_BATCH_DELETE],
            )
        )

    def _copy_skc_ids_from_batch_menu(self, page) -> None:
        self._copy_ids_from_batch_menu(page, TEXT_SKC_ID)

    def _copy_ids_from_batch_menu(self, page, id_text: str) -> None:
        self.log(f"正在通过批量复制菜单复制 {id_text}。")
        if not self._click_goods_toolbar_more(page):
            raise RuntimeError("没有找到或无法点击“更多”按钮。")
        page.wait_for_timeout(500)

        if not self._activate_popup_item(page, TEXT_BATCH_COPY_ID, exact=False):
            raise RuntimeError("没有找到“批量复制ID”菜单项。")
        page.wait_for_timeout(500)

        if not self._activate_popup_item(page, id_text, exact=True):
            raise RuntimeError(f"没有找到“{id_text}”菜单项。")
        page.wait_for_timeout(900)

    def _click_button_by_text(self, page, text: str) -> bool:
        result = page.evaluate(
            """
            ([targetText]) => {
              const buttons = [...document.querySelectorAll('button,[role="button"]')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === targetText);
              const button = buttons[0];
              if (!button) return false;
              button.scrollIntoView({ block: 'center', inline: 'center' });
              return rectOf(button);

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function rectOf(el) {
                const rect = el.getBoundingClientRect();
                return { x: rect.left + 14, y: rect.top + rect.height / 2 };
              }
            }
            """,
            [text],
        )
        if result:
            page.mouse.click(result["x"], result["y"])
            return True
        return False

    def _click_goods_toolbar_more(self, page) -> bool:
        self._scroll_to_goods_action_area(page)
        result = page.evaluate(
            """
            ([moreText, batchDeleteText]) => {
              const buttons = [...document.querySelectorAll('button,[role="button"]')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === moreText)
                .map(el => ({ el, score: scoreButton(el, batchDeleteText) }))
                .filter(item => item.score > 0)
                .sort((a, b) => b.score - a.score || a.el.getBoundingClientRect().top - b.el.getBoundingClientRect().top);
              const button = buttons[0]?.el || null;
              if (!button) return null;
              button.scrollIntoView({ block: 'center', inline: 'center' });
              fireMouse(button);
              button.click?.();
              const rect = button.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function scoreButton(el, batchDeleteText) {
                const rect = el.getBoundingClientRect();
                const text = ancestorText(el);
                let score = 0;
                if (text.includes(batchDeleteText)) score += 1000;
                if (text.includes('下载查询结果')) score += 500;
                if (text.includes('新建商品')) score += 300;
                if (rect.top < window.innerHeight * 0.65) score += 100;
                if (text.includes('机会商品') || text.includes('发布同款')) score -= 1000;
                return score;
              }
              function ancestorText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                }
                return norm(chunks.join(' '));
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function fireMouse(el) {
                const rect = el.getBoundingClientRect();
                const clientX = rect.left + rect.width / 2;
                const clientY = rect.top + rect.height / 2;
                for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                  el.dispatchEvent(new MouseEvent(type, {
                    bubbles: true,
                    cancelable: true,
                    view: window,
                    clientX,
                    clientY,
                  }));
                }
              }
            }
            """,
            [TEXT_MORE, TEXT_BATCH_DELETE],
        )
        if not result:
            return False
        page.wait_for_timeout(300)
        return True

    def _scroll_to_goods_action_area(self, page) -> None:
        page.evaluate(
            """
            ([batchDeleteText]) => {
              const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const candidates = [...document.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent).includes(batchDeleteText))
                .sort((a, b) => area(a) - area(b));
              const target = candidates[0];
              if (target) target.scrollIntoView({ block: 'center', inline: 'nearest' });

              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_BATCH_DELETE],
        )
        page.wait_for_timeout(300)

    def _activate_popup_item(self, page, text: str, exact: bool) -> bool:
        result = page.evaluate(
            """
            ([targetText, exact]) => {
              const item = findPopupItem(targetText, exact);
              if (!item) return false;
              const target = item.closest('[role="option"],[role="menuitem"],li,button') || item;
              fireMouse(target);
              target.click?.();
              return true;

              function findPopupItem(targetText, exact) {
                const targetCompact = compact(targetText);
                const candidates = [...document.querySelectorAll('body *')]
                  .filter(visible)
                  .filter(isPopupRelated)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    return exact ? compact(text) === targetCompact : compact(text).includes(targetCompact);
                  })
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isPopupRelated(el) {
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  const style = getComputedStyle(node);
                  const cls = String(node.className || '');
                  const role = node.getAttribute?.('role') || '';
                  if (style.position === 'fixed' || style.position === 'absolute') return true;
                  if (/dropdown|popover|popper|portal|menu|cascader/i.test(cls)) return true;
                  if (/menu|menuitem|listbox|option/.test(role)) return true;
                }
                return false;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
              function fireMouse(el) {
                el.scrollIntoView({ block: 'center', inline: 'center' });
                const rect = el.getBoundingClientRect();
                const clientX = rect.left + rect.width / 2;
                const clientY = rect.top + rect.height / 2;
                for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                  el.dispatchEvent(new MouseEvent(type, {
                    bubbles: true,
                    cancelable: true,
                    view: window,
                    clientX,
                    clientY,
                  }));
                }
              }
            }
            """,
            [text, exact],
        )
        return bool(result)

    def _wait_for_popup_item(self, page, text: str, exact: bool, timeout: int = 3000) -> bool:
        try:
            page.wait_for_function(
                """
                ([targetText, exact]) => {
                  const targetCompact = compact(targetText);
                  return [...document.querySelectorAll('body *')]
                    .filter(visible)
                    .filter(isPopupRelated)
                    .some(el => {
                      const text = norm(el.innerText || el.textContent);
                      return exact ? compact(text) === targetCompact : compact(text).includes(targetCompact);
                    });

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isPopupRelated(el) {
                    let node = el;
                    for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                      const style = getComputedStyle(node);
                      const cls = String(node.className || '');
                      const role = node.getAttribute?.('role') || '';
                      if (style.position === 'fixed' || style.position === 'absolute') return true;
                      if (/dropdown|popover|popper|portal|menu|cascader/i.test(cls)) return true;
                      if (/menu|menuitem|listbox|option/.test(role)) return true;
                    }
                    return false;
                  }
                }
                """,
                arg=[text, exact],
                timeout=timeout,
            )
            return True
        except Exception:  # noqa: BLE001 - caller decides whether to retry.
            return False

    def _activate_visible_text_item(self, page, text: str, exact: bool = False) -> bool:
        result = page.evaluate(
            """
            ([targetText, exact]) => {
              const targetCompact = compact(targetText);
              const candidates = [...document.querySelectorAll('body *')]
                .filter(visible)
                .filter(isPopupRelated)
                .filter(el => {
                  const text = norm(el.innerText || el.textContent);
                  if (!text) return false;
                  return exact ? compact(text) === targetCompact : compact(text).includes(targetCompact);
                })
                .sort((a, b) => {
                  return area(a) - area(b);
                });
              const item = candidates[0];
              if (!item) return false;
              const target = item.closest('[role="option"],[role="menuitem"],li,button') || item;
              fireMouse(target);
              target.click?.();
              return true;

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isPopupRelated(el) {
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  const style = getComputedStyle(node);
                  const cls = String(node.className || '');
                  const role = node.getAttribute?.('role') || '';
                  if (style.position === 'fixed' || style.position === 'absolute') return true;
                  if (/dropdown|popover|popper|portal|menu|select|option/i.test(cls)) return true;
                  if (/menu|menuitem|listbox|option/.test(role)) return true;
                }
                return false;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
              function fireMouse(el) {
                el.scrollIntoView({ block: 'center', inline: 'center' });
                const rect = el.getBoundingClientRect();
                const clientX = rect.left + rect.width / 2;
                const clientY = rect.top + rect.height / 2;
                for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                  el.dispatchEvent(new MouseEvent(type, {
                    bubbles: true,
                    cancelable: true,
                    view: window,
                    clientX,
                    clientY,
                  }));
                }
              }
            }
            """,
            [text, exact],
        )
        return bool(result)

    def _read_clipboard_text(self, page) -> str:
        try:
            return str(
                page.evaluate(
                    """
                    async () => {
                      if (!navigator.clipboard || !navigator.clipboard.readText) return '';
                      return await navigator.clipboard.readText();
                    }
                    """
                )
                or ""
            )
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("批量复制后无法读取剪贴板内容。") from exc

    def _parse_copied_ids(self, text: str) -> list[str]:
        ids = re.findall(r"\d{5,}", text)
        result: list[str] = []
        seen: set[str] = set()
        for item in ids:
            if item not in seen:
                seen.add(item)
                result.append(item)
        return result

    def _ensure_template_group_page(self, page) -> None:
        if "agentseller.temu.com/sample/clothing-set" not in page.url:
            raise RuntimeError(f"当前不在套版组管理页：{page.url}")
        try:
            page.wait_for_function(
                """
                ([titleText, addText, sameMaterialText]) => {
                  const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
                  const visible = el => {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  };
                  const text = norm(document.body.innerText || '');
                  const visibleAdds = [...document.querySelectorAll('a,button,span,div')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent) === addText);
                  const urlLooksRight = location.href.includes('/sample/clothing-set');
                  const strongText =
                    text.includes(titleText)
                    || text.includes(sameMaterialText)
                    || text.includes('同面料套版')
                    || text.includes('底版SKC')
                    || text.includes('SKC ID');
                  return urlLooksRight
                    && text.length > 0
                    && (strongText || visibleAdds.length > 0 || document.querySelectorAll('table').length > 0);
                }
                """,
                arg=[TEXT_TEMPLATE_GROUP, TEXT_ADD, TEXT_ADD_SAME_MATERIAL],
                timeout=8000,
            )
        except Exception as exc:  # noqa: BLE001
            signals = self._read_template_group_signals(page)
            if signals.get("urlLooksRight") and signals.get("bodyLength", 0) > 0:
                self.log(f"套版组页面文字信号较弱，但 URL 和页面内容已就绪，继续尝试操作。信号：{signals}")
                return
            raise RuntimeError(f"未检测到套版组管理内容。页面信号：{signals}") from exc

    def _click_first_add_button(self, page) -> None:
        if self._has_same_material_modal(page):
            self.log("检测到添加同面料套版弹窗已打开，继续填写。")
            return

        rect = page.evaluate(
            """
            ([addText]) => {
              const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              };
              const candidates = [...document.querySelectorAll('a,button,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === addText)
                .map(el => el.closest('a,button') || el)
                .filter(visible)
                .sort((a, b) => {
                  const scoreA = score(a);
                  const scoreB = score(b);
                  if (scoreA !== scoreB) return scoreA - scoreB;
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function score(el) {
                const text = norm(el.closest('tr,table,div')?.innerText || '');
                let value = 100;
                if (el.closest('tr')) value -= 60;
                if (el.closest('table')) value -= 30;
                if (text.includes('底版SKC') || text.includes('同面料套版') || text.includes('SKC ID')) value -= 20;
                return value;
              }
            }
            """,
            [TEXT_ADD],
        )
        if not rect:
            raise RuntimeError("没有找到套版组管理中的“添加”按钮。")
        page.mouse.click(rect["x"], rect["y"])
        self._wait_for_page_ready(page, timeout=5000)
        page.wait_for_timeout(600)
        page.wait_for_function(
            """
            ([modalTitle]) => {
              return Boolean(findModal(modalTitle));

              function findModal(title) {
                const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                  .filter(visible)
                  .filter(el => {
                    const text = String(el.innerText || el.textContent || '');
                    if (!text.includes(title)) return false;
                    if (![...el.querySelectorAll('textarea,input')].some(visible)) return false;
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    const cls = String(el.className || '');
                    const role = el.getAttribute?.('role') || '';
                    const modalLike = role === 'dialog'
                      || el.getAttribute('aria-modal') === 'true'
                      || /modal|dialog|drawer|panel|portal/i.test(cls)
                      || style.position === 'fixed'
                      || Number(style.zIndex || 0) >= 100;
                    const notWholePage = rect.width < window.innerWidth * 0.92 || rect.left > window.innerWidth * 0.25;
                    return modalLike && notWholePage;
                  })
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            arg=[TEXT_ADD_SAME_MATERIAL],
            timeout=10000,
        )

    def _has_same_material_modal(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([modalTitle]) => {
                  return Boolean(findModal(modalTitle));

                  function findModal(title) {
                    const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                      .filter(visible)
                      .filter(el => {
                        const text = String(el.innerText || el.textContent || '');
                        if (!text.includes(title)) return false;
                        if (![...el.querySelectorAll('textarea,input')].some(visible)) return false;
                        const rect = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        const cls = String(el.className || '');
                        const role = el.getAttribute?.('role') || '';
                        const modalLike = role === 'dialog'
                          || el.getAttribute('aria-modal') === 'true'
                          || /modal|dialog|drawer|panel|portal/i.test(cls)
                          || style.position === 'fixed'
                          || Number(style.zIndex || 0) >= 100;
                        const notWholePage = rect.width < window.innerWidth * 0.92 || rect.left > window.innerWidth * 0.25;
                        return modalLike && notWholePage;
                      })
                      .sort((a, b) => area(a) - area(b));
                    return candidates[0] || null;
                  }
                  function visible(el) {
                  const rect = el.getBoundingClientRect();
                  const style = getComputedStyle(el);
                  return rect.width > 0 && rect.height > 0
                    && style.display !== 'none'
                    && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_ADD_SAME_MATERIAL],
            )
        )

    def _fill_same_material_skc(self, page, skc_text: str) -> None:
        filled = page.evaluate(
            """
            ([value, modalTitle]) => {
              const scope = findModal(modalTitle);
              if (!scope) return false;
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              const inputs = [...scope.querySelectorAll('textarea,input')]
                .filter(visible)
                .filter(el => {
                  const placeholder = String(el.getAttribute('placeholder') || '');
                  const aria = String(el.getAttribute('aria-label') || '');
                  const near = String(el.closest('label,div,td,tr')?.innerText || '');
                  return placeholder.includes('SKC')
                    || aria.includes('SKC')
                    || near.includes(modalTitle)
                    || near.includes('套版SKC');
                })
                .sort((a, b) => score(a) - score(b));
              const input = inputs[0] || [...scope.querySelectorAll('textarea,input')].filter(visible)[0];
              if (!input) return false;
              input.scrollIntoView({ block: 'center', inline: 'center' });
              input.focus();
              setNativeValue(input, value);
              input.dispatchEvent(new Event('input', { bubbles: true }));
              input.dispatchEvent(new Event('change', { bubbles: true }));
              return true;

              function score(el) {
                const near = String(el.closest('label,div,td,tr')?.innerText || '');
                let value = 100;
                if (near.includes(modalTitle)) value -= 60;
                if (near.includes('同面料套版')) value -= 40;
                if (near.includes('套版SKC') || near.includes('SKC ID')) value -= 20;
                return value;
              }
              function setNativeValue(el, newValue) {
                const proto = el instanceof HTMLTextAreaElement ? HTMLTextAreaElement.prototype : HTMLInputElement.prototype;
                const setter = Object.getOwnPropertyDescriptor(proto, 'value')?.set;
                if (setter) {
                  setter.call(el, newValue);
                } else {
                  el.value = newValue;
                }
              }
              function findModal(title) {
                const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                  .filter(visible)
                  .filter(el => {
                    const text = String(el.innerText || el.textContent || '');
                    if (!text.includes(title)) return false;
                    if (![...el.querySelectorAll('textarea,input')].some(visible)) return false;
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    const cls = String(el.className || '');
                    const role = el.getAttribute?.('role') || '';
                    const modalLike = role === 'dialog'
                      || el.getAttribute('aria-modal') === 'true'
                      || /modal|dialog|drawer|panel|portal/i.test(cls)
                      || style.position === 'fixed'
                      || Number(style.zIndex || 0) >= 100;
                    const notWholePage = rect.width < window.innerWidth * 0.92 || rect.left > window.innerWidth * 0.25;
                    return modalLike && notWholePage;
                  })
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [skc_text, TEXT_ADD_SAME_MATERIAL],
        )
        if not filled:
            raise RuntimeError("没有找到“添加同面料套版 SKC”输入框。")
        page.wait_for_timeout(300)

    def _click_modal_button(self, page, text: str) -> None:
        rect = page.evaluate(
            """
            ([targetText, modalTitle, submitText, searchText]) => {
              const scope = findModal(modalTitle);
              if (!scope) return null;
              const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              const candidates = [...scope.querySelectorAll('button,a,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === targetText)
                .map(el => el.closest('button,a') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .sort((a, b) => {
                  const scoreA = score(a);
                  const scoreB = score(b);
                  if (scoreA !== scoreB) return scoreA - scoreB;
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  if (targetText === submitText) return br.top - ar.top || ar.left - br.left;
                  return ar.top - br.top || ar.left - br.left;
                });
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function score(el) {
                const near = ancestorText(el);
                let value = 100;
                if (near.includes(modalTitle)) value -= 60;
                if (near.includes('同面料套版')) value -= 30;
                if (targetText === searchText && isBesideSkcInput(el)) value -= 100;
                return value;
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true';
              }
              function isBesideSkcInput(el) {
                const input = findSkcInput(scope);
                if (!input) return false;
                const inputRect = input.getBoundingClientRect();
                const rect = el.getBoundingClientRect();
                const verticalOverlap = rect.top < inputRect.bottom && rect.bottom > inputRect.top;
                return verticalOverlap && rect.left >= inputRect.left;
              }
              function findSkcInput(scope) {
                const candidates = [...scope.querySelectorAll('textarea,input')]
                  .filter(visible)
                  .filter(el => {
                    const placeholder = String(el.getAttribute('placeholder') || '');
                    const aria = String(el.getAttribute('aria-label') || '');
                    const near = String(el.closest('label,div,td,tr')?.innerText || '');
                    return placeholder.includes('SKC') || aria.includes('SKC') || near.includes('套版SKC') || near.includes('SKC');
                  })
                  .sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
                return candidates[0] || null;
              }
              function ancestorText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                  const role = node.getAttribute?.('role') || '';
                  if (role === 'dialog') break;
                }
                return norm(chunks.join(' '));
              }
              function findModal(title) {
                const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                  .filter(visible)
                  .filter(el => {
                    const text = String(el.innerText || el.textContent || '');
                    if (!text.includes(title)) return false;
                    if (![...el.querySelectorAll('textarea,input')].some(visible)) return false;
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    const cls = String(el.className || '');
                    const role = el.getAttribute?.('role') || '';
                    const modalLike = role === 'dialog'
                      || el.getAttribute('aria-modal') === 'true'
                      || /modal|dialog|drawer|panel|portal/i.test(cls)
                      || style.position === 'fixed'
                      || Number(style.zIndex || 0) >= 100;
                    const notWholePage = rect.width < window.innerWidth * 0.92 || rect.left > window.innerWidth * 0.25;
                    return modalLike && notWholePage;
                  })
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [text, TEXT_ADD_SAME_MATERIAL, TEXT_SUBMIT, TEXT_SEARCH],
        )
        if not rect:
            raise RuntimeError(f"没有找到“{text}”按钮。")
        page.mouse.click(rect["x"], rect["y"])

    def _wait_for_modal_search_results(self, page, timeout: int = 18000) -> int:
        try:
            page.wait_for_function(
                """
                ([modalTitle]) => {
                  const text = lowerResultText(modalTitle);
                  if (!text) return false;
                  if (text.includes('请查询 SKC ID 并添加')) return false;
                  if (text.includes('无数据') || text.includes('暂无数据')) return false;
                  return /SKC ID[:：]\\s*\\d+/.test(text);
                }

                function lowerResultText(modalTitle) {
                  const modal = findModal(modalTitle);
                  if (!modal) return '';
                  const anchor = findSkcInput(modal);
                  if (!anchor) return '';
                  const rect = anchor.getBoundingClientRect();
                  return [...modal.querySelectorAll('tr,[role="row"],td,div,span')]
                    .filter(visible)
                    .filter(el => el.getBoundingClientRect().top > rect.bottom + 8)
                    .map(el => String(el.innerText || el.textContent || ''))
                    .join(' ')
                    .replace(/\\s+/g, ' ')
                    .trim();
                }
                function findSkcInput(scope) {
                  const candidates = [...scope.querySelectorAll('textarea,input')]
                    .filter(visible)
                    .filter(el => {
                      const placeholder = String(el.getAttribute('placeholder') || '');
                      const aria = String(el.getAttribute('aria-label') || '');
                      const near = String(el.closest('label,div,td,tr')?.innerText || '');
                      return placeholder.includes('SKC') || aria.includes('SKC') || near.includes('套版SKC') || near.includes('SKC');
                    })
                    .sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
                  return candidates[0] || null;
                }
                function findModal(title) {
                  const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                    .filter(visible)
                    .filter(el => {
                      const text = String(el.innerText || el.textContent || '');
                      if (!text.includes(title)) return false;
                      if (![...el.querySelectorAll('textarea,input')].some(visible)) return false;
                      const rect = el.getBoundingClientRect();
                      const style = getComputedStyle(el);
                      const cls = String(el.className || '');
                      const role = el.getAttribute?.('role') || '';
                      const modalLike = role === 'dialog'
                        || el.getAttribute('aria-modal') === 'true'
                        || /modal|dialog|drawer|panel|portal/i.test(cls)
                        || style.position === 'fixed'
                        || Number(style.zIndex || 0) >= 100;
                      const notWholePage = rect.width < window.innerWidth * 0.92 || rect.left > window.innerWidth * 0.25;
                      return modalLike && notWholePage;
                    })
                    .sort((a, b) => area(a) - area(b));
                  return candidates[0] || null;
                }
                function visible(el) {
                  const rect = el.getBoundingClientRect();
                  const style = getComputedStyle(el);
                  return rect.width > 0 && rect.height > 0
                    && style.display !== 'none'
                    && style.visibility !== 'hidden';
                }
                function area(el) {
                  const rect = el.getBoundingClientRect();
                  return rect.width * rect.height;
                }
                """,
                arg=[TEXT_ADD_SAME_MATERIAL],
                timeout=timeout,
            )
        except Exception as exc:  # noqa: BLE001
            signals = self._read_modal_search_signals(page)
            raise RuntimeError(f"点击“搜索”后没有检测到下方结果区刷新，已停止提交。弹窗信号：{signals}") from exc

        count = self._count_modal_result_skc(page)
        if count <= 0:
            signals = self._read_modal_search_signals(page)
            raise RuntimeError(f"搜索后下方结果区没有可提交的 SKC 结果，已停止提交。弹窗信号：{signals}")
        return count

    def _count_modal_result_skc(self, page) -> int:
        return int(
            page.evaluate(
                """
                ([modalTitle]) => {
                  const modal = findModal(modalTitle);
                  if (!modal) return 0;
                  const anchor = findSkcInput(modal);
                  if (!anchor) return 0;
                  const rect = anchor.getBoundingClientRect();
                  const text = [...modal.querySelectorAll('tr,[role="row"],td,div,span')]
                    .filter(visible)
                    .filter(el => el.getBoundingClientRect().top > rect.bottom + 8)
                    .map(el => String(el.innerText || el.textContent || ''))
                    .join(' ');
                  const ids = [...text.matchAll(/SKC ID[:：]\\s*(\\d+)/g)].map(match => match[1]);
                  return new Set(ids).size;
                  function findSkcInput(scope) {
                    const candidates = [...scope.querySelectorAll('textarea,input')]
                      .filter(visible)
                      .filter(el => {
                        const placeholder = String(el.getAttribute('placeholder') || '');
                        const aria = String(el.getAttribute('aria-label') || '');
                        const near = String(el.closest('label,div,td,tr')?.innerText || '');
                        return placeholder.includes('SKC') || aria.includes('SKC') || near.includes('套版SKC') || near.includes('SKC');
                      })
                      .sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
                    return candidates[0] || null;
                  }
                  function findModal(title) {
                    const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                      .filter(visible)
                      .filter(el => {
                        const text = String(el.innerText || el.textContent || '');
                        if (!text.includes(title)) return false;
                        if (![...el.querySelectorAll('textarea,input')].some(visible)) return false;
                        const rect = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        const cls = String(el.className || '');
                        const role = el.getAttribute?.('role') || '';
                        const modalLike = role === 'dialog'
                          || el.getAttribute('aria-modal') === 'true'
                          || /modal|dialog|drawer|panel|portal/i.test(cls)
                          || style.position === 'fixed'
                          || Number(style.zIndex || 0) >= 100;
                        const notWholePage = rect.width < window.innerWidth * 0.92 || rect.left > window.innerWidth * 0.25;
                        return modalLike && notWholePage;
                      })
                      .sort((a, b) => area(a) - area(b));
                    return candidates[0] || null;
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_ADD_SAME_MATERIAL],
            )
            or 0
        )

    def _body_text(self, page) -> str:
        return str(
            page.evaluate(
                """
                () => String(document.body.innerText || '').replace(/\\s+/g, ' ').trim()
                """
            )
            or ""
        )

    def _read_page_signals(self, page) -> dict:
        try:
            return dict(
                page.evaluate(
                    """
                    ([goodsListText, goodsInfoText, moreText]) => {
                      const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
                      const text = norm(document.body.innerText || '');
                      const buttons = [...document.querySelectorAll('button,[role="button"]')]
                        .map(el => norm(el.innerText || el.textContent));
                      return {
                        url: location.href,
                        urlLooksRight: location.href.includes('/goods/list'),
                        bodyLength: text.length,
                        hasGoodsListText: text.includes(goodsListText),
                        hasGoodsInfoText: text.includes(goodsInfoText),
                        hasSkcText: text.includes('SKC ID'),
                        hasSkuText: text.includes('SKU ID'),
                        tableCount: document.querySelectorAll('table').length,
                        hasMoreButton: buttons.includes(moreText),
                        buttons: buttons.slice(0, 12),
                      };
                    }
                    """,
                    [TEXT_GOODS_LIST, TEXT_GOODS_INFO, TEXT_MORE],
                )
            )
        except Exception as exc:  # noqa: BLE001
            return {"url": page.url, "error": str(exc)}

    def _read_template_group_signals(self, page) -> dict:
        try:
            return dict(
                page.evaluate(
                    """
                    ([titleText, addText, sameMaterialText]) => {
                      const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
                      const visible = el => {
                        const rect = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        return rect.width > 0 && rect.height > 0
                          && style.display !== 'none'
                          && style.visibility !== 'hidden';
                      };
                      const text = norm(document.body.innerText || '');
                      const actionTexts = [...document.querySelectorAll('button,a,[role="button"],span')]
                        .filter(visible)
                        .map(el => norm(el.innerText || el.textContent))
                        .filter(Boolean);
                      return {
                        url: location.href,
                        urlLooksRight: location.href.includes('/sample/clothing-set'),
                        bodyLength: text.length,
                        hasTemplateGroupText: text.includes(titleText),
                        hasAddSameMaterialText: text.includes(sameMaterialText) || text.includes('同面料套版'),
                        hasBaseSkcText: text.includes('底版SKC'),
                        hasSkcText: text.includes('SKC ID') || text.includes('SKC'),
                        tableCount: document.querySelectorAll('table').length,
                        inputCount: document.querySelectorAll('textarea,input').length,
                        visibleAddCount: actionTexts.filter(item => item === addText).length,
                        actions: actionTexts.slice(0, 16),
                      };
                    }
                    """,
                    [TEXT_TEMPLATE_GROUP, TEXT_ADD, TEXT_ADD_SAME_MATERIAL],
                )
            )
        except Exception as exc:  # noqa: BLE001
            return {"url": page.url, "error": str(exc)}

    def _read_modal_search_signals(self, page) -> dict:
        try:
            return dict(
                page.evaluate(
                    """
                    ([modalTitle]) => {
                      const modal = findModal(modalTitle);
                      if (!modal) return { hasModal: false };
                      const input = findSkcInput(modal);
                      const rect = input?.getBoundingClientRect?.();
                      const lowerText = input
                        ? [...modal.querySelectorAll('tr,[role="row"],td,div,span')]
                            .filter(visible)
                            .filter(el => el.getBoundingClientRect().top > rect.bottom + 8)
                            .map(el => String(el.innerText || el.textContent || ''))
                            .join(' ')
                            .replace(/\\s+/g, ' ')
                            .trim()
                        : '';
                      const actions = [...modal.querySelectorAll('button,a,[role="button"],span')]
                        .filter(visible)
                        .map(el => String(el.innerText || el.textContent || '').replace(/\\s+/g, ' ').trim())
                        .filter(Boolean)
                        .slice(0, 20);
                      const ids = [...lowerText.matchAll(/SKC ID[:：]\\s*(\\d+)/g)].map(match => match[1]);
                      return {
                        hasModal: true,
                        hasInput: Boolean(input),
                        inputValueLength: input ? String(input.value || '').length : 0,
                        lowerTextLength: lowerText.length,
                        lowerHasQueryHint: lowerText.includes('请查询 SKC ID 并添加'),
                        lowerHasNoData: lowerText.includes('无数据') || lowerText.includes('暂无数据'),
                        lowerSkcCount: new Set(ids).size,
                        actions,
                      };

                      function findSkcInput(scope) {
                        const candidates = [...scope.querySelectorAll('textarea,input')]
                          .filter(visible)
                          .filter(el => {
                            const placeholder = String(el.getAttribute('placeholder') || '');
                            const aria = String(el.getAttribute('aria-label') || '');
                            const near = String(el.closest('label,div,td,tr')?.innerText || '');
                            return placeholder.includes('SKC') || aria.includes('SKC') || near.includes('套版SKC') || near.includes('SKC');
                          })
                          .sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
                        return candidates[0] || null;
                      }
                      function findModal(title) {
                        const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                          .filter(visible)
                          .filter(el => {
                            const text = String(el.innerText || el.textContent || '');
                            if (!text.includes(title)) return false;
                            if (![...el.querySelectorAll('textarea,input')].some(visible)) return false;
                            const rect = el.getBoundingClientRect();
                            const style = getComputedStyle(el);
                            const cls = String(el.className || '');
                            const role = el.getAttribute?.('role') || '';
                            const modalLike = role === 'dialog'
                              || el.getAttribute('aria-modal') === 'true'
                              || /modal|dialog|drawer|panel|portal/i.test(cls)
                              || style.position === 'fixed'
                              || Number(style.zIndex || 0) >= 100;
                            const notWholePage = rect.width < window.innerWidth * 0.92 || rect.left > window.innerWidth * 0.25;
                            return modalLike && notWholePage;
                          })
                          .sort((a, b) => area(a) - area(b));
                        return candidates[0] || null;
                      }
                      function visible(el) {
                        const rect = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        return rect.width > 0 && rect.height > 0
                          && style.display !== 'none'
                          && style.visibility !== 'hidden';
                      }
                      function area(el) {
                        const rect = el.getBoundingClientRect();
                        return rect.width * rect.height;
                      }
                    }
                    """,
                    [TEXT_ADD_SAME_MATERIAL],
                )
            )
        except Exception as exc:  # noqa: BLE001
            return {"url": page.url, "error": str(exc)}

    def _ensure_compliance_info_page(self, page) -> None:
        if "agentseller.temu.com/govern/information-supplementation" not in page.url:
            raise RuntimeError(f"当前不在商品合规信息页面：{page.url}")
        try:
            page.wait_for_function(
                """
                ([titleText, uploadText]) => {
                  const text = document.body.innerText || '';
                  return text.includes(titleText) || text.includes(uploadText);
                }
                """,
                arg=[TEXT_COMPLIANCE_INFO, TEXT_BATCH_UPLOAD_COMPLIANCE],
                timeout=10000,
            )
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("未检测到商品合规信息页面内容。") from exc

    def _ensure_compliant_live_photo_files(self) -> None:
        missing = [
            str(path)
            for path in (COMPLIANT_FRONT_IMAGE_PATH, COMPLIANT_SIDE_IMAGE_PATH)
            if not path.exists()
        ]
        if missing:
            raise RuntimeError("没有找到商品合规图文件：" + "；".join(missing))

    def _ensure_compliant_live_photos_page(self, page) -> None:
        if "agentseller.temu.com/govern/compliant-live-photos" not in page.url:
            raise RuntimeError(f"当前不在商品实拍图页面：{page.url}")
        try:
            page.wait_for_function(
                """
                ([pageText, batchText, uploadText]) => {
                  const text = String(document.body.innerText || '');
                  const urlLooksRight = location.href.includes('/govern/compliant-live-photos');
                  return urlLooksRight
                    && text.length > 0
                    && (text.includes(pageText) || text.includes(batchText) || text.includes(uploadText));
                }
                """,
                arg=[TEXT_COMPLIANT_LIVE_PHOTOS, TEXT_BATCH_UPLOAD_PACKAGE_PHOTOS, TEXT_UPLOAD_LIVE_PHOTO],
                timeout=10000,
            )
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("未检测到商品实拍图页面内容。") from exc

    def _open_compliant_live_photos_batch_page(self, page) -> None:
        if "compliant-live-photos-batch" in page.url:
            self.log("当前已在商品合规图批量上传页面。")
            return
        self.log("正在打开商品合规图批量上传页面。")
        for attempt in range(1, 6):
            if not self._click_button_by_text(page, TEXT_BATCH_UPLOAD_PACKAGE_PHOTOS):
                raise RuntimeError("没有找到“具有相同包装的商品批量上传或更新”按钮。")
            page.wait_for_timeout(600)
            if not self._wait_for_popup_item(page, TEXT_UPLOAD_LIVE_PHOTO, exact=True, timeout=4000):
                self.log(f"第 {attempt} 次没有等到“上传实拍图”菜单项，准备重试。")
                page.keyboard.press("Escape")
                page.wait_for_timeout(800)
                continue
            if not self._activate_popup_item(page, TEXT_UPLOAD_LIVE_PHOTO, exact=True):
                self.log(f"第 {attempt} 次检测到“上传实拍图”但点击失败，准备重试。")
                page.keyboard.press("Escape")
                page.wait_for_timeout(800)
                continue
            try:
                page.wait_for_url("**/govern/compliant-live-photos-batch**", timeout=12000)
                self._wait_for_page_ready(page, timeout=10000)
                return
            except Exception:  # noqa: BLE001 - retry trigger when navigation was slow or did not start.
                self.log(f"第 {attempt} 次点击“上传实拍图”后没有进入批量上传页，准备重试。")
                page.wait_for_timeout(900)
        raise RuntimeError("多次尝试后仍没有进入商品合规图批量上传页面。")

    def _ensure_compliant_live_photos_batch_page(self, page) -> None:
        if "agentseller.temu.com/govern/compliant-live-photos-batch" not in page.url:
            raise RuntimeError(f"当前不在商品合规图批量上传页面：{page.url}")
        try:
            page.wait_for_function(
                """
                ([categoryText, historyText, submitText]) => {
                  const text = String(document.body.innerText || '');
                  return location.href.includes('/govern/compliant-live-photos-batch')
                    && text.includes(categoryText)
                    && (text.includes(historyText) || text.includes(submitText));
                }
                """,
                arg=[TEXT_PRODUCT_CATEGORY, TEXT_HISTORY, TEXT_CONFIRM_SUBMIT],
                timeout=12000,
            )
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("未检测到商品合规图批量上传页面内容。") from exc

    def _select_first_live_photo_history(self, page) -> None:
        self.log("正在选择第一条历史记录。")
        for attempt in range(1, 6):
            if self._click_first_live_photo_history_tag(page):
                self._wait_for_page_ready(page, timeout=12000)
                page.wait_for_timeout(1500)
                if self._live_photo_has_product_table(page):
                    return
            self.log(f"第 {attempt} 次选择历史记录后还未检测到商品列表，继续等待/重试。")
            page.wait_for_timeout(1200)
        raise RuntimeError("选择历史记录后没有检测到商品合规图商品列表。")

    def _click_first_live_photo_history_tag(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([historyText]) => {
                  const label = [...document.querySelectorAll('span,div,label')]
                    .filter(visible)
                    .find(el => norm(el.innerText || el.textContent) === historyText);
                  if (!label) return false;
                  const row = label.closest('.rocket-row-flex,.rocket-form-field-item') || label.parentElement;
                  const tags = [...(row || document).querySelectorAll('.rocket-tag,button,a,span')]
                    .filter(visible)
                    .filter(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text && text !== historyText && !text.includes('请选择');
                    })
                    .sort((a, b) => {
                      const ar = a.getBoundingClientRect();
                      const br = b.getBoundingClientRect();
                      const aTag = String(a.className || '').includes('rocket-tag') ? 0 : 1;
                      const bTag = String(b.className || '').includes('rocket-tag') ? 0 : 1;
                      return aTag - bTag || area(a) - area(b) || ar.left - br.left;
                    });
                  const target = tags[0];
                  if (!target) return false;
                  fireMouse(target);
                  target.click?.();
                  return true;

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                  function fireMouse(el) {
                    el.scrollIntoView({ block: 'center', inline: 'center' });
                    const rect = el.getBoundingClientRect();
                    const clientX = rect.left + rect.width / 2;
                    const clientY = rect.top + rect.height / 2;
                    for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                      el.dispatchEvent(new MouseEvent(type, {
                        bubbles: true,
                        cancelable: true,
                        view: window,
                        clientX,
                        clientY,
                      }));
                    }
                  }
                }
                """,
                [TEXT_HISTORY],
            )
        )

    def _live_photo_has_product_table(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                () => {
                  const text = String(document.body.innerText || '');
                  return text.includes('商品信息')
                    || text.includes('实拍图识别类型')
                    || text.includes('识别状态')
                    || document.querySelectorAll('table,.rocket-table').length > 0;
                }
                """
            )
        )

    def _set_live_photo_page_size(self, page, size: int) -> None:
        self.log(f"正在设置商品合规图每页 {size} 条。")
        rect = None
        for attempt in range(1, 9):
            current_text = self._read_live_photo_page_size_text(page)
            if str(size) in current_text:
                self.log(f"商品合规图当前已经是每页 {size} 条。")
                return
            rect = self._locate_live_photo_page_size_control(page)
            if rect:
                break
            self.log(f"第 {attempt} 次没有找到商品合规图分页条数控件，继续等待。")
            page.wait_for_timeout(800)
        if not rect:
            raise RuntimeError("没有找到商品合规图分页条数控件。")
        for attempt in range(1, 5):
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(600)
            if self._click_live_photo_page_size_option(page, size):
                break
            self.log(f"第 {attempt} 次没有点中 {size}条/页 选项，准备重新展开分页下拉。")
            page.keyboard.press("Escape")
            page.wait_for_timeout(500)
        else:
            raise RuntimeError(f"没有找到 {size}条/页 选项。")
        self._wait_for_page_ready(page, timeout=8000)
        page.wait_for_timeout(1200)
        if str(size) in self._read_live_photo_page_size_text(page):
            return
        self._fill_live_photo_page_size_control(page, size)
        self._wait_for_page_ready(page, timeout=8000)
        page.wait_for_timeout(1200)
        if str(size) not in self._read_live_photo_page_size_text(page):
            raise RuntimeError(f"分页条数切换后未检测到 {size}条/页。")

    def _click_live_photo_page_size_option(self, page, size: int) -> bool:
        targets = page_size_option_labels(size)
        for _attempt in range(5):
            clicked = page.evaluate(
                """
                ([targetTexts]) => {
                  const targets = targetTexts.map(compact);
                  const candidates = [...document.querySelectorAll('body *')]
                    .filter(visible)
                    .filter(el => targets.includes(compact(el.innerText || el.textContent || el.value)))
                    .sort((a, b) => {
                      const ar = a.getBoundingClientRect();
                      const br = b.getBoundingClientRect();
                      const aOption = isOption(a) ? 0 : 1;
                      const bOption = isOption(b) ? 0 : 1;
                      return aOption - bOption || area(a) - area(b) || br.top - ar.top;
                    });
                  const target = candidates[0];
                  if (!target) return false;
                  fireMouse(target);
                  target.click?.();
                  return true;

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isOption(el) {
                    const cls = String(el.className || '');
                    const role = el.getAttribute?.('role') || '';
                    return /option|item|select|dropdown/i.test(cls) || /option|menuitem/.test(role);
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                  function fireMouse(el) {
                    el.scrollIntoView({ block: 'center', inline: 'center' });
                    const rect = el.getBoundingClientRect();
                    const clientX = rect.left + rect.width / 2;
                    const clientY = rect.top + rect.height / 2;
                    for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                      el.dispatchEvent(new MouseEvent(type, {
                        bubbles: true,
                        cancelable: true,
                        view: window,
                        clientX,
                        clientY,
                      }));
                    }
                  }
                }
                """,
                [targets],
            )
            if clicked:
                page.wait_for_timeout(700)
                return True
            page.wait_for_timeout(500)
        return False

    def _locate_live_photo_page_size_control(self, page):
        return page.evaluate(
            """
            () => {
              const candidates = [...document.querySelectorAll('.rocket-pagination-options-size-changer,.rocket-pagination-options .rocket-select,input,button,span,div')]
                .filter(visible)
                .filter(el => /\\d+\\s*条\\/页/.test(norm(el.innerText || el.textContent || el.value)))
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  const aSelect = /rocket-select/.test(String(a.className || '')) ? 0 : 1;
                  const bSelect = /rocket-select/.test(String(b.className || '')) ? 0 : 1;
                  return aSelect - bSelect || br.top - ar.top || br.left - ar.left || area(a) - area(b);
                });
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """
        )

    def _fill_live_photo_page_size_control(self, page, size: int) -> None:
        rect = self._locate_live_photo_page_size_control(page)
        if not rect:
            return
        page.mouse.click(rect["x"], rect["y"])
        page.keyboard.press("Control+A")
        page.keyboard.type(str(size), delay=20)
        page.keyboard.press("Enter")

    def _read_live_photo_page_size_text(self, page) -> str:
        return str(
            page.evaluate(
                """
                () => {
                  const target = [...document.querySelectorAll('.rocket-pagination-options-size-changer,.rocket-pagination-options')]
                    .filter(visible)
                    .map(el => norm((el.innerText || el.textContent || '') + ' ' + (el.value || '')))
                    .find(text => /\\d+\\s*条\\/页/.test(text));
                  return target || '';

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """
            )
            or ""
        )

    def _select_all_live_photo_products(self, page) -> int:
        self.log("正在全选商品合规图商品。")
        for attempt in range(1, 5):
            rect = page.evaluate(
                """
                () => {
                  const table = [...document.querySelectorAll('.rocket-table,.rocket-table-container,table')]
                    .filter(visible)
                    .filter(el => area(el) > 1000)
                    .sort((a, b) => {
                      const at = norm(a.innerText || a.textContent);
                      const bt = norm(b.innerText || b.textContent);
                      const as = scoreTable(a, at);
                      const bs = scoreTable(b, bt);
                      return bs - as || a.getBoundingClientRect().top - b.getBoundingClientRect().top;
                    })[0];
                  const scope = table || document;
                  const header = [...scope.querySelectorAll('.rocket-table-thead,thead')][0] || scope;
                  const target = [...header.querySelectorAll('input[type="checkbox"],[role="checkbox"],.rocket-checkbox,label')]
                    .filter(visible)
                    .sort((a, b) => a.getBoundingClientRect().left - b.getBoundingClientRect().left)[0];
                  if (!target) return null;
                  target.scrollIntoView({ block: 'center', inline: 'center' });
                  const rect = target.getBoundingClientRect();
                  return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

                  function scoreTable(el, text) {
                    let score = 0;
                    if (text.includes('商品信息')) score += 100;
                    if (text.includes('实拍图识别类型')) score += 80;
                    if (text.includes('识别状态')) score += 80;
                    if (text.includes('售卖影响及建议')) score += 40;
                    return score;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """
            )
            if not rect:
                page.wait_for_timeout(800)
                continue
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(1200)
            selected_count = self._read_live_photo_selected_count(page)
            if selected_count > 0:
                self.log(f"已全选商品合规图商品：{selected_count} 个。")
                return selected_count
            self.log(f"第 {attempt} 次点击全选后未检测到已选数量，准备重试。")
        raise RuntimeError("已点击商品合规图全选，但没有检测到已选商品数量。")

    def _read_live_photo_selected_count(self, page) -> int:
        return int(
            page.evaluate(
                """
                () => {
                  const text = String(document.body.innerText || '');
                  const patterns = [
                    /已选[:：]?\\s*(\\d+)/,
                    /选中[:：]?\\s*(\\d+)/,
                    /共选择\\s*(\\d+)/
                  ];
                  for (const pattern of patterns) {
                    const match = text.match(pattern);
                    if (match) return Number(match[1]);
                  }
                  const checked = [...document.querySelectorAll('input[type="checkbox"]')]
                    .filter(input => input.checked)
                    .filter(input => !ancestorText(input).includes('是否要更新'));
                  return checked.length;

                  function ancestorText(el) {
                    const chunks = [];
                    let node = el;
                    for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                      chunks.push(node.innerText || node.textContent || '');
                    }
                    return chunks.join(' ');
                  }
                }
                """
            )
            or 0
        )

    def _upload_live_photo_image(self, page, label_text: str, file_path: Path) -> None:
        if not file_path.exists():
            raise RuntimeError(f"商品合规图文件不存在：{file_path}")
        if self._live_photo_upload_has_file(page, label_text):
            self.log(f"{label_text}已存在上传文件，跳过重复上传。")
            return
        self.log(f"正在上传{label_text}：{file_path}")
        for attempt in range(1, 4):
            if self._set_live_photo_file_input_near_label(page, label_text, file_path):
                if self._wait_for_live_photo_upload_file(page, label_text, timeout=6000):
                    return
            self.log(f"第 {attempt} 次没有在“{label_text}”对应上传框中检测到文件，准备重试。")
            page.wait_for_timeout(800)
        raise RuntimeError(f"上传{label_text}后没有检测到文件已选中。")

    def _wait_for_live_photo_upload_file(self, page, label_text: str, timeout: int = 6000) -> bool:
        end_time = page.evaluate("Date.now()") + timeout
        while page.evaluate("Date.now()") < end_time:
            if self._live_photo_upload_has_file(page, label_text):
                return True
            page.wait_for_timeout(500)
        return False

    def _set_live_photo_file_input_near_label(self, page, label_text: str, file_path: Path) -> bool:
        handle = page.evaluate_handle(
            """
            ([labelText]) => {
              const card = findUploadCard(labelText);
              if (!card) return null;
              const emptyInputs = [...card.querySelectorAll('input[type="file"]')]
                .filter(input => !input.files || input.files.length === 0);
              return emptyInputs[0] || null;

              function findUploadCard(text) {
                const labels = [...document.querySelectorAll('span,div,label')]
                  .filter(visible)
                  .filter(el => compact(el.innerText || el.textContent).includes(compact(text)))
                  .sort((a, b) => area(a) - area(b));
                for (const label of labels) {
                  const card = closestCard(label, text);
                  if (card) return card;
                }
                return null;
              }

              function closestCard(label, text) {
                let node = label;
                for (let depth = 0; node && depth < 7; depth += 1, node = node.parentElement) {
                  const text = norm(node.innerText || node.textContent);
                  const compactText = compact(text);
                  const rect = node.getBoundingClientRect();
                  if (
                    compactText.includes(compact(labelText))
                    && node.querySelector('input[type="file"]')
                    && area(node) >= 500
                    && rect.width <= 240
                    && rect.height <= 220
                  ) {
                    return node;
                  }
                }
                return null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [label_text],
        )
        element = handle.as_element()
        if not element:
            return False
        element.set_input_files(str(file_path), timeout=10000)
        return True

    def _live_photo_upload_trigger_rect(self, page, label_text: str):
        return page.evaluate(
            """
            ([labelText]) => {
              const label = [...document.querySelectorAll('span,div,label')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent).includes(labelText))
                .sort((a, b) => area(a) - area(b))[0];
              if (!label) return null;
              const labelRect = label.getBoundingClientRect();
              const candidates = [...document.querySelectorAll('button,[role="button"],.rocket-upload,span,div')]
                .filter(visible)
                .filter(el => {
                  const text = norm(el.innerText || el.textContent);
                  if (!(text.includes(labelText) || text.includes('上传'))) return false;
                  const rect = el.getBoundingClientRect();
                  return Math.abs((rect.left + rect.width / 2) - (labelRect.left + labelRect.width / 2)) < 220
                    && Math.abs((rect.top + rect.height / 2) - (labelRect.top + labelRect.height / 2)) < 160;
                })
                .sort((a, b) => area(a) - area(b));
              const target = candidates[0] || label;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [label_text],
        )

    def _live_photo_upload_has_file(self, page, label_text: str) -> bool:
        return bool(
            page.evaluate(
                """
                ([labelText]) => {
                  const label = [...document.querySelectorAll('span,div,label')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent).includes(labelText))
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!label) return false;
                  let node = label;
                  for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                    const text = norm(node.innerText || node.textContent);
                    if (text.includes('商品外包装实拍图') && text.includes(labelText) && !text.includes(`${labelText} (0/`)) {
                      return true;
                    }
                    if (text.includes(labelText) && (
                      /\\.(jpg|jpeg|png|webp)/i.test(text)
                      || /\\(\\s*[1-9]\\d*\\s*\\/\\s*\\d+\\s*\\)/.test(text)
                      || text.includes('重新上传')
                      || text.includes('上传成功')
                    )) {
                      return true;
                    }
                  }
                  return false;

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [label_text],
            )
        )

    def _select_live_photo_update_all_option(self, page) -> None:
        self.log("正在选择更新所选全部商品的包装实拍图。")
        for attempt in range(1, 5):
            rect = page.evaluate(
                """
                ([targetText]) => {
                  window.scrollTo({ top: document.body.scrollHeight, behavior: 'instant' });
                  const target = [...document.querySelectorAll('label,span,div')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent).includes(targetText))
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!target) return null;
                  const clickable = target.closest('label') || target;
                  clickable.scrollIntoView({ block: 'center', inline: 'center' });
                  const rect = clickable.getBoundingClientRect();
                  return { x: rect.left + Math.min(18, rect.width / 2), y: rect.top + rect.height / 2 };

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_UPDATE_SELECTED_PACKAGE_PHOTOS],
            )
            if rect:
                page.mouse.click(rect["x"], rect["y"])
                page.wait_for_timeout(700)
                if self._live_photo_update_all_selected(page):
                    return
            self.log(f"第 {attempt} 次选择更新选项未确认，准备重试。")
            page.wait_for_timeout(800)
        raise RuntimeError("没有成功选中“更新所选全部商品的包装实拍图”。")

    def _live_photo_update_all_selected(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([targetText]) => {
                  const target = [...document.querySelectorAll('label,span,div')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent).includes(targetText))
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!target) return false;
                  let node = target;
                  for (let depth = 0; node && depth < 6; depth += 1, node = node.parentElement) {
                    const input = node.querySelector?.('input[type="radio"]');
                    if (input && input.checked) return true;
                    const cls = String(node.className || '');
                    if (/radio.*checked|checked.*radio/i.test(cls)) return true;
                    if (node.getAttribute?.('aria-checked') === 'true') return true;
                  }
                  return false;

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_UPDATE_SELECTED_PACKAGE_PHOTOS],
            )
        )

    def _submit_live_photo_upload(self, page) -> None:
        self._click_global_button(page, TEXT_CONFIRM_SUBMIT, timeout=10000)
        self._wait_for_page_ready(page, timeout=12000)
        page.wait_for_timeout(1200)
        if self._click_global_button_if_visible(page, TEXT_CONFIRM, timeout=5000):
            self.log("已点击商品合规图提交确认。")
        if self._click_global_button_if_visible(page, TEXT_I_KNOW, timeout=12000):
            self.log("已点击“我知道了”。")

    def _ensure_jit_product_select_page(self, page) -> None:
        if "agentseller.temu.com/newon/product-select" not in page.url:
            raise RuntimeError(f"当前不在JIT商品选择页面：{page.url}")
        try:
            page.wait_for_function(
                """
                ([titleText, adjustText, goodsInfoText]) => {
                  const text = String(document.body.innerText || '');
                  return location.href.includes('/newon/product-select')
                    && text.length > 0
                    && (text.includes(titleText) || text.includes(adjustText) || text.includes(goodsInfoText));
                }
                """,
                arg=[TEXT_NEWON_PRODUCT_SELECT, TEXT_BATCH_ADJUST_JIT, TEXT_GOODS_INFO],
                timeout=12000,
            )
        except Exception as exc:  # noqa: BLE001
            signals = self._read_jit_product_select_signals(page)
            raise RuntimeError(f"未检测到JIT商品选择页面内容。页面信号：{signals}") from exc

    def _confirm_pending_jit_product_info_if_needed(self, page) -> int:
        pending_count = self._read_jit_pending_product_info_count(page)
        if pending_count <= 0:
            self.log("未检测到需要确认商品信息的商品，继续开通JIT管理。")
            return 0

        self.log(f"检测到商品信息待确认：{pending_count} 个，准备先批量确认商品信息。")
        self._activate_jit_pending_product_info_filter(page)
        self._set_jit_product_page_size(page, JIT_PRODUCT_INFO_CONFIRM_PAGE_SIZE)
        self._activate_jit_pending_product_info_filter(page)
        self._ensure_jit_pending_product_info_list_active(page)
        selected_count = self._select_all_jit_products(page)
        if not self._jit_batch_confirm_product_info_enabled(page):
            self.log("全选后“批量确认商品信息”仍不可用，准备重新点击“商品信息待确认”后再全选。")
            self._activate_jit_pending_product_info_filter(page)
            self._ensure_jit_pending_product_info_list_active(page)
            selected_count = self._select_all_jit_products(page)
        self._click_batch_confirm_product_info(page)
        self._confirm_pending_product_info_modal_if_visible(page)
        self._wait_for_page_ready(page, timeout=15000)
        page.wait_for_timeout(1500)
        if self._click_global_button_if_visible(page, TEXT_I_KNOW, timeout=6000):
            self.log("已点击“我知道了”。")
        elif self._click_visible_text_as_button(page, TEXT_I_KNOW):
            self.log("已点击“我知道了”。")
        self.log(f"已批量确认商品信息，选中商品：{selected_count} 个，准备刷新页面后继续开通JIT管理。")
        page.reload(wait_until="domcontentloaded")
        self._wait_for_page_ready(page, timeout=15000)
        self._ensure_jit_product_select_page(page)
        return selected_count

    def _read_jit_pending_product_info_count(self, page) -> int:
        return int(
            page.evaluate(
                """
                ([targetText]) => {
                  const targetCompact = compact(targetText);
                  const candidates = [...document.querySelectorAll('div,span,button,[role="button"]')]
                    .filter(visible)
                    .map(el => ({
                      text: norm(el.innerText || el.textContent),
                      compactText: compact(el.innerText || el.textContent),
                      area: area(el),
                    }))
                    .filter(item => item.compactText.includes(targetCompact))
                    .sort((a, b) => a.area - b.area);
                  for (const item of candidates) {
                    const match = item.compactText.match(new RegExp(targetCompact + '(\\\\d+)'));
                    if (match) return Number(match[1] || 0);
                  }
                  const bodyMatch = compact(document.body.innerText || '').match(new RegExp(targetCompact + '(\\\\d+)'));
                  return Number(bodyMatch?.[1] || 0);

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_PENDING_PRODUCT_INFO_CONFIRM],
            )
            or 0
        )

    def _activate_jit_pending_product_info_filter(self, page) -> None:
        if self._click_jit_quick_filter(page, TEXT_PENDING_PRODUCT_INFO_CONFIRM):
            self._wait_for_page_ready(page, timeout=12000)
            page.wait_for_timeout(1000)
            return
        raise RuntimeError("检测到存在商品信息待确认，但没有成功点击“商品信息待确认”筛选项。")

    def _ensure_jit_pending_product_info_list_active(self, page) -> None:
        if self._jit_pending_product_info_list_active(page):
            return
        raise RuntimeError("已切换每页 100 条，但当前列表不是“商品信息待确认”列表，已停止全选。")

    def _jit_pending_product_info_list_active(self, page) -> bool:
        if not self._jit_pending_product_info_filter_selected(page):
            return False
        return bool(
            page.evaluate(
                """
                () => {
                  const text = norm(document.body.innerText || '');
                  const tableText = readTableText();
                  if (tableText.includes('上新待确认') || tableText.includes('去确认')) return true;
                  if (text.includes('全部 0') && text.includes('商品信息待确认')) return true;
                  return false;

                  function readTableText() {
                    const tables = [...document.querySelectorAll('[class*="TB_innerMiddle"],[class*="TB_body"],tbody,.TB_outerWrapper_5-120-1,[class*="TB_outerWrapper"]')]
                      .filter(visible)
                      .map(el => norm(el.innerText || el.textContent))
                      .sort((a, b) => b.length - a.length);
                    return tables[0] || '';
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """
            )
        )

    def _jit_pending_product_info_filter_selected(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([targetText]) => {
                  const targetCompact = compact(targetText);
                  const candidates = [...document.querySelectorAll('div,span,button,[role="button"]')]
                    .filter(visible)
                    .filter(el => compact(el.innerText || el.textContent).includes(targetCompact))
                    .map(el => el.closest('[class*="card"],button,[role="button"]') || el)
                    .filter(visible)
                    .filter((el, index, list) => list.indexOf(el) === index)
                    .sort((a, b) => area(a) - area(b));
                  for (const candidate of candidates) {
                    if (selected(candidate)) return true;
                  }
                  return false;

                  function selected(el) {
                    let node = el;
                    for (let depth = 0; node && depth < 4; depth += 1, node = node.parentElement) {
                      const style = getComputedStyle(node);
                      const cls = String(node.className || '');
                      if (node.getAttribute?.('aria-selected') === 'true') return true;
                      if (/active|selected|current/i.test(cls)) return true;
                      if (isBlue(style.color) || isBlue(style.borderColor) || isBlue(style.outlineColor)) return true;
                      if (String(style.border || '').includes('64, 124, 255')) return true;
                      if (String(style.boxShadow || '').includes('64, 124, 255')) return true;
                    }
                    return false;
                  }
                  function isBlue(value) {
                    const nums = String(value || '').match(/\\d+(?:\\.\\d+)?/g)?.map(Number) || [];
                    return nums.length >= 3 && nums[2] > 180 && nums[0] < 120 && nums[1] > 80 && nums[1] < 170;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_PENDING_PRODUCT_INFO_CONFIRM],
            )
        )

    def _click_jit_quick_filter(self, page, text: str) -> bool:
        rect = page.evaluate(
            """
            ([targetText]) => {
              const targetCompact = compact(targetText);
              const candidates = [...document.querySelectorAll('div,span,button,[role="button"]')]
                .filter(visible)
                .filter(el => compact(el.innerText || el.textContent).includes(targetCompact))
                .map(el => el.closest('[class*="card"],button,[role="button"]') || el)
                .filter(visible)
                .filter((el, index, list) => list.indexOf(el) === index)
                .sort((a, b) => score(a) - score(b));
              const target = candidates[0] || null;
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function score(el) {
                const rect = el.getBoundingClientRect();
                let value = rect.width * rect.height;
                if (String(el.className || '').includes('card')) value -= 10000;
                return value;
              }
            }
            """,
            [text],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        return True

    def _jit_batch_confirm_product_info_enabled(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([targetText]) => {
                  const targetCompact = compact(targetText);
                  return [...document.querySelectorAll('button,[role="button"],a')]
                    .filter(visible)
                    .filter(el => compact(el.innerText || el.textContent) === targetCompact)
                    .some(el => !isDisabled(el));

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isDisabled(el) {
                    return el.disabled === true
                      || el.getAttribute('disabled') !== null
                      || el.getAttribute('aria-disabled') === 'true'
                      || /disabled/.test(String(el.className || ''));
                  }
                }
                """,
                [TEXT_BATCH_CONFIRM_PRODUCT_INFO],
            )
        )

    def _click_batch_confirm_product_info(self, page) -> None:
        self.log("正在点击“批量确认商品信息”。")
        for attempt in range(1, 5):
            if self._click_enabled_button_by_text(page, TEXT_BATCH_CONFIRM_PRODUCT_INFO):
                page.wait_for_timeout(1200)
                return
            self.log(f"第 {attempt} 次没有点到“批量确认商品信息”，准备重试。")
            page.wait_for_timeout(700)
        raise RuntimeError("没有成功点击“批量确认商品信息”。")

    def _confirm_pending_product_info_modal_if_visible(self, page) -> None:
        for _attempt in range(1, 4):
            if self._click_pending_product_info_confirm_button(page):
                self._wait_for_page_ready(page, timeout=15000)
                page.wait_for_timeout(1200)
                return
            page.wait_for_timeout(600)

    def _click_pending_product_info_confirm_button(self, page) -> bool:
        rect = page.evaluate(
            """
            ([confirmText, modalSelector]) => {
              const modal = [...document.querySelectorAll(modalSelector)]
                .filter(visible)
                .filter(el => {
                  const text = norm(el.innerText || el.textContent);
                  return text.includes(confirmText)
                    && (text.includes('商品信息') || text.includes('确认商品') || text.includes('批量确认'));
                })
                .sort((a, b) => area(a) - area(b))[0];
              if (!modal) return null;
              const scope = modal;
              const button = [...scope.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === confirmText)
                .map(el => el.closest('button,[role="button"],a') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .filter((el, index, list) => list.indexOf(el) === index)
                .sort((a, b) => area(a) - area(b))[0];
              if (!button) return null;
              button.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = button.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''));
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_CONFIRM, jit_confirm_modal_selector()],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        return True

    def _click_enabled_button_by_text(self, page, text: str) -> bool:
        rect = page.evaluate(
            """
            ([targetText]) => {
              const targetCompact = compact(targetText);
              const button = [...document.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => compact(el.innerText || el.textContent) === targetCompact)
                .map(el => el.closest('button,[role="button"],a') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .filter((el, index, list) => list.indexOf(el) === index)
                .sort((a, b) => area(a) - area(b))[0];
              if (!button) return null;
              button.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = button.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''));
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [text],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        return True

    def _set_jit_product_page_size(self, page, size: int) -> None:
        if size <= 0:
            raise RuntimeError("JIT分页数量必须大于 0。")
        self.log(f"正在设置JIT商品选择每页 {size} 条。")
        for attempt in range(1, 6):
            current = self._read_jit_product_page_size(page)
            if current == size:
                self.log(f"JIT商品选择当前已经是每页 {size} 条。")
                return
            rect = self._locate_jit_product_page_size_control(page)
            if not rect:
                self.log(f"第 {attempt} 次没有找到JIT分页条数控件，继续等待。")
                page.wait_for_timeout(900)
                continue
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(600)
            if self._click_jit_page_size_option(page, size):
                self._wait_for_page_ready(page, timeout=10000)
                page.wait_for_timeout(1200)
                if self._read_jit_product_page_size(page) == size:
                    return
            self._fill_jit_page_size_control(page, size)
            self._wait_for_page_ready(page, timeout=10000)
            page.wait_for_timeout(1200)
            if self._read_jit_product_page_size(page) == size:
                return
            self.log(f"第 {attempt} 次设置JIT分页条数后还未生效，准备重试。")
        raise RuntimeError(f"JIT分页条数切换后未检测到 {size} 条/页。")

    def _read_jit_product_page_size(self, page) -> int:
        return int(
            page.evaluate(
                """
                () => {
                  const inputs = [...document.querySelectorAll('input')]
                    .filter(visible)
                    .filter(input => /^\\d+$/.test(input.value || ''))
                    .map(input => ({ input, value: Number(input.value), score: score(input) }))
                    .filter(item => item.value > 0 && item.value <= 500)
                    .sort((a, b) => b.score - a.score);
                  return Number(inputs[0]?.value || 0);

                  function score(input) {
                    const text = nearbyText(input);
                    const rect = input.getBoundingClientRect();
                    let value = 0;
                    if (text.includes('每页')) value += 100;
                    if (text.includes('共有')) value += 80;
                    if (text.includes('条')) value += 60;
                    if (String(input.className || '').includes('IPT_input')) value += 20;
                    value += Math.max(0, rect.top / 1000);
                    value += Math.max(0, rect.left / 1000);
                    return value;
                  }
                  function nearbyText(el) {
                    const chunks = [];
                    let node = el;
                    for (let depth = 0; node && depth < 5; depth += 1, node = node.parentElement) {
                      chunks.push(node.innerText || node.textContent || '');
                    }
                    return norm(chunks.join(' '));
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """
            )
            or 0
        )

    def _locate_jit_product_page_size_control(self, page):
        return page.evaluate(
            """
            () => {
              window.scrollTo({ top: document.body.scrollHeight, behavior: 'instant' });
              const candidates = [...document.querySelectorAll('input,.PGT_sizeChanger_5-120-1,.ST_outerWrapper_5-120-1,li,div')]
                .filter(visible)
                .filter(el => {
                  const text = nearbyText(el);
                  const value = el.value || '';
                  return text.includes('每页') || /^\\d+$/.test(value);
                })
                .map(el => {
                  const target = el.closest('.PGT_sizeChanger_5-120-1,.ST_outerWrapper_5-120-1,li') || el;
                  return { el: target, score: score(target) };
                })
                .filter(item => visible(item.el))
                .sort((a, b) => b.score - a.score);
              const target = candidates[0]?.el || null;
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function score(el) {
                const text = nearbyText(el);
                const rect = el.getBoundingClientRect();
                let value = 0;
                if (text.includes('每页')) value += 100;
                if (text.includes('共有')) value += 70;
                if (text.includes('条')) value += 60;
                if (String(el.className || '').includes('PGT_sizeChanger')) value += 80;
                if (String(el.className || '').includes('ST_outerWrapper')) value += 60;
                value += Math.max(0, rect.top / 1000);
                value += Math.max(0, rect.left / 1000);
                return value;
              }
              function nearbyText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 4; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                }
                return norm(chunks.join(' '));
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
            }
            """
        )

    def _click_jit_page_size_option(self, page, size: int) -> bool:
        targets = [f"{size} 条/页", f"{size}条/页", str(size)]
        for target in targets:
            if self._activate_visible_text_item(page, target, exact=True):
                return True
        return False

    def _fill_jit_page_size_control(self, page, size: int) -> None:
        rect = self._locate_jit_product_page_size_control(page)
        if not rect:
            return
        page.mouse.click(rect["x"], rect["y"])
        page.keyboard.press("Control+A")
        page.keyboard.type(str(size), delay=20)
        page.keyboard.press("Enter")

    def _select_all_jit_products(self, page) -> int:
        self.log("正在全选JIT商品选择列表。")
        for attempt in range(1, 6):
            rect = self._locate_jit_header_checkbox(page)
            if not rect:
                self.log(f"第 {attempt} 次没有找到JIT表头全选框，继续等待。")
                page.wait_for_timeout(900)
                continue
            if rect.get("checked"):
                return self._read_jit_selected_count_or_page_size(page)
            page.mouse.move(rect["x"], rect["y"])
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(1000)
            if self._jit_header_checkbox_checked(page):
                selected_count = self._read_jit_selected_count_or_page_size(page)
                self.log(f"已全选JIT商品：{selected_count} 个。")
                return selected_count
            self._click_jit_checkbox_by_dom_index(page, int(rect.get("index", 0)))
            page.wait_for_timeout(1000)
            if self._jit_header_checkbox_checked(page):
                selected_count = self._read_jit_selected_count_or_page_size(page)
                self.log(f"已全选JIT商品：{selected_count} 个。")
                return selected_count
            self.log(f"第 {attempt} 次点击JIT全选后未确认选中，准备重试。")
        raise RuntimeError("已尝试点击JIT表头全选框，但没有检测到选中状态。")

    def _locate_jit_header_checkbox(self, page):
        return page.evaluate(
            """
            ([goodsInfoText]) => {
              window.scrollTo({ top: 0, behavior: 'instant' });
              const items = [...document.querySelectorAll(
                'label[data-testid="beast-core-checkbox"],[data-testid="beast-core-checkbox"],input[type="checkbox"],[role="checkbox"]'
              )]
                .map((el, index) => ({ el, index }))
                .filter(item => visible(item.el))
                .map(item => {
                  const target = item.el.closest('label') || item.el;
                  const rect = target.getBoundingClientRect();
                  const text = ancestorText(target);
                  const inHeader = Boolean(target.closest('thead,[class*="thead"],[class*="THead"],[class*="header"],[class*="Header"],th'));
                  return {
                    index: item.index,
                    x: rect.left + rect.width / 2,
                    y: rect.top + rect.height / 2,
                    checked: isChecked(item.el) || isChecked(target),
                    inHeader,
                    score: scoreCandidate(target, rect, text, goodsInfoText, inHeader),
                  };
                })
                .filter(item => item.inHeader)
                .filter(item => item.y > 40 && item.y < window.innerHeight - 40)
                .sort((a, b) => b.score - a.score || a.y - b.y || a.x - b.x);
              return items[0] || null;

              function scoreCandidate(el, rect, text, goodsInfoText, inHeader) {
                let score = 0;
                if (inHeader) score += 2000;
                if (text.includes(goodsInfoText)) score += 800;
                if (!text.includes('SPU') && !text.includes('SKC')) score += 160;
                if (rect.left < 520) score += 140;
                if (rect.top < 760) score += 120;
                if (String(el.className || '').includes('CBX')) score += 80;
                score -= Math.abs(rect.left - 293) * 0.4;
                score -= Math.abs(rect.top - 414) * 0.2;
                return score;
              }
              function ancestorText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                  const text = chunks.join(' ');
                  if (text.includes(goodsInfoText)) break;
                }
                return norm(chunks.join(' '));
              }
              function isChecked(el) {
                return el.getAttribute('data-checked') === 'true'
                  || el.getAttribute('aria-checked') === 'true'
                  || el.checked === true
                  || el.querySelector?.('input[type="checkbox"]')?.checked === true
                  || /checked/i.test(String(el.className || ''));
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden'
                  && rect.bottom >= 0
                  && rect.right >= 0
                  && rect.top <= window.innerHeight
                  && rect.left <= window.innerWidth;
              }
            }
            """,
            [TEXT_GOODS_INFO],
        )

    def _click_jit_checkbox_by_dom_index(self, page, index: int) -> None:
        page.evaluate(
            """
            (targetIndex) => {
              const items = [...document.querySelectorAll(
                'label[data-testid="beast-core-checkbox"],[data-testid="beast-core-checkbox"],input[type="checkbox"],[role="checkbox"]'
              )];
              const el = items[targetIndex];
              if (!el) return false;
              const target = el.closest('label') || el;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              const clientX = rect.left + rect.width / 2;
              const clientY = rect.top + rect.height / 2;
              for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                target.dispatchEvent(new MouseEvent(type, {
                  bubbles: true,
                  cancelable: true,
                  view: window,
                  clientX,
                  clientY,
                }));
              }
              target.click?.();
              return true;
            }
            """,
            index,
        )

    def _jit_header_checkbox_checked(self, page) -> bool:
        rect = self._locate_jit_header_checkbox(page)
        return bool(rect and rect.get("checked"))

    def _read_jit_selected_count_or_page_size(self, page) -> int:
        count = int(
            page.evaluate(
                """
                () => {
                  const text = String(document.body.innerText || '');
                  const patterns = [
                    /已选[:：]?\\s*(\\d+)/,
                    /选中[:：]?\\s*(\\d+)/,
                    /共选择\\s*(\\d+)/
                  ];
                  for (const pattern of patterns) {
                    const match = text.match(pattern);
                    if (match) return Number(match[1]);
                  }
                  return 0;
                }
                """
            )
            or 0
        )
        if count > 0:
            return count
        return self._read_jit_product_page_size(page) or JIT_TARGET_PAGE_SIZE

    def _open_batch_adjust_jit_menu(self, page) -> None:
        self.log("正在打开“批量调整JIT”菜单。")
        for attempt in range(1, 5):
            if not self._click_batch_adjust_jit_control(page) and not self._click_button_by_text(page, TEXT_BATCH_ADJUST_JIT):
                if not self._click_visible_text_as_button(page, TEXT_BATCH_ADJUST_JIT):
                    raise RuntimeError("没有找到或无法点击“批量调整JIT”按钮。")
            page.wait_for_timeout(700)
            if self._wait_for_jit_popup_item(page, TEXT_BATCH_OPEN_JIT, timeout=5000):
                return
            if self._wait_for_popup_item(page, TEXT_BATCH_OPEN_JIT, exact=True, timeout=1200):
                return
            self.log(f"第 {attempt} 次没有等到“批量开通JIT”菜单项，准备重试。")
            page.keyboard.press("Escape")
            page.wait_for_timeout(800)
        raise RuntimeError("已点击“批量调整JIT”，但没有检测到“批量开通JIT”菜单项。")

    def _click_batch_open_jit(self, page) -> None:
        self.log("正在点击“批量开通JIT”。")
        for attempt in range(1, 4):
            if self._activate_jit_popup_item(page, TEXT_BATCH_OPEN_JIT):
                page.wait_for_timeout(1000)
                return
            if self._activate_popup_item(page, TEXT_BATCH_OPEN_JIT, exact=True):
                page.wait_for_timeout(1000)
                return
            self.log(f"第 {attempt} 次点击“批量开通JIT”失败，准备重试。")
            page.wait_for_timeout(700)
        raise RuntimeError("没有成功点击“批量开通JIT”。")

    def _click_batch_adjust_jit_control(self, page) -> bool:
        result = page.evaluate(
            """
            ([targetText]) => {
              const targetCompact = compact(targetText);
              const candidates = [...document.querySelectorAll(
                '.ST_outerWrapper_5-120-1,.ST_head_5-120-1,.ST_selectValueSingle_5-120-1,'
                + '[class*="ST_outerWrapper"],[class*="ST_head"],[class*="ST_selectValueSingle"],'
                + 'button,[role="button"],span,div'
              )]
                .filter(visible)
                .filter(el => compact(el.innerText || el.textContent) === targetCompact)
                .map(el => el.closest('.ST_outerWrapper_5-120-1,[class*="ST_outerWrapper"],button,[role="button"]') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .filter((el, index, list) => list.indexOf(el) === index)
                .sort((a, b) => score(a) - score(b))[0];
              if (!candidates) return null;
              candidates.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = candidates.getBoundingClientRect();
              const clientX = rect.left + rect.width / 2;
              const clientY = rect.top + rect.height / 2;
              for (const eventType of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                candidates.dispatchEvent(new MouseEvent(eventType, {
                  bubbles: true,
                  cancelable: true,
                  view: window,
                  clientX,
                  clientY,
                }));
              }
              candidates.click?.();
              return { x: clientX, y: clientY };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''));
              }
              function score(el) {
                const rect = el.getBoundingClientRect();
                let value = rect.width * rect.height;
                const cls = String(el.className || '');
                if (/ST_outerWrapper/.test(cls)) value -= 10000;
                if (/ST_head/.test(cls)) value -= 5000;
                return value;
              }
            }
            """,
            [TEXT_BATCH_ADJUST_JIT],
        )
        if not result:
            return False
        page.mouse.move(result["x"], result["y"])
        page.mouse.down()
        page.mouse.up()
        return True

    def _wait_for_jit_popup_item(self, page, text: str, timeout: int = 3000) -> bool:
        try:
            page.wait_for_function(
                """
                ([targetText]) => {
                  const targetCompact = compact(targetText);
                  return Boolean(findJitPopupItem(targetCompact));

                  function findJitPopupItem(targetCompact) {
                    return [...document.querySelectorAll(
                      '.PT_outerWrapper_5-120-1,.ST_dropdown_5-120-1,.cIL_item_5-120-1,'
                      + '[class*="PT_outerWrapper"],[class*="ST_dropdown"],[class*="cIL_item"],'
                      + 'li[role="option"],[role="option"],ul[role="listbox"],[role="listbox"]'
                    )]
                      .filter(visible)
                      .filter(isJitPopupRelated)
                      .find(el => compact(el.innerText || el.textContent) === targetCompact) || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isJitPopupRelated(el) {
                    let node = el;
                    for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                      const cls = String(node.className || '');
                      const role = node.getAttribute?.('role') || '';
                      if (/ST_dropdown|PT_outerWrapper|PT_portal|cIL_item/.test(cls)) return true;
                      if (/listbox|option/.test(role)) return true;
                    }
                    return false;
                  }
                }
                """,
                [text],
                timeout=timeout,
            )
            return True
        except Exception:  # noqa: BLE001
            return False

    def _activate_jit_popup_item(self, page, text: str) -> bool:
        rect = page.evaluate(
            """
            ([targetText]) => {
              const targetCompact = compact(targetText);
              const item = [...document.querySelectorAll(
                '.ST_outerWrapper_5-120-1,.PT_outerWrapper_5-120-1,.ST_dropdown_5-120-1,'
                + '.cIL_item_5-120-1,[class*="ST_outerWrapper"],[class*="PT_outerWrapper"],'
                + '[class*="ST_dropdown"],[class*="cIL_item"],li[role="option"],[role="option"]'
              )]
                .filter(visible)
                .filter(isJitPopupRelated)
                .filter(el => compact(el.innerText || el.textContent) === targetCompact)
                .map(el => el.closest('li[role="option"],[role="option"],.cIL_item_5-120-1,[class*="cIL_item"],li,button') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .filter((el, index, list) => list.indexOf(el) === index)
                .sort((a, b) => area(a) - area(b))[0];
              if (!item) return null;
              item.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = item.getBoundingClientRect();
              const clientX = rect.left + rect.width / 2;
              const clientY = rect.top + rect.height / 2;
              for (const eventType of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                item.dispatchEvent(new MouseEvent(eventType, {
                  bubbles: true,
                  cancelable: true,
                  view: window,
                  clientX,
                  clientY,
                }));
              }
              item.click?.();
              return { x: clientX, y: clientY };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''));
              }
              function isJitPopupRelated(el) {
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  const cls = String(node.className || '');
                  const role = node.getAttribute?.('role') || '';
                  if (/ST_dropdown|PT_outerWrapper|PT_portal|cIL_item/.test(cls)) return true;
                  if (/listbox|option/.test(role)) return true;
                }
                return false;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [text],
        )
        if not rect:
            return False
        page.mouse.move(rect["x"], rect["y"])
        page.mouse.down()
        page.mouse.up()
        return True

    def _wait_for_jit_confirm_modal(self, page, timeout: int = 12000) -> int:
        try:
            page.wait_for_function(
                """
                ([confirmTexts, modalSelector]) => {
                  const dialogs = [...document.querySelectorAll(modalSelector)]
                    .filter(visible)
                    .filter(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text.includes('开通JIT')
                        && text.includes('SKC')
                        && confirmTexts.some(confirmText => text.includes(confirmText))
                        && /确认下列\\s*\\d+\\s*个?SKC/.test(text);
                    });
                  return dialogs.length > 0;

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                arg=[list(TEXT_JIT_CONFIRM_OPTIONS), jit_confirm_modal_selector()],
                timeout=timeout,
            )
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("点击“批量开通JIT”后没有检测到确认弹窗。") from exc
        count = self._read_jit_confirm_modal_count(page)
        self.log(f"已打开JIT开通确认弹窗，弹窗商品数量：{count or '未知'}。")
        return count

    def _has_jit_confirm_modal(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([modalSelector]) => {
                  return Boolean(findJitConfirmModal());

                  function findJitConfirmModal() {
                    return [...document.querySelectorAll(modalSelector)]
                      .filter(visible)
                      .find(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text.includes('开通JIT')
                        && text.includes('SKC')
                        && /确认下列\\s*\\d+\\s*个?SKC/.test(text);
                      }) || null;
                  }

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                [jit_confirm_modal_selector()],
            )
        )

    def _read_jit_confirm_modal_count(self, page) -> int:
        return int(
            page.evaluate(
                """
                ([modalSelector]) => {
                  const dialogs = [...document.querySelectorAll(modalSelector)]
                    .filter(visible)
                    .map(el => ({
                      text: norm(el.innerText || el.textContent),
                      area: area(el),
                    }))
                    .filter(item => item.text.includes('开通JIT') && item.text.includes('确认下列'))
                    .sort((a, b) => a.area - b.area);
                  const text = dialogs[0]?.text || '';
                  const match = text.match(/确认下列\\s*(\\d+)\\s*个?SKC[^?？]*全部开通JIT/);
                  return Number(match?.[1] || 0);

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [jit_confirm_modal_selector()],
            )
            or 0
        )

    def _confirm_jit_open(self, page) -> None:
        if not self._click_jit_confirm_button(page):
            raise RuntimeError("没有找到JIT确认弹窗中的“确认”按钮。")
        self._wait_for_page_ready(page, timeout=15000)
        page.wait_for_timeout(1800)
        if self._click_global_button_if_visible(page, TEXT_I_KNOW, timeout=15000):
            self.log("已点击“我知道了”。")
            return
        if self._click_visible_text_as_button(page, TEXT_I_KNOW):
            self.log("已点击“我知道了”。")

    def _confirm_all_jit_open_modals(self, page) -> None:
        confirmed = 0
        for attempt in range(1, 5):
            self._scroll_jit_confirm_modal_to_bottom(page)
            if not self._has_jit_confirm_modal(page):
                break
            if not self._click_jit_confirm_button(page):
                raise RuntimeError("检测到JIT确认弹窗，但没有找到可点击的“确认”按钮。")
            confirmed += 1
            self._wait_for_page_ready(page, timeout=15000)
            page.wait_for_timeout(1600)
            if self._click_global_button_if_visible(page, TEXT_I_KNOW, timeout=6000):
                self.log("已点击“我知道了”。")
            elif self._click_visible_text_as_button(page, TEXT_I_KNOW):
                self.log("已点击“我知道了”。")
            self._scroll_jit_confirm_modal_to_bottom(page)
            if not self._has_jit_confirm_modal(page):
                break
            self.log(f"第 {attempt} 次确认后仍检测到JIT确认弹窗，继续处理残留弹窗。")
        if self._has_jit_confirm_modal(page):
            raise RuntimeError("JIT确认后仍存在未关闭的确认弹窗，请人工检查页面状态。")
        self.log(f"已处理JIT确认弹窗：{confirmed} 个。")

    def _click_jit_confirm_button(self, page) -> bool:
        self._scroll_jit_confirm_modal_to_bottom(page)
        rect = page.evaluate(
            """
            ([confirmTexts, modalSelector]) => {
              const modal = findJitConfirmModal();
              if (!modal) return null;
              const modalRoot = modal.closest('[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-dialog,.MDL_outerWrapper_5-120-1,.MDL_container_5-120-1,.MDL_innerWrapper_5-120-1,[class*="MDL_outerWrapper"],[class*="MDL_container"],[class*="MDL_innerWrapper"]')
                || modal.parentElement
                || modal;
              const button = [...modalRoot.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => confirmTexts.includes(norm(el.innerText || el.textContent)))
                .map(el => el.closest('button,[role="button"],a') || el)
                .filter((el, index, list) => list.indexOf(el) === index)
                .filter(el => visible(el) && !isDisabled(el))
                .sort((a, b) => area(a) - area(b))[0];
              if (!button) return null;
              button.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = button.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findJitConfirmModal() {
                const candidates = [...document.querySelectorAll(modalSelector)]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    return text.includes('开通JIT')
                      && text.includes('SKC')
                      && confirmTexts.some(confirmText => text.includes(confirmText))
                      && /确认下列\\s*\\d+\\s*个?SKC/.test(text);
                  })
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''));
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [list(TEXT_JIT_CONFIRM_OPTIONS), jit_confirm_modal_selector()],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        return True

    def _scroll_jit_confirm_modal_to_bottom(self, page) -> None:
        page.evaluate(
            """
            ([modalSelector]) => {
              for (const modal of findJitConfirmModals()) {
                const modalRoot = modal.closest('[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-dialog,.MDL_outerWrapper_5-120-1,.MDL_container_5-120-1,.MDL_innerWrapper_5-120-1,[class*="MDL_outerWrapper"],[class*="MDL_container"],[class*="MDL_innerWrapper"]')
                  || modal.parentElement
                  || modal;
                const nodes = [modal, modalRoot, ...modalRoot.querySelectorAll('*')].filter(visible);
                for (const node of nodes) {
                  if (node.scrollHeight > node.clientHeight + 8) {
                    node.scrollTop = node.scrollHeight;
                  }
                }
                modalRoot.scrollTop = modalRoot.scrollHeight;
              }

              function findJitConfirmModals() {
                return [...document.querySelectorAll(modalSelector)]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    return text.includes('开通JIT')
                      && text.includes('SKC')
                      && /确认下列\\s*\\d+\\s*个?SKC/.test(text);
                  })
                  .sort((a, b) => area(a) - area(b));
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [jit_confirm_modal_selector()],
        )
        page.wait_for_timeout(300)

    def _ensure_stock_sale_manage_page(self, page) -> None:
        if "agentseller.temu.com/stock/fully-mgt/sale-manage/main" not in page.url:
            raise RuntimeError(f"当前不在销售管理库存页面：{page.url}")
        try:
            page.wait_for_function(
                """
                ([saleManageText, settingText, goodsInfoText]) => {
                  const text = String(document.body.innerText || '');
                  return location.href.includes('/stock/fully-mgt/sale-manage/main')
                    && text.length > 0
                    && (text.includes(saleManageText) || text.includes(settingText))
                    && text.includes(goodsInfoText);
                }
                """,
                arg=[TEXT_STOCK_SALE_MANAGE, TEXT_EXPECTED_ARRIVAL_AREA_SETTING, TEXT_GOODS_INFO],
                timeout=12000,
            )
        except Exception as exc:  # noqa: BLE001
            signals = self._read_stock_sale_manage_signals(page)
            raise RuntimeError(f"未检测到销售管理库存页面内容。页面信号：{signals}") from exc

    def _open_expected_arrival_area_drawer(self, page) -> None:
        if self._has_expected_arrival_drawer(page):
            self.log("检测到期望到货区域设置弹窗已打开，继续操作。")
            return
        self.log("正在打开期望到货区域设置弹窗。")
        for attempt in range(1, 5):
            if not self._click_button_by_text(page, TEXT_EXPECTED_ARRIVAL_AREA_SETTING):
                if not self._click_visible_text_as_button(page, TEXT_EXPECTED_ARRIVAL_AREA_SETTING):
                    raise RuntimeError("没有找到或无法点击“期望到货区域设置”按钮。")
            try:
                self._wait_for_expected_arrival_drawer(page, timeout=12000)
                page.wait_for_timeout(1200)
                return
            except Exception:  # noqa: BLE001
                self.log(f"第 {attempt} 次点击后没有检测到期望到货区域弹窗，准备重试。")
                page.wait_for_timeout(900)
        raise RuntimeError("多次尝试后仍没有打开“期望到货区域设置”弹窗。")

    def _has_expected_arrival_drawer(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([titleText, allText, selectedText]) => {
                  return Boolean(findDrawer(titleText, allText, selectedText));

                  function findDrawer(titleText, allText, selectedText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .find(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes(allText)
                          && text.includes(selectedText)
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      }) || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, TEXT_ALL_JIT_CUSTOM_PRODUCTS, TEXT_SELECTED_JIT_CUSTOM_PRODUCTS],
            )
        )

    def _wait_for_expected_arrival_drawer(self, page, timeout: int = 10000) -> None:
        page.wait_for_function(
            """
            ([titleText, allText, selectedText]) => {
              return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                .filter(visible)
                .some(el => {
                  const text = norm(el.innerText || el.textContent);
                  const rect = el.getBoundingClientRect();
                  return text.includes(titleText)
                    && text.includes(allText)
                    && text.includes(selectedText)
                    && rect.width > window.innerWidth * 0.35
                    && rect.height > window.innerHeight * 0.45;
                });

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
            }
            """,
            arg=[TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, TEXT_ALL_JIT_CUSTOM_PRODUCTS, TEXT_SELECTED_JIT_CUSTOM_PRODUCTS],
            timeout=timeout,
        )

    def _set_expected_arrival_page_size(self, page, size: int) -> None:
        if size <= 0:
            raise RuntimeError("期望到货区域分页数量必须大于 0。")
        self.log(f"正在设置期望到货区域弹窗每页 {size} 条。")
        for attempt in range(1, 6):
            current = self._read_expected_arrival_page_size(page)
            if current == size:
                self.log(f"期望到货区域弹窗当前已经是每页 {size} 条。")
                return
            rect = self._locate_expected_arrival_page_size_control(page)
            if not rect:
                self.log(f"第 {attempt} 次没有找到期望到货区域分页条数控件，继续等待。")
                page.wait_for_timeout(900)
                continue
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(600)
            if self._click_expected_arrival_page_size_option(page, size):
                self._wait_for_page_ready(page, timeout=10000)
                page.wait_for_timeout(1200)
                if self._read_expected_arrival_page_size(page) == size:
                    return
            self._fill_expected_arrival_page_size_control(page, size)
            self._wait_for_page_ready(page, timeout=10000)
            page.wait_for_timeout(1200)
            if self._read_expected_arrival_page_size(page) == size:
                return
            self.log(f"第 {attempt} 次设置期望到货区域分页后还未生效，准备重试。")
        raise RuntimeError(f"期望到货区域分页条数切换后未检测到 {size} 条/页。")

    def _read_expected_arrival_page_size(self, page) -> int:
        return int(
            page.evaluate(
                """
                ([titleText]) => {
                  const drawer = findDrawer(titleText);
                  if (!drawer) return 0;
                  const leftPanel = findLeftPanel(drawer);
                  const inputs = [...leftPanel.querySelectorAll('input')]
                    .filter(visible)
                    .filter(input => /^\\d+$/.test(input.value || ''))
                    .map(input => ({ input, value: Number(input.value), score: score(input) }))
                    .filter(item => item.value > 0 && item.value <= 500)
                    .sort((a, b) => b.score - a.score);
                  return Number(inputs[0]?.value || 0);

                  function findDrawer(titleText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes('全部JIT/定制品')
                          && text.includes('已选JIT/定制品')
                          && text.includes('期望到货区域')
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function findLeftPanel(drawer) {
                    const panels = [...drawer.querySelectorAll('div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        const drawerRect = drawer.getBoundingClientRect();
                        return text.includes('共有')
                          && text.includes('每页')
                          && text.includes('序号')
                          && text.includes('商品信息')
                          && rect.left < drawerRect.left + drawerRect.width * 0.58;
                      })
                      .sort((a, b) => area(a) - area(b));
                    return panels[0] || drawer;
                  }
                  function score(input) {
                    const text = nearbyText(input);
                    const rect = input.getBoundingClientRect();
                    let value = 0;
                    if (text.includes('每页')) value += 100;
                    if (text.includes('共有')) value += 80;
                    if (text.includes('条')) value += 60;
                    if (String(input.className || '').includes('IPT_input')) value += 20;
                    value += Math.max(0, rect.top / 1000);
                    return value;
                  }
                  function nearbyText(el) {
                    const chunks = [];
                    let node = el;
                    for (let depth = 0; node && depth < 5; depth += 1, node = node.parentElement) {
                      chunks.push(node.innerText || node.textContent || '');
                    }
                    return norm(chunks.join(' '));
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA],
            )
            or 0
        )

    def _locate_expected_arrival_page_size_control(self, page):
        return page.evaluate(
            """
            ([titleText]) => {
              const drawer = findDrawer(titleText);
              if (!drawer) return null;
              const leftPanel = findLeftPanel(drawer);
              const candidates = [...leftPanel.querySelectorAll('input,.PGT_sizeChanger_5-120-1,.ST_outerWrapper_5-120-1,li,div')]
                .filter(visible)
                .filter(el => {
                  const text = nearbyText(el);
                  const value = el.value || '';
                  return text.includes('每页') || /^\\d+$/.test(value);
                })
                .map(el => {
                  const target = el.closest('.PGT_sizeChanger_5-120-1,.ST_outerWrapper_5-120-1,li') || el;
                  return { el: target, score: score(target) };
                })
                .filter(item => visible(item.el))
                .sort((a, b) => b.score - a.score);
              const target = candidates[0]?.el || null;
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findDrawer(titleText) {
                return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    const rect = el.getBoundingClientRect();
                    return text.includes(titleText)
                      && text.includes('全部JIT/定制品')
                      && text.includes('已选JIT/定制品')
                      && text.includes('期望到货区域')
                      && rect.width > window.innerWidth * 0.35
                      && rect.height > window.innerHeight * 0.45;
                  })
                  .sort((a, b) => area(a) - area(b))[0] || null;
              }
              function findLeftPanel(drawer) {
                const panels = [...drawer.querySelectorAll('div')]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    const rect = el.getBoundingClientRect();
                    const drawerRect = drawer.getBoundingClientRect();
                    return text.includes('共有')
                      && text.includes('每页')
                      && text.includes('序号')
                      && text.includes('商品信息')
                      && rect.left < drawerRect.left + drawerRect.width * 0.58;
                  })
                  .sort((a, b) => area(a) - area(b));
                return panels[0] || drawer;
              }
              function score(el) {
                const text = nearbyText(el);
                const rect = el.getBoundingClientRect();
                let value = 0;
                if (text.includes('每页')) value += 100;
                if (text.includes('共有')) value += 70;
                if (text.includes('条')) value += 60;
                if (String(el.className || '').includes('PGT_sizeChanger')) value += 80;
                if (String(el.className || '').includes('ST_outerWrapper')) value += 60;
                value += Math.max(0, rect.top / 1000);
                return value;
              }
              function nearbyText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 4; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                }
                return norm(chunks.join(' '));
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA],
        )

    def _click_expected_arrival_page_size_option(self, page, size: int) -> bool:
        for target in (f"{size} 条/页", f"{size}条/页", str(size)):
            if self._activate_visible_text_item(page, target, exact=True):
                return True
        return False

    def _fill_expected_arrival_page_size_control(self, page, size: int) -> None:
        rect = self._locate_expected_arrival_page_size_control(page)
        if not rect:
            return
        page.mouse.click(rect["x"], rect["y"])
        page.keyboard.press("Control+A")
        page.keyboard.type(str(size), delay=20)
        page.keyboard.press("Enter")

    def _filter_expected_arrival_current_area(self, page, option_text: str) -> None:
        self.log(f"正在筛选当前期望到货区域：{option_text}。")
        for attempt in range(1, 5):
            if not self._expected_arrival_current_area_selected(page, option_text):
                rect = self._locate_expected_arrival_current_area_select(page)
                if not rect:
                    raise RuntimeError("没有找到“当前期望到货区域”筛选下拉框。")
                page.mouse.click(rect["x"], rect["y"])
                page.wait_for_timeout(600)
                if not self._click_expected_arrival_current_area_option(page, option_text):
                    self.log(f"第 {attempt} 次没有点中“{option_text}”选项，准备重试。")
                    page.wait_for_timeout(600)
                    continue
                page.wait_for_timeout(900)
                if not self._expected_arrival_current_area_selected(page, option_text):
                    self.log(f"第 {attempt} 次点击“{option_text}”后未检测到筛选值，准备重试。")
                    continue
            self._click_expected_arrival_drawer_query(page)
            self._wait_for_page_ready(page, timeout=10000)
            page.wait_for_timeout(1500)
            return
        raise RuntimeError(f"没有成功筛选当前期望到货区域“{option_text}”。")

    def _click_expected_arrival_current_area_option(self, page, option_text: str) -> bool:
        rect = page.evaluate(
            """
            ([optionText]) => {
              const dropdowns = [...document.querySelectorAll('body *')]
                .filter(visible)
                .filter(el => {
                  const cls = String(el.className || '');
                  const role = el.getAttribute?.('role') || '';
                  const text = norm(el.innerText || el.textContent);
                  return text.includes(optionText)
                    && (/dropdown|popover|popper|portal|select/i.test(cls) || /listbox|option|menu/.test(role));
                })
                .sort((a, b) => area(a) - area(b));
              for (const dropdown of dropdowns) {
                const candidates = [...dropdown.querySelectorAll('[role="option"],li,label,span,div')]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent) === optionText)
                  .sort((a, b) => area(a) - area(b));
                const item = candidates[0];
                if (!item) continue;
                const target = clickableAncestor(item, dropdown);
                target.scrollIntoView({ block: 'center', inline: 'center' });
                const rect = target.getBoundingClientRect();
                return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };
              }
              return null;

              function clickableAncestor(item, dropdown) {
                let node = item;
                let best = item;
                for (let depth = 0; node && node !== dropdown && depth < 5; depth += 1, node = node.parentElement) {
                  const cls = String(node.className || '');
                  const role = node.getAttribute?.('role') || '';
                  const text = norm(node.innerText || node.textContent);
                  if (text === optionText && (/option|item|cell|row/i.test(cls) || /option|menuitem/.test(role))) {
                    return node;
                  }
                  if (text === optionText && area(node) >= area(best)) {
                    best = node;
                  }
                }
                return best;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [option_text],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        return True

    def _locate_expected_arrival_current_area_select(self, page):
        return page.evaluate(
            """
            ([titleText, labelText]) => {
              const drawer = findDrawer(titleText);
              if (!drawer) return null;
              const labels = [...drawer.querySelectorAll('span,div,label')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === labelText)
                .filter(el => el.getBoundingClientRect().top < window.innerHeight * 0.55)
                .map(el => ({ el, score: labelScore(el, drawer) }))
                .sort((a, b) => b.score - a.score || area(a.el) - area(b.el));
              const label = labels[0]?.el || null;
              if (!label) return null;
              const labelRect = label.getBoundingClientRect();
              const candidates = [...drawer.querySelectorAll('.ST_outerWrapper_5-120-1,[class*="select"],[class*="Select"],input,div')]
                .filter(visible)
                .filter(el => {
                  const rect = el.getBoundingClientRect();
                  const text = norm(el.innerText || el.textContent || el.value);
                  const closeToLabel = Math.abs((rect.top + rect.height / 2) - (labelRect.top + labelRect.height / 2)) < 35
                    && rect.left > labelRect.right - 20;
                  return closeToLabel && (text.includes('全部') || text.includes('按照历史') || el.querySelector?.('input'));
                })
                .sort((a, b) => {
                  const aSelect = /ST_outerWrapper|select/i.test(String(a.className || '')) ? 0 : 1;
                  const bSelect = /ST_outerWrapper|select/i.test(String(b.className || '')) ? 0 : 1;
                  return aSelect - bSelect || area(a) - area(b);
                });
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findDrawer(titleText) {
                return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    const rect = el.getBoundingClientRect();
                    return text.includes(titleText)
                      && text.includes('全部JIT/定制品')
                      && text.includes('已选JIT/定制品')
                      && text.includes('期望到货区域')
                      && rect.width > window.innerWidth * 0.35
                      && rect.height > window.innerHeight * 0.45;
                  })
                  .sort((a, b) => area(a) - area(b))[0] || null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
              function labelScore(el, drawer) {
                const rect = el.getBoundingClientRect();
                const drawerRect = drawer.getBoundingClientRect();
                let score = 0;
                if (rect.top > drawerRect.top + 240 && rect.top < drawerRect.top + 380) score += 1000;
                if (rect.left > drawerRect.left + 250 && rect.left < drawerRect.left + 520) score += 500;
                if (rect.top < drawerRect.top + 430) score += 100;
                return score;
              }
            }
            """,
            [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, "当前期望到货区域"],
        )

    def _expected_arrival_current_area_selected(self, page, option_text: str) -> bool:
        return bool(
            page.evaluate(
                """
                ([titleText, labelText, optionText]) => {
                  const drawer = findDrawer(titleText);
                  if (!drawer) return false;
                  const labels = [...drawer.querySelectorAll('span,div,label')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent) === labelText)
                    .filter(el => el.getBoundingClientRect().top < window.innerHeight * 0.55)
                    .map(el => ({ el, score: labelScore(el, drawer) }))
                    .sort((a, b) => b.score - a.score || area(a.el) - area(b.el));
                  const label = labels[0]?.el || null;
                  if (!label) return false;
                  const labelRect = label.getBoundingClientRect();
                  const optionCompact = compact(optionText);
                  return [...drawer.querySelectorAll('span,div,input')]
                    .filter(visible)
                    .some(el => {
                      const rect = el.getBoundingClientRect();
                      const text = norm(el.innerText || el.textContent || el.value);
                      const textCompact = compact(text).replace(/[.。…]+$/g, '');
                      const looksSelected = text.includes(optionText)
                        || (textCompact.length >= 6 && optionCompact.includes(textCompact))
                        || textCompact.includes(optionCompact);
                      return looksSelected
                        && Math.abs((rect.top + rect.height / 2) - (labelRect.top + labelRect.height / 2)) < 45
                        && rect.left > labelRect.right - 20;
                    });

                  function findDrawer(titleText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes('全部JIT/定制品')
                          && text.includes('已选JIT/定制品')
                          && text.includes('期望到货区域')
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                  function labelScore(el, drawer) {
                    const rect = el.getBoundingClientRect();
                    const drawerRect = drawer.getBoundingClientRect();
                    let score = 0;
                    if (rect.top > drawerRect.top + 240 && rect.top < drawerRect.top + 380) score += 1000;
                    if (rect.left > drawerRect.left + 250 && rect.left < drawerRect.left + 520) score += 500;
                    if (rect.top < drawerRect.top + 430) score += 100;
                    return score;
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, "当前期望到货区域", option_text],
            )
        )

    def _click_expected_arrival_drawer_query(self, page) -> None:
        rect = page.evaluate(
            """
            ([titleText, queryText]) => {
              const drawer = findDrawer(titleText);
              if (!drawer) return null;
              const buttons = [...drawer.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === queryText)
                .map(el => el.closest('button,[role="button"],a') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .filter(el => {
                  const rect = el.getBoundingClientRect();
                  const drawerRect = drawer.getBoundingClientRect();
                  return rect.left >= drawerRect.left
                    && rect.right <= drawerRect.right
                    && rect.top > drawerRect.top + 120
                    && rect.top < drawerRect.top + drawerRect.height * 0.65;
                })
                .sort((a, b) => area(a) - area(b));
              const target = buttons[0] || findGlobalQueryButton(titleText, queryText);
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findDrawer(titleText) {
                  return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.Drawer_visible_5-120-1,.Drawer_content_5-120-1,.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                    .filter(visible)
                    .filter(el => {
                      const text = norm(el.innerText || el.textContent);
                      const rect = el.getBoundingClientRect();
                      return text.includes(titleText)
                        && text.includes('全部JIT/定制品')
                        && text.includes('已选JIT/定制品')
                        && text.includes('期望到货区域')
                        && rect.width > window.innerWidth * 0.35
                        && rect.height > window.innerHeight * 0.45;
                    })
                    .sort((a, b) => drawerScore(b) - drawerScore(a) || area(a) - area(b))[0] || null;
                }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true';
              }
              function findGlobalQueryButton(titleText, queryText) {
                const seen = new Set();
                return [...document.querySelectorAll('button,[role="button"],a,span,div')]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent) === queryText)
                  .map(el => el.closest('button,[role="button"],a') || el)
                  .filter(el => {
                    if (!el || seen.has(el)) return false;
                    seen.add(el);
                    return true;
                  })
                  .filter(visible)
                  .filter(el => !isDisabled(el))
                  .filter(el => {
                    const rect = el.getBoundingClientRect();
                    if (rect.left < window.innerWidth * 0.45 || rect.top < 120 || rect.top > window.innerHeight * 0.65) {
                      return false;
                    }
                    let node = el;
                    for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                      const text = norm(node.innerText || node.textContent);
                      if (text.includes('当前期望到货区域') && text.includes('历史发货地')) return true;
                    }
                    return false;
                  })
                  .sort((a, b) => area(a) - area(b))[0] || null;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
              function drawerScore(el) {
                const cls = String(el.className || '');
                let score = 0;
                if (/Drawer_content|Drawer_outerWrapper|Drawer_visible/i.test(cls)) score += 1000;
                if (el.getBoundingClientRect().left > window.innerWidth * 0.35) score += 400;
                return score;
              }
            }
            """,
            [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, TEXT_QUERY],
        )
        if not rect:
            raise RuntimeError("没有找到期望到货区域弹窗内的“查询”按钮。")
        page.mouse.click(rect["x"], rect["y"])

    def _select_all_expected_arrival_products(self, page) -> int:
        self.log("正在全选期望到货区域商品。")
        for attempt in range(1, 6):
            rect = self._locate_expected_arrival_header_checkbox(page)
            if not rect:
                self.log(f"第 {attempt} 次没有找到期望到货区域表头全选框，继续等待。")
                page.wait_for_timeout(900)
                continue
            if rect.get("checked"):
                selected_count = self._read_expected_arrival_selected_count(page)
                return selected_count or self._read_expected_arrival_page_size(page)
            page.mouse.move(rect["x"], rect["y"])
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(1200)
            selected_count = self._read_expected_arrival_selected_count(page)
            if selected_count > 0:
                self.log(f"已全选期望到货区域商品：{selected_count} 个。")
                return selected_count
            if self._expected_arrival_header_checkbox_checked(page):
                selected_count = self._read_expected_arrival_page_size(page) or TARGET_PAGE_SIZE
                self.log(f"已检测到期望到货区域表头全选状态：{selected_count} 个。")
                return selected_count
            self._click_expected_arrival_checkbox_by_dom_index(page, int(rect.get("index", 0)))
            page.wait_for_timeout(1200)
            selected_count = self._read_expected_arrival_selected_count(page)
            if selected_count > 0:
                self.log(f"已全选期望到货区域商品：{selected_count} 个。")
                return selected_count
            if self._expected_arrival_header_checkbox_checked(page):
                selected_count = self._read_expected_arrival_page_size(page) or TARGET_PAGE_SIZE
                self.log(f"已检测到期望到货区域表头全选状态：{selected_count} 个。")
                return selected_count
            self.log(f"第 {attempt} 次点击期望到货区域全选后未检测到已选数量，准备重试。")
        raise RuntimeError("已尝试点击期望到货区域表头全选框，但没有检测到选中数量。")

    def _expected_arrival_header_checkbox_checked(self, page) -> bool:
        rect = self._locate_expected_arrival_header_checkbox(page)
        return bool(rect and rect.get("checked"))

    def _locate_expected_arrival_header_checkbox(self, page):
        return page.evaluate(
            """
            ([titleText]) => {
              const drawer = findDrawer(titleText);
              if (!drawer) return null;
              const leftPanel = findLeftPanel(drawer);
              const items = [...leftPanel.querySelectorAll(
                'label[data-testid="beast-core-checkbox"],[data-testid="beast-core-checkbox"],input[type="checkbox"],[role="checkbox"]'
              )]
                .map((el, index) => ({ el, index: globalIndex(el) }))
                .filter(item => visible(item.el))
                .map(item => {
                  const target = item.el.closest('label') || item.el;
                  const rect = target.getBoundingClientRect();
                  const text = ancestorText(target);
                  const inHeader = Boolean(target.closest('thead,[class*="thead"],[class*="THead"],[class*="header"],[class*="Header"],th'));
                  const topElement = document.elementFromPoint(rect.left + rect.width / 2, rect.top + rect.height / 2);
                  return {
                    index: item.index,
                    x: rect.left + rect.width / 2,
                    y: rect.top + rect.height / 2,
                    checked: isChecked(item.el) || isChecked(target) || isChecked(topElement) || isChecked(topElement?.closest?.('label')),
                    inHeader,
                    score: scoreCandidate(target, rect, text, inHeader),
                  };
                })
                .filter(item => item.inHeader)
                .filter(item => item.y > 80 && item.y < window.innerHeight - 100)
                .sort((a, b) => b.score - a.score || a.y - b.y || a.x - b.x);
              return items[0] || null;

              function findDrawer(titleText) {
                return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    const rect = el.getBoundingClientRect();
                    return text.includes(titleText)
                      && text.includes('全部JIT/定制品')
                      && text.includes('已选JIT/定制品')
                      && text.includes('期望到货区域')
                      && rect.width > window.innerWidth * 0.35
                      && rect.height > window.innerHeight * 0.45;
                  })
                  .sort((a, b) => area(a) - area(b))[0] || null;
              }
              function findLeftPanel(drawer) {
                const panels = [...drawer.querySelectorAll('div')]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    const rect = el.getBoundingClientRect();
                    const drawerRect = drawer.getBoundingClientRect();
                    return text.includes('序号')
                      && text.includes('商品信息')
                      && text.includes('历史发货地')
                      && text.includes('当前期望到货区域')
                      && rect.left < drawerRect.left + drawerRect.width * 0.58;
                  })
                  .sort((a, b) => area(a) - area(b));
                return panels[0] || drawer;
              }
              function globalIndex(el) {
                const all = [...document.querySelectorAll(
                  'label[data-testid="beast-core-checkbox"],[data-testid="beast-core-checkbox"],input[type="checkbox"],[role="checkbox"]'
                )];
                return all.indexOf(el);
              }
              function scoreCandidate(el, rect, text, inHeader) {
                let score = 0;
                if (inHeader) score += 2000;
                if (text.includes('序号')) score += 700;
                if (text.includes('商品信息')) score += 700;
                if (text.includes('历史发货地')) score += 500;
                if (text.includes('SKC ID')) score -= 800;
                if (String(el.className || '').includes('CBX')) score += 80;
                score -= rect.left * 0.05;
                score -= Math.abs(rect.top - 430) * 0.1;
                return score;
              }
              function ancestorText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                  const text = chunks.join(' ');
                  if (text.includes('序号') && text.includes('商品信息')) break;
                }
                return norm(chunks.join(' '));
              }
              function isChecked(el) {
                return el.getAttribute('data-checked') === 'true'
                  || el.getAttribute('aria-checked') === 'true'
                  || el.checked === true
                  || el.querySelector?.('input[type="checkbox"]')?.checked === true
                  || /checked/i.test(String(el.className || ''));
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden'
                  && rect.bottom >= 0
                  && rect.right >= 0
                  && rect.top <= window.innerHeight
                  && rect.left <= window.innerWidth;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA],
        )

    def _click_expected_arrival_checkbox_by_dom_index(self, page, index: int) -> None:
        page.evaluate(
            """
            (targetIndex) => {
              const items = [...document.querySelectorAll(
                'label[data-testid="beast-core-checkbox"],[data-testid="beast-core-checkbox"],input[type="checkbox"],[role="checkbox"]'
              )];
              const el = items[targetIndex];
              if (!el) return false;
              const target = el.closest('label') || el;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              const clientX = rect.left + rect.width / 2;
              const clientY = rect.top + rect.height / 2;
              for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                target.dispatchEvent(new MouseEvent(type, {
                  bubbles: true,
                  cancelable: true,
                  view: window,
                  clientX,
                  clientY,
                }));
              }
              target.click?.();
              return true;
            }
            """,
            index,
        )

    def _read_expected_arrival_selected_count(self, page) -> int:
        return int(
            page.evaluate(
                """
                ([titleText]) => {
                  const drawer = findDrawer(titleText);
                  if (!drawer) return 0;
                  const text = norm(drawer.innerText || drawer.textContent);
                  const patterns = [
                    /已选中\\s*(\\d+)\\s*个JIT\\/定制品商品/,
                    /已选[:：]?\\s*(\\d+)/,
                    /选中[:：]?\\s*(\\d+)/
                  ];
                  for (const pattern of patterns) {
                    const match = text.match(pattern);
                    if (match) return Number(match[1]);
                  }
                  return 0;

                  function findDrawer(titleText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes('全部JIT/定制品')
                          && text.includes('已选JIT/定制品')
                          && text.includes('期望到货区域')
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA],
            )
            or 0
        )

    def _select_expected_arrival_area(self, page, area_text: str) -> None:
        self.log(f"正在选择期望到货区域：{area_text}。")
        for attempt in range(1, 5):
            clicked = page.evaluate(
                """
                ([titleText, areaText]) => {
                  const drawer = findDrawer(titleText);
                  if (!drawer) return false;
                  const bottom = [...drawer.querySelectorAll('div')]
                    .filter(visible)
                    .filter(el => {
                      const text = norm(el.innerText || el.textContent);
                      const rect = el.getBoundingClientRect();
                      return text.includes('已选中')
                        && text.includes('期望到货区域')
                        && text.includes(areaText)
                        && rect.top > window.innerHeight * 0.65;
                    })
                    .sort((a, b) => area(a) - area(b))[0] || drawer;
                  const target = [...bottom.querySelectorAll('label,[role="radio"],input[type="radio"],span,div')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent || el.value) === areaText)
                    .map(el => el.closest('label,[role="radio"]') || el)
                    .filter(visible)
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!target) return false;
                  fireMouse(target);
                  target.click?.();
                  return true;

                  function findDrawer(titleText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes('全部JIT/定制品')
                          && text.includes('已选JIT/定制品')
                          && text.includes('期望到货区域')
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                  function fireMouse(el) {
                    el.scrollIntoView({ block: 'center', inline: 'center' });
                    const rect = el.getBoundingClientRect();
                    const clientX = rect.left + rect.width / 2;
                    const clientY = rect.top + rect.height / 2;
                    for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                      el.dispatchEvent(new MouseEvent(type, {
                        bubbles: true,
                        cancelable: true,
                        view: window,
                        clientX,
                        clientY,
                      }));
                    }
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, area_text],
            )
            page.wait_for_timeout(700)
            if clicked and self._expected_arrival_area_selected(page, area_text):
                return
            self.log(f"第 {attempt} 次选择“{area_text}”后未确认选中，准备重试。")
        raise RuntimeError(f"没有成功选中期望到货区域“{area_text}”。")

    def _expected_arrival_area_selected(self, page, area_text: str) -> bool:
        return bool(
            page.evaluate(
                """
                ([titleText, areaText]) => {
                  const drawer = findDrawer(titleText);
                  if (!drawer) return false;
                  const labels = [...drawer.querySelectorAll('label,[role="radio"],input[type="radio"],span,div')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent || el.value) === areaText);
                  for (const label of labels) {
                    let node = label.closest('label,[role="radio"]') || label;
                    for (let depth = 0; node && depth < 6; depth += 1, node = node.parentElement) {
                      const input = node.querySelector?.('input[type="radio"]');
                      if (input && input.checked) return true;
                      if (node.checked === true) return true;
                      if (node.getAttribute?.('aria-checked') === 'true') return true;
                      if (/checked/i.test(String(node.className || ''))) return true;
                    }
                  }
                  return false;

                  function findDrawer(titleText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes('全部JIT/定制品')
                          && text.includes('已选JIT/定制品')
                          && text.includes('期望到货区域')
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, area_text],
            )
        )

    def _click_expected_arrival_drawer_confirm(self, page) -> None:
        self.log("正在点击期望到货区域弹窗内的确认。")
        for attempt in range(1, 4):
            rect = page.evaluate(
                """
                ([titleText, confirmText]) => {
                  const drawer = findDrawer(titleText);
                  if (!drawer) return null;
                  const buttons = [...drawer.querySelectorAll('button,[role="button"],a,span,div')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent) === confirmText)
                    .map(el => el.closest('button,[role="button"],a') || el)
                    .filter(visible)
                    .filter(el => !isDisabled(el))
                    .filter(el => el.getBoundingClientRect().top > window.innerHeight * 0.65)
                    .sort((a, b) => area(a) - area(b));
                  const target = buttons[0];
                  if (!target) return null;
                  target.scrollIntoView({ block: 'center', inline: 'center' });
                  const rect = target.getBoundingClientRect();
                  return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

                  function findDrawer(titleText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes('全部JIT/定制品')
                          && text.includes('已选JIT/定制品')
                          && text.includes('期望到货区域')
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isDisabled(el) {
                    return el.disabled === true
                      || el.getAttribute('disabled') !== null
                      || el.getAttribute('aria-disabled') === 'true';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, TEXT_CONFIRM],
            )
            if rect:
                page.mouse.click(rect["x"], rect["y"])
                page.wait_for_timeout(1000)
                return
            self.log(f"第 {attempt} 次没有找到弹窗底部确认按钮，准备重试。")
            page.wait_for_timeout(700)
        raise RuntimeError("没有找到期望到货区域弹窗底部“确认”按钮。")

    def _wait_for_expected_arrival_confirm_modal(self, page, timeout: int = 12000) -> int:
        try:
            page.wait_for_function(
                """
                ([areaText, confirmText]) => {
                  return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_alert_5-120-1,.rocket-modal,div')]
                    .filter(visible)
                    .some(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text.includes(areaText)
                        && text.includes(confirmText)
                        && text.includes('确认将')
                        && text.includes('期望到货区域')
                        && /确认将\\s*\\d+\\s*个/.test(text);
                    });

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                arg=[TEXT_YIWU, TEXT_CONFIRM],
                timeout=timeout,
            )
        except Exception as exc:  # noqa: BLE001
            if self._expected_arrival_visible_rows_already_area(page, TEXT_YIWU):
                raise RuntimeError(f"当前页商品的期望到货区域已经是{TEXT_YIWU}，没有可修改为{TEXT_YIWU}的商品。") from exc
            raise RuntimeError("点击期望到货区域确认后没有检测到二次确认弹窗。") from exc
        count = self._read_expected_arrival_confirm_count(page)
        self.log(f"已打开期望到货区域二次确认弹窗，弹窗商品数量：{count or '未知'}。")
        return count

    def _expected_arrival_visible_rows_already_area(self, page, area_text: str) -> bool:
        return bool(
            page.evaluate(
                """
                ([titleText, areaText]) => {
                  const drawer = findDrawer(titleText);
                  if (!drawer) return false;
                  const rows = [...drawer.querySelectorAll('tbody tr,[data-testid="beast-core-table-body-tr"]')]
                    .filter(visible)
                    .map(row => norm(row.innerText || row.textContent))
                    .filter(text => text.includes('SKC ID') && text.includes('SPU ID'));
                  if (!rows.length) return false;
                  return rows.every(text => text.endsWith(areaText) || text.includes(`-- ${areaText}`));

                  function findDrawer(titleText) {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_innerWrapper_5-120-1,.MDL_content_5-120-1,div')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        const rect = el.getBoundingClientRect();
                        return text.includes(titleText)
                          && text.includes('全部JIT/定制品')
                          && text.includes('已选JIT/定制品')
                          && text.includes('期望到货区域')
                          && rect.width > window.innerWidth * 0.35
                          && rect.height > window.innerHeight * 0.45;
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_MODIFY_EXPECTED_ARRIVAL_AREA, area_text],
            )
        )

    def _has_expected_arrival_confirm_modal(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([areaText]) => {
                  return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_alert_5-120-1,.rocket-modal,div')]
                    .filter(visible)
                    .some(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text.includes(areaText)
                        && text.includes('确认将')
                        && text.includes('期望到货区域')
                        && /确认将\\s*\\d+\\s*个/.test(text);
                    });

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                [TEXT_YIWU],
            )
        )

    def _read_expected_arrival_confirm_count(self, page) -> int:
        return int(
            page.evaluate(
                """
                ([areaText]) => {
                  const dialogs = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_alert_5-120-1,.rocket-modal,div')]
                    .filter(visible)
                    .map(el => ({
                      text: norm(el.innerText || el.textContent),
                      area: area(el),
                    }))
                    .filter(item => item.text.includes(areaText)
                      && item.text.includes('确认将')
                      && item.text.includes('期望到货区域')
                      && /确认将\\s*\\d+\\s*个/.test(item.text))
                    .sort((a, b) => a.area - b.area);
                  const text = dialogs[0]?.text || '';
                  const match = text.match(/确认将\\s*(\\d+)\\s*个/);
                  return Number(match?.[1] || 0);

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_YIWU],
            )
            or 0
        )

    def _confirm_expected_arrival_area(self, page) -> None:
        rect = page.evaluate(
            """
            ([areaText, confirmText]) => {
              const dialogs = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.MDL_outerWrapper_5-120-1,.MDL_alert_5-120-1,.rocket-modal,div')]
                .filter(visible)
                .filter(el => {
                  const text = norm(el.innerText || el.textContent);
                  return text.includes(areaText)
                    && text.includes('确认将')
                    && text.includes('期望到货区域')
                    && /确认将\\s*\\d+\\s*个/.test(text);
                })
                .map(el => ({
                  el,
                  button: findConfirmButton(el),
                  modalRank: modalRank(el),
                  area: area(el),
                }))
                .sort((a, b) => Number(Boolean(b.button)) - Number(Boolean(a.button))
                  || b.modalRank - a.modalRank
                  || a.area - b.area);
              const dialog = dialogs[0]?.el || null;
              const button = dialogs[0]?.button || findGlobalConfirmButton();
              if (!button) return null;
              button.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = button.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true';
              }
              function findConfirmButton(scope) {
                const seen = new Set();
                return [...scope.querySelectorAll('button,[role="button"],a,span,div')]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent) === confirmText)
                  .map(el => el.closest('button,[role="button"],a') || el)
                  .filter(el => {
                    if (!el || seen.has(el)) return false;
                    seen.add(el);
                    return true;
                  })
                  .filter(visible)
                  .filter(el => !isDisabled(el))
                  .sort((a, b) => area(a) - area(b))[0] || null;
              }
              function findGlobalConfirmButton() {
                const seen = new Set();
                return [...document.querySelectorAll('button,[role="button"],a,span,div')]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent) === confirmText)
                  .map(el => el.closest('button,[role="button"],a') || el)
                  .filter(el => {
                    if (!el || seen.has(el)) return false;
                    seen.add(el);
                    return true;
                  })
                  .filter(visible)
                  .filter(el => !isDisabled(el))
                  .filter(el => {
                    let node = el;
                    for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                      const text = norm(node.innerText || node.textContent);
                      if (text.includes(areaText)
                        && text.includes('确认将')
                        && text.includes('期望到货区域')
                        && /确认将\\s*\\d+\\s*个/.test(text)) {
                        return true;
                      }
                    }
                    return false;
                  })
                  .sort((a, b) => area(a) - area(b))[0] || null;
              }
              function modalRank(el) {
                const cls = String(el.className || '');
                if (el.getAttribute('role') === 'dialog' || el.getAttribute('aria-modal') === 'true') return 3;
                if (/MDL_|modal|Modal|alert|Alert/i.test(cls)) return 2;
                return 0;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_YIWU, TEXT_CONFIRM],
        )
        if not rect:
            raise RuntimeError("没有找到期望到货区域二次确认弹窗中的“确认”按钮。")
        page.mouse.click(rect["x"], rect["y"])
        self._wait_for_page_ready(page, timeout=15000)
        page.wait_for_timeout(1500)
        if self._click_global_button_if_visible(page, TEXT_I_KNOW, timeout=8000):
            self.log("已点击“我知道了”。")

    def _read_stock_sale_manage_signals(self, page) -> dict:
        try:
            return dict(
                page.evaluate(
                    """
                    ([settingText, goodsInfoText]) => {
                      const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
                      const text = norm(document.body.innerText || '');
                      const actions = [...document.querySelectorAll('button,a,[role="button"],span')]
                        .filter(visible)
                        .map(el => norm(el.innerText || el.textContent))
                        .filter(Boolean);
                      return {
                        url: location.href,
                        urlLooksRight: location.href.includes('/stock/fully-mgt/sale-manage/main'),
                        bodyLength: text.length,
                        hasSettingText: text.includes(settingText),
                        hasGoodsInfoText: text.includes(goodsInfoText),
                        checkboxCount: document.querySelectorAll('[data-testid="beast-core-checkbox"],input[type="checkbox"],[role="checkbox"]').length,
                        actionTexts: actions.slice(0, 20),
                      };

                      function visible(el) {
                        const rect = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        return rect.width > 0 && rect.height > 0
                          && style.display !== 'none'
                          && style.visibility !== 'hidden';
                      }
                    }
                    """,
                    [TEXT_EXPECTED_ARRIVAL_AREA_SETTING, TEXT_GOODS_INFO],
                )
            )
        except Exception as exc:  # noqa: BLE001
            return {"url": page.url, "error": str(exc)}

    def _click_visible_text_as_button(self, page, text: str) -> bool:
        rect = page.evaluate(
            """
            ([targetText]) => {
              const target = [...document.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === targetText)
                .map(el => el.closest('button,[role="button"],a') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .sort((a, b) => area(a) - area(b))[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [text],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        return True

    def _read_jit_product_select_signals(self, page) -> dict:
        try:
            return dict(
                page.evaluate(
                    """
                    ([titleText, adjustText, goodsInfoText]) => {
                      const norm = value => String(value || '').replace(/\\s+/g, ' ').trim();
                      const text = norm(document.body.innerText || '');
                      const actions = [...document.querySelectorAll('button,a,[role="button"],span')]
                        .filter(visible)
                        .map(el => norm(el.innerText || el.textContent))
                        .filter(Boolean);
                      return {
                        url: location.href,
                        urlLooksRight: location.href.includes('/newon/product-select'),
                        bodyLength: text.length,
                        hasTitleText: text.includes(titleText),
                        hasAdjustJitText: text.includes(adjustText),
                        hasGoodsInfoText: text.includes(goodsInfoText),
                        checkboxCount: document.querySelectorAll('[data-testid="beast-core-checkbox"],input[type="checkbox"],[role="checkbox"]').length,
                        actionTexts: actions.slice(0, 20),
                      };

                      function visible(el) {
                        const rect = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        return rect.width > 0 && rect.height > 0
                          && style.display !== 'none'
                          && style.visibility !== 'hidden';
                      }
                    }
                    """,
                    [TEXT_NEWON_PRODUCT_SELECT, TEXT_BATCH_ADJUST_JIT, TEXT_GOODS_INFO],
                )
            )
        except Exception as exc:  # noqa: BLE001
            return {"url": page.url, "error": str(exc)}

    def _click_batch_upload_compliance(self, page) -> None:
        if self._has_compliance_upload_modal(page):
            self.log("检测到已有批量上传合规信息弹窗，先关闭后重新打开。")
            self._close_existing_compliance_upload_modal(page)
        self.log("正在打开批量上传合规信息弹窗。")
        for _attempt in range(3):
            if not self._click_button_by_text(page, TEXT_BATCH_UPLOAD_COMPLIANCE):
                raise RuntimeError("没有找到“批量上传合规信息”按钮。")
            try:
                self._wait_for_compliance_drawer(page, timeout=10000)
                page.wait_for_timeout(1200)
                return
            except Exception:  # noqa: BLE001 - retry click when drawer animation did not start.
                page.wait_for_timeout(800)
        raise RuntimeError("已点击“批量上传合规信息”，但没有检测到右侧上传抽屉。")

    def _close_existing_compliance_upload_modal(self, page) -> None:
        for _attempt in range(3):
            rect = page.evaluate(
                """
                ([modalTitle, cancelText, containerSelector]) => {
                  const drawer = [...document.querySelectorAll(containerSelector)]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent).includes(modalTitle))
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!drawer) return { closed: true };
                  const cancel = [...drawer.querySelectorAll('button,[role="button"],a,span')]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent) === cancelText)
                    .map(el => el.closest('button,[role="button"],a') || el)
                    .filter(visible)
                    .filter(el => !isDisabled(el))
                    .sort((a, b) => area(a) - area(b))[0];
                  const close = cancel || [...drawer.querySelectorAll('[aria-label*="关闭"],[aria-label*="Close"],.rocket-drawer-close,.rocket-modal-close,button,span')]
                    .filter(visible)
                    .filter(el => !isDisabled(el))
                    .filter(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text === '' || text === '×';
                    })
                    .sort((a, b) => b.getBoundingClientRect().top - a.getBoundingClientRect().top || b.getBoundingClientRect().right - a.getBoundingClientRect().right)[0];
                  if (!close) return null;
                  close.scrollIntoView({ block: 'center', inline: 'center' });
                  const rect = close.getBoundingClientRect();
                  return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isDisabled(el) {
                    return el.disabled === true
                      || el.getAttribute('disabled') !== null
                      || el.getAttribute('aria-disabled') === 'true';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, TEXT_CANCEL, compliance_modal_container_selector()],
            )
            if isinstance(rect, dict) and rect.get("closed"):
                return
            if rect:
                page.mouse.click(rect["x"], rect["y"])
                page.wait_for_timeout(800)
                if self._click_global_button_if_visible(page, TEXT_CONFIRM, timeout=1500):
                    page.wait_for_timeout(800)
                if not self._has_compliance_upload_modal(page):
                    return
            page.keyboard.press("Escape")
            page.wait_for_timeout(800)
            if not self._has_compliance_upload_modal(page):
                return
        raise RuntimeError("检测到已有批量上传合规信息弹窗，但未能关闭。")

    def _open_identifier_upload_dialog(self, page) -> None:
        self.log("正在打开批量上传商品识别码弹窗。")
        for attempt in range(1, 6):
            if self._has_identifier_upload_dialog(page):
                page.wait_for_timeout(800)
                return
            if not self._click_button_by_text(page, TEXT_BATCH_UPLOAD_IDENTIFIER):
                raise RuntimeError("没有找到或无法点击“批量上传商品识别码”按钮。")
            page.wait_for_timeout(500)
            if not self._wait_for_popup_item(page, TEXT_UPLOAD_FILE, exact=True, timeout=4000):
                self.log(f"第 {attempt} 次没有等到“上传文件”菜单项，准备重试。")
                page.keyboard.press("Escape")
                page.wait_for_timeout(900)
                continue
            if not self._activate_popup_item(page, TEXT_UPLOAD_FILE, exact=True):
                self.log(f"第 {attempt} 次检测到“上传文件”菜单项但点击失败，准备重试。")
                page.keyboard.press("Escape")
                page.wait_for_timeout(900)
                continue
            try:
                self._wait_for_identifier_upload_dialog(page, timeout=12000)
                page.wait_for_timeout(800)
                return
            except Exception:  # noqa: BLE001 - upload dialog may still be rendering; retry from trigger.
                self.log(f"第 {attempt} 次点击“上传文件”后没有检测到上传弹窗，准备重试。")
                page.wait_for_timeout(900)
        raise RuntimeError("多次尝试后仍没有打开“批量上传商品识别码 > 上传文件”弹窗。")

    def _has_identifier_upload_dialog(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([titleText, uploadText]) => {
                  return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.rocket-modal,div')]
                    .filter(visible)
                    .some(el => norm(el.innerText || el.textContent).includes(titleText)
                      && norm(el.innerText || el.textContent).includes(uploadText));

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                [TEXT_IDENTIFIER_UPLOAD_MODAL_TITLE, TEXT_UPLOAD_FILE],
            )
        )

    def _wait_for_identifier_upload_dialog(self, page, timeout: int = 10000) -> None:
        page.wait_for_function(
            """
            ([titleText, uploadText]) => {
              return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.rocket-modal,div')]
                .filter(visible)
                .some(el => norm(el.innerText || el.textContent).includes(titleText)
                  && norm(el.innerText || el.textContent).includes(uploadText));

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
            }
            """,
            arg=[TEXT_IDENTIFIER_UPLOAD_MODAL_TITLE, TEXT_UPLOAD_FILE],
            timeout=timeout,
        )

    def _upload_identifier_file_to_dialog(self, page, file_path: Path) -> None:
        if not file_path.exists():
            raise RuntimeError(f"上传文件不存在：{file_path}")
        self.log(f"正在上传文件：{file_path}")
        file_input = page.locator("input[type='file']").last
        file_input.set_input_files(str(file_path), timeout=10000)
        page.wait_for_timeout(1200)

    def _wait_for_identifier_file_parsed(self, page, timeout: int = 60000) -> None:
        page.wait_for_function(
            """
            ([startImportText]) => {
              const button = [...document.querySelectorAll('button,[role="button"]')]
                .filter(visible)
                .find(el => compact(el.innerText || el.textContent) === compact(startImportText));
              return Boolean(button && !isDisabled(button));

              function compact(value) {
                return String(value || '').replace(/\\s+/g, '').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true';
              }
            }
            """,
            arg=[TEXT_START_IMPORT],
            timeout=timeout,
        )
        self.log("商品识别码文件解析完成，已检测到“开始导入”。")

    def _submit_identifier_import(self, page) -> None:
        self._click_global_button(page, TEXT_START_IMPORT)
        self._wait_for_page_ready(page, timeout=12000)
        page.wait_for_timeout(1200)
        if self._click_global_button_if_visible(page, TEXT_IMPORT_CONFIRMED, timeout=12000):
            self.log("已点击“我已确认数据无误，确认导入”。")
        elif self._click_global_button_if_visible(page, TEXT_CONFIRM, timeout=5000):
            self.log("已点击导入确认。")
        self._wait_for_identifier_import_finished(page)
        if self._click_global_button_if_visible(page, TEXT_I_KNOW, timeout=12000):
            self.log("已点击“我知道了”。")
        self._close_identifier_upload_dialog(page)

    def _wait_for_identifier_import_finished(self, page, timeout: int = 60000) -> None:
        try:
            page.wait_for_function(
                """
                () => {
                  const text = String(document.body.innerText || '');
                  if (text.includes('导入完成')) return true;
                  if (text.includes('导入成功')) return true;
                  return false;
                }
                """,
                timeout=timeout,
            )
        except Exception:  # noqa: BLE001 - Some imports close quickly or only show a toast.
            self.log("未明确检测到导入完成文案，继续尝试关闭上传弹窗。")
        page.wait_for_timeout(1200)

    def _close_identifier_upload_dialog(self, page) -> None:
        if self._click_identifier_dialog_close(page):
            self.log("已关闭上传文件弹窗。")
            return
        if self._click_global_button_if_visible(page, "关闭", timeout=3000):
            self.log("已关闭上传文件弹窗。")

    def _click_identifier_dialog_close(self, page) -> bool:
        rect = page.evaluate(
            """
            ([titleText]) => {
              const dialogs = [...document.querySelectorAll('[role="dialog"],.rocket-modal,.rocket-modal-wrap,.rocket-modal-content')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent).includes(titleText))
                .sort((a, b) => area(a) - area(b));
              const dialog = dialogs[0];
              if (!dialog) return null;
              const close = [...dialog.querySelectorAll('button,[aria-label*="close"],[class*="close"],[class*="Close"]')]
                .filter(visible)
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return br.top - ar.top || br.left - ar.left;
                })[0];
              if (!close) return null;
              const rect = close.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_IDENTIFIER_UPLOAD_MODAL_TITLE],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(700)
        return True

    def _open_inventory_setting_upload_dialog(self, page) -> None:
        self.log("正在打开批量导入库存设置上传弹窗。")
        base_file_input_count = self._file_input_count(page)
        for attempt in range(1, 6):
            if not self._click_goods_toolbar_more(page):
                raise RuntimeError("没有找到或无法点击“更多”按钮。")
            page.wait_for_timeout(500)
            if not self._activate_popup_item(page, TEXT_BATCH_IMPORT_INVENTORY, exact=False):
                self.log(f"第 {attempt} 次没有找到“批量导入库存”菜单项，准备重试。")
                page.wait_for_timeout(800)
                continue
            page.wait_for_timeout(500)
            if not self._wait_for_popup_item(page, TEXT_BATCH_IMPORT_INVENTORY_SETTING, exact=True, timeout=4000):
                self.log(f"第 {attempt} 次没有等到“批量导入-库存设置”菜单项，准备重试。")
                page.wait_for_timeout(800)
                continue
            if not self._activate_popup_item(page, TEXT_BATCH_IMPORT_INVENTORY_SETTING, exact=True):
                self.log(f"第 {attempt} 次检测到“批量导入-库存设置”但点击失败，准备重试。")
                page.wait_for_timeout(800)
                continue
            try:
                self._wait_for_upload_file_input(page, minimum_count=base_file_input_count + 1, timeout=12000)
                page.wait_for_timeout(800)
                return
            except Exception:  # noqa: BLE001 - upload dialog can render slowly.
                self.log(f"第 {attempt} 次点击库存设置上传后没有检测到文件上传控件，准备重试。")
                page.wait_for_timeout(900)
        raise RuntimeError("多次尝试后仍没有打开“批量导入-库存设置”上传弹窗。")

    def _file_input_count(self, page) -> int:
        return int(page.locator("input[type='file']").count())

    def _wait_for_upload_file_input(self, page, minimum_count: int = 1, timeout: int = 10000) -> None:
        page.wait_for_function(
            """
            ([minimumCount]) => document.querySelectorAll('input[type="file"]').length >= minimumCount
            """,
            arg=[minimum_count],
            timeout=timeout,
        )

    def _wait_for_inventory_setting_edit_modal(self, page, timeout: int = 30000) -> None:
        page.wait_for_function(
            """
            ([saveText, cancelText]) => {
              const bodyText = norm(document.body.innerText || '');
              const recognized = bodyText.includes('修改库存') && bodyText.includes('SKU ID') && bodyText.includes('修改后库存');
              const save = [...document.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === saveText)
                .map(el => el.closest('button,[role="button"],a') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))[0];
              const cancel = [...document.querySelectorAll('button,[role="button"],a,span,div')]
                .filter(visible)
                .some(el => norm(el.innerText || el.textContent) === cancelText);
              return Boolean(recognized && save && cancel);

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true';
              }
            }
            """,
            arg=["保存", TEXT_CANCEL],
            timeout=timeout,
        )
        self.log("库存设置文件识别完成，已打开“修改库存”弹窗，准备点击“保存”。")

    def _save_inventory_setting_import(self, page) -> None:
        if not self._click_global_button_if_visible(page, "保存", timeout=10000):
            raise RuntimeError("没有找到库存设置弹窗中的“保存”按钮。")
        self._wait_for_page_ready(page, timeout=10000)
        self._wait_for_inventory_setting_save_finished(page)
        self._close_identifier_upload_dialog(page)

    def _wait_for_inventory_setting_save_finished(self, page, timeout: int = 30000) -> None:
        try:
            page.wait_for_function(
                """
                () => {
                  const bodyText = norm(document.body.innerText || '');
                  if (bodyText.includes('保存成功')) return true;
                  if (bodyText.includes('上传成功')) return true;
                  if (bodyText.includes('导入成功')) return true;
                  if (bodyText.includes('修改成功')) return true;
                  return !hasInventorySettingEditModal();

                  function hasInventorySettingEditModal() {
                    return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-dialog,.rocket-modal-content')]
                      .filter(visible)
                      .some(el => {
                        const text = norm(el.innerText || el.textContent);
                        return text.includes('修改库存')
                          && text.includes('SKU ID')
                          && text.includes('修改后库存');
                      });
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                timeout=timeout,
            )
            self.log("库存设置保存完成，已检测到完成提示或弹窗关闭。")
        except Exception:  # noqa: BLE001 - avoid leaving the worker stuck on an uncertain toast.
            self.log("未明确检测到库存设置保存完成提示，继续执行弹窗清理。")
        page.wait_for_timeout(800)

    def _has_inventory_setting_edit_modal(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                () => {
                  return [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-dialog,.rocket-modal-content')]
                    .filter(visible)
                    .some(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text.includes('修改库存')
                        && text.includes('SKU ID')
                        && text.includes('修改后库存');
                    });

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """
            )
        )

    def _has_compliance_upload_modal(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([modalTitle, containerSelector]) => {
                  return [...document.querySelectorAll(containerSelector)]
                    .filter(visible)
                    .some(el => norm(el.innerText || el.textContent).includes(modalTitle)
                      && norm(el.innerText || el.textContent).includes('选择合规信息类型'));

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
            )
        )

    def _wait_for_compliance_drawer(self, page, timeout: int = 10000) -> None:
        page.wait_for_function(
            """
            ([modalTitle, containerSelector]) => {
              return [...document.querySelectorAll(containerSelector)]
                .filter(visible)
                .some(el => norm(el.innerText || el.textContent).includes(modalTitle)
                  && norm(el.innerText || el.textContent).includes('选择合规信息类型'));

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
            }
            """,
            arg=[TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
            timeout=timeout,
        )

    def _select_compliance_modal_dropdown(
        self,
        page,
        label_text: str,
        option_text: str,
        option_fallback: str | None = None,
    ) -> None:
        if self._select_compliance_dropdown_in_drawer(page, label_text, option_text, option_fallback):
            return

        rect = None
        for _attempt in range(6):
            rect = page.evaluate(
                """
                ([modalTitle, labelText]) => {
                  const modal = findModal(modalTitle);
                  if (!modal) return null;
                  const target = findFieldControl(modal, labelText);
                  if (!target) return null;
                  target.scrollIntoView({ block: 'center', inline: 'center' });
                  const rect = target.getBoundingClientRect();
                  return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

                  function findFieldControl(scope, label) {
                    const rows = [...scope.querySelectorAll('label,div,section,td,tr')]
                      .filter(visible)
                      .filter(el => norm(el.innerText || el.textContent).includes(label))
                      .filter(el => el.querySelector('input,[role="combobox"],[class*="select"],[class*="Select"]'))
                      .sort((a, b) => area(a) - area(b));
                    for (const row of rows) {
                      const control = bestControl(row, label);
                      if (control) return control;
                    }

                    const labels = [...scope.querySelectorAll('label,span,div')]
                      .filter(visible)
                      .filter(el => norm(el.innerText || el.textContent).includes(label))
                      .sort((a, b) => area(a) - area(b));
                    for (const labelEl of labels) {
                      const labelRect = labelEl.getBoundingClientRect();
                      const control = [...scope.querySelectorAll('input,[role="combobox"],[class*="select"],[class*="Select"],[class*="selector"],[class*="Selector"],div')]
                        .filter(visible)
                        .filter(el => {
                          const rect = el.getBoundingClientRect();
                          if (rect.width < 80 || rect.height < 20) return false;
                          if (norm(el.innerText || el.textContent) === label) return false;
                          const sameLine = rect.top < labelRect.bottom + 24 && rect.bottom > labelRect.top - 24;
                          const toRight = rect.left > labelRect.left + 20;
                          return sameLine && toRight;
                        })
                        .sort((a, b) => {
                          const ar = a.getBoundingClientRect();
                          const br = b.getBoundingClientRect();
                          return area(b) - area(a) || ar.left - br.left;
                        })[0];
                      if (control) return control;
                    }
                    return null;
                  }
                  function bestControl(row, label) {
                    const controls = [...row.querySelectorAll('input,[role="combobox"],[class*="select"],[class*="Select"],[class*="selector"],[class*="Selector"]')]
                      .filter(visible)
                      .filter(el => {
                        const rect = el.getBoundingClientRect();
                        if (rect.width < 80 || rect.height < 20) return false;
                        const text = norm(el.innerText || el.textContent || el.value || el.getAttribute('placeholder'));
                        return text !== label;
                      })
                      .sort((a, b) => area(b) - area(a));
                    return controls[0] || null;
                  }
                  function findModal(title) {
                    const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                      .filter(visible)
                      .filter(el => isModalContainer(el, title))
                      .sort((a, b) => area(a) - area(b));
                    return candidates[0] || null;
                  }
                  function isModalContainer(el, title) {
                    const text = String(el.innerText || el.textContent || '');
                    if (!text.includes(title)) return false;
                    if (!el.querySelector('input,button,[role="button"],table')) return false;
                    const style = getComputedStyle(el);
                    const cls = String(el.className || '');
                    const role = el.getAttribute?.('role') || '';
                    const modalLike = role === 'dialog'
                      || el.getAttribute('aria-modal') === 'true'
                      || /modal|dialog|drawer|panel|portal/i.test(cls)
                      || style.position === 'fixed'
                      || style.position === 'absolute'
                      || Number(style.zIndex || 0) >= 100;
                    const rect = el.getBoundingClientRect();
                    const notMainPage = rect.width < window.innerWidth * 0.96 || rect.left > window.innerWidth * 0.2;
                    return area(el) > 10000 && modalLike && notMainPage;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, label_text],
            )
            if rect:
                break
            page.wait_for_timeout(1000)
        if not rect:
            raise RuntimeError(f"没有找到“{label_text}”下拉框。")
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(500)
        selected = (
            self._activate_compliance_option_near_dropdown(page, label_text, option_text, exact=False)
            or
            self._activate_popup_item(page, option_text, exact=False)
        )
        if not selected:
            page.keyboard.press("Control+A")
            page.keyboard.type(option_text, delay=20)
            page.wait_for_timeout(700)
            selected = (
                self._activate_compliance_option_near_dropdown(page, label_text, option_text, exact=False)
                or
                self._activate_popup_item(page, option_text, exact=False)
            )
        if not selected and option_fallback:
            page.keyboard.press("Control+A")
            page.keyboard.type(option_fallback, delay=20)
            page.wait_for_timeout(700)
            selected = (
                self._activate_compliance_option_near_dropdown(page, label_text, option_fallback, exact=False)
                or
                self._activate_popup_item(page, option_fallback, exact=False)
            )
        if not selected:
            raise RuntimeError(f"没有找到“{option_text}”选项。")
        page.wait_for_timeout(700)
        if not self._compliance_dropdown_has_value(page, label_text, option_text, option_fallback):
            raise RuntimeError(f"已尝试选择“{option_text}”，但“{label_text}”仍未显示为已选择。")

    def _select_compliance_dropdown_in_drawer(
        self,
        page,
        label_text: str,
        option_text: str,
        option_fallback: str | None = None,
    ) -> bool:
        try:
            drawer = self._compliance_drawer(page)
            input_id = ""
            for _attempt in range(40):
                input_id = self._read_compliance_label_for(page, label_text)
                if input_id:
                    break
                page.wait_for_timeout(500)
            if not input_id:
                self.log(f"抽屉内暂未出现“{label_text}”字段。")
                return False

            rect = self._compliance_select_rect(page, input_id)
            if not rect:
                return False

            for _attempt in range(6):
                selected = self._click_compliance_dropdown_option(page, option_text)
                if not selected and option_fallback:
                    selected = self._click_compliance_dropdown_option(page, option_fallback)
                if selected:
                    page.wait_for_timeout(700)
                    if self._compliance_drawer_dropdown_has_value(page, label_text, option_text, option_fallback):
                        return True

                page.mouse.click(rect["x"], rect["y"])
                page.wait_for_timeout(500)
                selected = self._click_compliance_dropdown_option(page, option_text)
                if not selected and option_fallback:
                    selected = self._click_compliance_dropdown_option(page, option_fallback)
                if not selected:
                    for search_text in self._compliance_dropdown_search_terms(option_text, option_fallback):
                        self._search_compliance_dropdown(page, rect, search_text)
                        selected = self._click_compliance_dropdown_option(page, option_text)
                        if not selected and option_fallback:
                            selected = self._click_compliance_dropdown_option(page, option_fallback)
                        if selected:
                            break
                if selected:
                    page.wait_for_timeout(700)
                    if self._compliance_drawer_dropdown_has_value(page, label_text, option_text, option_fallback):
                        return True
            return False
        except Exception as exc:  # noqa: BLE001 - old coordinate path remains as fallback.
            self.log(f"抽屉内选择“{label_text}”未成功，改用备用选择方式：{exc}")
            return False

    def _compliance_dropdown_search_terms(
        self,
        option_text: str,
        option_fallback: str | None = None,
    ) -> list[str]:
        terms = [option_text]
        shortened = re.sub(r"[（(].*?[）)]", "", option_text).strip()
        if shortened and shortened not in terms:
            terms.append(shortened)
        if option_fallback and option_fallback not in terms:
            terms.append(option_fallback)
        return terms

    def _search_compliance_dropdown(self, page, rect: dict, search_text: str) -> None:
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(200)
        page.keyboard.press("Control+A")
        page.keyboard.type(search_text, delay=20)
        page.wait_for_timeout(700)

    def _compliance_drawer(self, page):
        page.wait_for_function(
            """
            ([modalTitle, containerSelector]) => {
              return [...document.querySelectorAll(containerSelector)]
                .filter(visible)
                .some(el => norm(el.innerText || el.textContent).includes(modalTitle));

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
            }
            """,
            arg=[TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
            timeout=10000,
        )
        return page.locator(compliance_modal_container_selector()).filter(has_text=TEXT_COMPLIANCE_MODAL_TITLE).last

    def _compliance_select_rect(self, page, input_id: str) -> dict | None:
        return page.evaluate(
            """
            ([modalTitle, inputId, containerSelector]) => {
              const drawer = [...document.querySelectorAll(containerSelector)]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent).includes(modalTitle))
                .filter(el => el.querySelector(`input#${CSS.escape(inputId)}`))
                .sort((a, b) => area(a) - area(b))[0];
              if (!drawer) return null;
              const input = [...drawer.querySelectorAll('input')]
                .filter(visible)
                .filter(el => el.id === inputId)
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                })[0];
              const select = input?.closest('.rocket-select');
              if (!select || !visible(select)) return null;
              select.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = select.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, input_id, compliance_modal_container_selector()],
        )

    def _read_compliance_label_for(self, page, label_text: str) -> str:
        return str(
            page.evaluate(
                """
                ([modalTitle, labelText, containerSelector]) => {
                  const drawer = [...document.querySelectorAll(containerSelector)]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent).includes(modalTitle))
                    .filter(el => {
                      const target = compact(labelText);
                      return [...el.querySelectorAll('label')]
                        .filter(visible)
                        .some(label => compact(label.innerText || label.textContent) === target);
                    })
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!drawer) return '';
                  const target = compact(labelText);
                  const labels = [...drawer.querySelectorAll('label')]
                    .filter(visible)
                    .filter(el => compact(el.innerText || el.textContent) === target);
                  return labels[0]?.getAttribute('for') || '';

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, label_text, compliance_modal_container_selector()],
            )
            or ""
        )

    def _click_compliance_dropdown_option(self, page, option_text: str) -> bool:
        rect = page.evaluate(
            """
            ([targetText]) => {
              const target = compact(targetText);
              const dropdowns = [...document.querySelectorAll('.rocket-select-dropdown')]
                .filter(visible)
                .filter(el => !String(el.className || '').includes('rocket-select-dropdown-hidden'));
              const options = dropdowns.flatMap(dropdown =>
                [...dropdown.querySelectorAll('.rocket-select-item-option,[role="option"]')]
                  .filter(visible)
                  .filter(el => compact(el.innerText || el.textContent) === target)
              );
              const option = options.sort((a, b) => {
                const ar = a.getBoundingClientRect();
                const br = b.getBoundingClientRect();
                return ar.top - br.top || area(a) - area(b);
              })[0];
              if (!option) return null;
              option.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = option.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [option_text],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(300)
        return True

    def _activate_compliance_option_near_dropdown(
        self,
        page,
        label_text: str,
        option_text: str,
        exact: bool = False,
    ) -> bool:
        result = page.evaluate(
            """
            ([modalTitle, labelText, optionText, exact]) => {
              const modal = findModal(modalTitle);
              if (!modal) return false;
              const field = findFieldControl(modal, labelText);
              if (!field) return false;
              const fieldRect = field.getBoundingClientRect();
              const targetCompact = compact(optionText);
              const candidates = [...document.querySelectorAll('body *')]
                .filter(visible)
                .filter(el => {
                  const text = norm(el.innerText || el.textContent);
                  if (!text) return false;
                  const matched = exact ? compact(text) === targetCompact : compact(text).includes(targetCompact);
                  if (!matched) return false;
                  const rect = el.getBoundingClientRect();
                  const nearHorizontal = rect.right >= fieldRect.left - 80 && rect.left <= fieldRect.right + 80;
                  const belowOrOverlay = rect.top >= fieldRect.top - 8 && rect.top <= fieldRect.bottom + 420;
                  const smallEnough = area(el) < Math.max(60000, fieldRect.width * 260);
                  return nearHorizontal && belowOrOverlay && smallEnough;
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  const aPopup = isPopupRelated(a) ? 0 : 1;
                  const bPopup = isPopupRelated(b) ? 0 : 1;
                  if (aPopup !== bPopup) return aPopup - bPopup;
                  return Math.abs(ar.top - fieldRect.bottom) - Math.abs(br.top - fieldRect.bottom)
                    || area(a) - area(b);
                });
              const item = candidates[0];
              if (!item) return false;
              const target = item.closest('[role="option"],[role="menuitem"],li,button') || item;
              fireMouse(target);
              target.click?.();
              return true;

              function findFieldControl(scope, label) {
                const labels = [...scope.querySelectorAll('label,span,div')]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent).includes(label))
                  .sort((a, b) => area(a) - area(b));
                for (const labelEl of labels) {
                  const labelRect = labelEl.getBoundingClientRect();
                  const input = [...scope.querySelectorAll('input,[role="combobox"]')]
                    .filter(visible)
                    .filter(el => {
                      const rect = el.getBoundingClientRect();
                      if (rect.width < 80 || rect.height < 18) return false;
                      const sameLine = rect.top < labelRect.bottom + 26 && rect.bottom > labelRect.top - 26;
                      const toRight = rect.left > labelRect.left + 20;
                      return sameLine && toRight;
                    })
                    .sort((a, b) => {
                      const ar = a.getBoundingClientRect();
                      const br = b.getBoundingClientRect();
                      return ar.left - br.left;
                    })[0];
                  if (input) return input;
                }
                return null;
              }
              function findModal(title) {
                const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                  .filter(visible)
                  .filter(el => isModalContainer(el, title))
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function isModalContainer(el, title) {
                const text = String(el.innerText || el.textContent || '');
                if (!text.includes(title)) return false;
                if (!el.querySelector('input,button,[role="button"],table')) return false;
                const style = getComputedStyle(el);
                const cls = String(el.className || '');
                const role = el.getAttribute?.('role') || '';
                const modalLike = role === 'dialog'
                  || el.getAttribute('aria-modal') === 'true'
                  || /modal|dialog|drawer|panel|portal/i.test(cls)
                  || style.position === 'fixed'
                  || style.position === 'absolute'
                  || Number(style.zIndex || 0) >= 100;
                const rect = el.getBoundingClientRect();
                const notMainPage = rect.width < window.innerWidth * 0.96 || rect.left > window.innerWidth * 0.2;
                return area(el) > 10000 && modalLike && notMainPage;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isPopupRelated(el) {
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  const style = getComputedStyle(node);
                  const cls = String(node.className || '');
                  const role = node.getAttribute?.('role') || '';
                  if (style.position === 'fixed' || style.position === 'absolute') return true;
                  if (/dropdown|popover|popper|portal|menu|select|option/i.test(cls)) return true;
                  if (/menu|menuitem|listbox|option/.test(role)) return true;
                }
                return false;
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
              function fireMouse(el) {
                el.scrollIntoView({ block: 'center', inline: 'center' });
                const rect = el.getBoundingClientRect();
                const clientX = rect.left + rect.width / 2;
                const clientY = rect.top + rect.height / 2;
                for (const type of ['pointerover', 'mouseover', 'pointermove', 'mousemove', 'pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
                  el.dispatchEvent(new MouseEvent(type, {
                    bubbles: true,
                    cancelable: true,
                    view: window,
                    clientX,
                    clientY,
                  }));
                }
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, label_text, option_text, exact],
        )
        return bool(result)

    def _compliance_dropdown_has_value(
        self,
        page,
        label_text: str,
        option_text: str,
        option_fallback: str | None = None,
    ) -> bool:
        try:
            if self._compliance_drawer_dropdown_has_value(page, label_text, option_text, option_fallback):
                return True
        except Exception:  # noqa: BLE001 - keep the older DOM fallback below.
            pass

        return bool(
            page.evaluate(
                """
                ([modalTitle, labelText, optionText, fallbackText]) => {
                  const modal = findModal(modalTitle);
                  if (!modal) return false;
                  const control = findFieldControl(modal, labelText);
                  if (!control) return false;
                  const container = control.closest('label,div,section,td,tr') || control;
                  const text = compact(
                    (control.innerText || control.textContent || control.value || '')
                    + ' '
                    + (container.innerText || container.textContent || '')
                  );
                  const option = compact(optionText);
                  const fallback = compact(fallbackText || '');
                  return text.includes(option) || (fallback && text.includes(fallback));

                  function findFieldControl(scope, label) {
                    const rows = [...scope.querySelectorAll('label,div,section,td,tr')]
                      .filter(visible)
                      .filter(el => norm(el.innerText || el.textContent).includes(label))
                      .filter(el => el.querySelector('input,[role="combobox"],[class*="select"],[class*="Select"]'))
                      .sort((a, b) => area(a) - area(b));
                    for (const row of rows) {
                      const control = bestControl(row, label);
                      if (control) return control;
                    }
                    return null;
                  }
                  function bestControl(row, label) {
                    const controls = [...row.querySelectorAll('input,[role="combobox"],[class*="select"],[class*="Select"],[class*="selector"],[class*="Selector"],div')]
                      .filter(visible)
                      .filter(el => {
                        const rect = el.getBoundingClientRect();
                        if (rect.width < 80 || rect.height < 20) return false;
                        const text = norm(el.innerText || el.textContent || el.value || el.getAttribute('placeholder'));
                        return text !== label;
                      })
                      .sort((a, b) => area(b) - area(a));
                    return controls[0] || null;
                  }
                  function findModal(title) {
                    const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                      .filter(visible)
                      .filter(el => isModalContainer(el, title))
                      .sort((a, b) => area(a) - area(b));
                    return candidates[0] || null;
                  }
                  function isModalContainer(el, title) {
                    const text = String(el.innerText || el.textContent || '');
                    if (!text.includes(title)) return false;
                    if (!el.querySelector('input,button,[role="button"],table')) return false;
                    const style = getComputedStyle(el);
                    const cls = String(el.className || '');
                    const role = el.getAttribute?.('role') || '';
                    const modalLike = role === 'dialog'
                      || el.getAttribute('aria-modal') === 'true'
                      || /modal|dialog|drawer|panel|portal/i.test(cls)
                      || style.position === 'fixed'
                      || style.position === 'absolute'
                      || Number(style.zIndex || 0) >= 100;
                    const rect = el.getBoundingClientRect();
                    const notMainPage = rect.width < window.innerWidth * 0.96 || rect.left > window.innerWidth * 0.2;
                    return area(el) > 10000 && modalLike && notMainPage;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, label_text, option_text, option_fallback or ""],
            )
        )

    def _compliance_drawer_dropdown_has_value(
        self,
        page,
        label_text: str,
        option_text: str,
        option_fallback: str | None = None,
    ) -> bool:
        input_id = self._read_compliance_label_for(page, label_text)
        if not input_id:
            return False
        return bool(
            page.evaluate(
                """
                ([modalTitle, inputId, optionText, fallbackText, containerSelector]) => {
                  const drawer = [...document.querySelectorAll(containerSelector)]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent).includes(modalTitle))
                    .filter(el => el.querySelector(`input#${CSS.escape(inputId)}`))
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!drawer) return false;
                  const input = [...drawer.querySelectorAll('input')]
                    .filter(visible)
                    .filter(el => el.id === inputId)[0];
                  const select = input?.closest('.rocket-select');
                  if (!select) return false;
                  const text = compact(select.innerText || select.textContent || '');
                  const option = compact(optionText);
                  const fallback = compact(fallbackText || '');
                  return text.includes(option) || Boolean(fallback && text.includes(fallback));

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [
                    TEXT_COMPLIANCE_MODAL_TITLE,
                    input_id,
                    option_text,
                    option_fallback or "",
                    compliance_modal_container_selector(),
                ],
            )
        )

    def _click_compliance_modal_button(self, page, text: str) -> None:
        if self._click_compliance_drawer_button(page, text):
            page.wait_for_timeout(700)
            return
        if text == TEXT_SEARCH and self._click_compliance_drawer_button(page, "查询"):
            page.wait_for_timeout(700)
            return

        rect = page.evaluate(
            """
            ([modalTitle, targetText]) => {
              const modal = findModal(modalTitle);
              if (!modal) return null;
              const candidates = [...modal.querySelectorAll('button,a,[role="button"],span,div')]
                .filter(visible)
                .filter(el => norm(el.innerText || el.textContent) === targetText)
                .map(el => el.closest('button,a,[role="button"]') || el)
                .filter(visible)
                .filter(el => !isDisabled(el))
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return area(a) - area(b) || br.top - ar.top || ar.left - br.left;
                });
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findModal(title) {
                const candidates = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],div,section,aside')]
                  .filter(visible)
                  .filter(el => isModalContainer(el, title))
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function isModalContainer(el, title) {
                const text = String(el.innerText || el.textContent || '');
                if (!text.includes(title)) return false;
                if (!el.querySelector('input,button,[role="button"],table')) return false;
                const style = getComputedStyle(el);
                const cls = String(el.className || '');
                const role = el.getAttribute?.('role') || '';
                const modalLike = role === 'dialog'
                  || el.getAttribute('aria-modal') === 'true'
                  || /modal|dialog|drawer|panel|portal/i.test(cls)
                  || style.position === 'fixed'
                  || style.position === 'absolute'
                  || Number(style.zIndex || 0) >= 100;
                const rect = el.getBoundingClientRect();
                const notMainPage = rect.width < window.innerWidth * 0.96 || rect.left > window.innerWidth * 0.2;
                return area(el) > 10000 && modalLike && notMainPage;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, text],
        )
        if not rect:
            raise RuntimeError(f"没有找到“{text}”按钮。")
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(700)

    def _click_compliance_drawer_button(self, page, text: str) -> bool:
        try:
            drawer = self._compliance_drawer(page)
            target_compact = re.sub(r"\s+", "", text)
            buttons = drawer.locator("button,[role='button'],a")
            count = buttons.count()
            for index in range(count):
                button = buttons.nth(index)
                try:
                    label = re.sub(r"\s+", "", button.inner_text(timeout=1000))
                    if label == target_compact and button.is_visible() and button.is_enabled():
                        button.click(timeout=5000)
                        return True
                except Exception:  # noqa: BLE001 - continue scanning visible buttons.
                    continue
            return False
        except Exception as exc:  # noqa: BLE001
            self.log(f"抽屉内点击“{text}”未成功，改用备用点击方式：{exc}")
            return False

    def _submit_compliance_upload(self, page) -> None:
        self._click_compliance_modal_button(page, TEXT_CONFIRM_UPLOAD)
        self._wait_for_page_ready(page, timeout=10000)
        page.wait_for_timeout(1200)
        if self._click_global_button_if_visible(page, TEXT_I_KNOW, timeout=12000):
            self.log("已点击“我知道了”。")
        elif self._click_visible_text_as_button(page, TEXT_I_KNOW):
            self.log("已点击“我知道了”。")
        elif self._close_compliance_success_modal(page):
            self.log("已点击“我知道了”。")
        else:
            raise RuntimeError("确认上传后没有关闭上传成功提示，已停止后续合规步骤。")

    def _dismiss_existing_compliance_success_modal(self, page) -> None:
        if not self._has_compliance_success_modal(page):
            return
        self.log("检测到残留的合规上传成功提示，正在关闭。")
        if not self._close_compliance_success_modal(page, timeout=5000):
            raise RuntimeError("检测到残留的合规上传成功提示，但未能关闭，已停止执行以避免误操作。")

    def _close_compliance_success_modal(self, page, timeout: int = 12000) -> bool:
        end_time = page.evaluate("Date.now()") + timeout
        while page.evaluate("Date.now()") < end_time:
            rect = page.evaluate(
                """
                ([iKnowText]) => {
                  const modal = findSuccessModal();
                  if (!modal) return null;
                  const exact = [...modal.querySelectorAll('button,[role="button"],a,span')]
                    .filter(visible)
                    .filter(el => compact(el.innerText || el.textContent) === compact(iKnowText))
                    .map(el => el.closest('button,[role="button"],a') || el)
                    .filter(visible)
                    .filter(el => !isDisabled(el))
                    .sort((a, b) => scoreButton(a) - scoreButton(b) || area(a) - area(b))[0];
                  const primary = exact || [...modal.querySelectorAll('button,[role="button"]')]
                    .filter(visible)
                    .filter(el => !isDisabled(el))
                    .sort((a, b) => scoreButton(a) - scoreButton(b) || b.getBoundingClientRect().top - a.getBoundingClientRect().top)[0];
                  const close = primary || [...modal.querySelectorAll('[aria-label*="关闭"],[aria-label*="Close"],.rocket-modal-close,.rocket-dialog-close,.rocket-icon-close,button,span')]
                    .filter(visible)
                    .filter(el => !isDisabled(el))
                    .filter(el => compact(el.innerText || el.textContent) === '' || compact(el.innerText || el.textContent) === '×')
                    .sort((a, b) => b.getBoundingClientRect().top - a.getBoundingClientRect().top || b.getBoundingClientRect().right - a.getBoundingClientRect().right)[0];
                  if (!close) return null;
                  close.scrollIntoView({ block: 'center', inline: 'center' });
                  const rect = close.getBoundingClientRect();
                  return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

                  function findSuccessModal() {
                    const containers = [
                      ...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-modal-content,.rocket-dialog,.rocket-dialog-content')
                    ]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        if (!text.includes('上传成功') && !text.includes('合规信息已上传成功')) return false;
                        return text.includes(iKnowText)
                          || [...el.querySelectorAll('button,[role="button"],a,[aria-label*="关闭"],[aria-label*="Close"],.rocket-modal-close,.rocket-dialog-close')].some(visible);
                      })
                      .sort((a, b) => area(a) - area(b));
                    return containers[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isDisabled(el) {
                    return el.disabled === true
                      || el.getAttribute('disabled') !== null
                      || el.getAttribute('aria-disabled') === 'true';
                  }
                  function scoreButton(el) {
                    const text = compact(el.innerText || el.textContent);
                    const cls = String(el.className || '');
                    let score = 50;
                    if (text === compact(iKnowText)) score -= 100;
                    if (/primary|btn-primary|button-primary/.test(cls)) score -= 20;
                    return score;
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_I_KNOW],
            )
            if rect:
                page.mouse.click(rect["x"], rect["y"])
                page.wait_for_timeout(800)
                if not self._has_compliance_success_modal(page):
                    return True
            page.wait_for_timeout(500)
        return False

    def _has_compliance_success_modal(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                () => {
                  return Boolean(findSuccessModal());

                  function findSuccessModal() {
                    const containers = [...document.querySelectorAll('[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-modal-content,.rocket-dialog,.rocket-dialog-content')]
                    .filter(visible)
                    .filter(el => {
                      const text = String(el.innerText || el.textContent || '').replace(/\\s+/g, ' ').trim();
                      if (!text.includes('上传成功') && !text.includes('合规信息已上传成功')) return false;
                      return text.includes('我知道了')
                        || [...el.querySelectorAll('button,[role="button"],a,[aria-label*="关闭"],[aria-label*="Close"],.rocket-modal-close,.rocket-dialog-close')].some(visible);
                    })
                    .sort((a, b) => area(a) - area(b));
                    return containers[0] || null;
                  }

                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """
            )
        )

    def _click_global_button_if_visible(self, page, text: str, timeout: int = 5000) -> bool:
        end_time = page.evaluate("Date.now()") + timeout
        while page.evaluate("Date.now()") < end_time:
            rect = page.evaluate(
                """
                ([targetText]) => {
                  const target = compact(targetText);
                  const candidates = [...document.querySelectorAll('button,[role="button"],a,span,div')]
                    .filter(visible)
                    .filter(el => compact(el.innerText || el.textContent) === target)
                    .map(el => el.closest('button,[role="button"],a') || el)
                    .filter(visible)
                    .filter(el => !isDisabled(el))
                    .sort((a, b) => {
                      const ar = a.getBoundingClientRect();
                      const br = b.getBoundingClientRect();
                      const aDialog = inDialog(a) ? 0 : 1;
                      const bDialog = inDialog(b) ? 0 : 1;
                      return aDialog - bDialog || area(a) - area(b) || br.top - ar.top;
                    });
                  const button = candidates[0];
                  if (!button) return null;
                  button.scrollIntoView({ block: 'center', inline: 'center' });
                  const rect = button.getBoundingClientRect();
                  return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function isDisabled(el) {
                    return el.disabled === true
                      || el.getAttribute('disabled') !== null
                      || el.getAttribute('aria-disabled') === 'true';
                  }
                  function inDialog(el) {
                    return Boolean(el.closest('[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-dialog,.rocket-drawer'));
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [text],
            )
            if rect:
                page.mouse.click(rect["x"], rect["y"])
                page.wait_for_timeout(700)
                return True
            page.wait_for_timeout(500)
        return False

    def _click_transient_site_popup_close(self, page) -> bool:
        rect = page.evaluate(
            """
            () => {
              const closeTexts = [
                '我知道了', '知道了', '关闭', '稍后再说', '稍后处理', '暂不处理', '取消', '跳过',
                'Not now', 'Later', 'Close', 'Cancel'
              ];
              const protectedTexts = [
                '批量上传合规信息', '上传文件', '修改库存', '批量开通JIT',
                '确认下列', '添加同面料套版', '套版组管理'
              ];
              const popups = [...document.querySelectorAll(
                '[role="dialog"],[aria-modal="true"],.rocket-modal,.rocket-dialog,.rocket-popover,'
                + '.MDL_outerWrapper_5-120-1,.MDL_container_5-120-1,[class*="MDL_outerWrapper"],'
                + '[class*="MDL_container"],[class*="modal"],[class*="Modal"],[class*="dialog"],[class*="Dialog"]'
              )]
                .filter(visible)
                .filter(el => area(el) > 1000)
                .filter(el => area(el) < window.innerWidth * window.innerHeight * 0.75)
                .filter(el => {
                  const text = norm(el.innerText || el.textContent);
                  if (protectedTexts.some(item => text.includes(item))) return false;
                  return closeTexts.some(item => text.includes(item))
                    || [...el.querySelectorAll('[aria-label*="关闭"],[aria-label*="Close"],[class*="close"],[class*="Close"]')].some(visible);
                })
                .sort((a, b) => area(a) - area(b));
              const popup = popups[0];
              if (!popup) return null;
              const button = findCloseButton(popup, closeTexts);
              if (!button) return null;
              button.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = button.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findCloseButton(scope, texts) {
                const textButton = [...scope.querySelectorAll('button,[role="button"],a,span,div')]
                  .filter(visible)
                  .filter(el => texts.includes(norm(el.innerText || el.textContent)))
                  .map(el => el.closest('button,[role="button"],a') || el)
                  .filter(visible)
                  .filter(el => !isDisabled(el))
                  .sort((a, b) => buttonScore(a, texts) - buttonScore(b, texts) || area(a) - area(b))[0];
                if (textButton) return textButton;
                return [...scope.querySelectorAll('[aria-label*="关闭"],[aria-label*="Close"],.rocket-modal-close,.rocket-dialog-close,.rocket-icon-close,button,span')]
                  .filter(visible)
                  .filter(el => !isDisabled(el))
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    return text === '' || text === '×' || text === 'x' || text === 'X';
                  })
                  .sort((a, b) => b.getBoundingClientRect().top - a.getBoundingClientRect().top || b.getBoundingClientRect().right - a.getBoundingClientRect().right)[0] || null;
              }
              function buttonScore(el, texts) {
                const text = norm(el.innerText || el.textContent);
                let score = 100;
                if (text === '我知道了' || text === '知道了' || text === '关闭') score -= 80;
                if (text === '稍后再说' || text === '稍后处理' || text === '暂不处理') score -= 60;
                if (text === '取消' || text === 'Cancel') score -= 20;
                if (texts.includes(text)) score -= 10;
                const cls = String(el.className || '');
                if (/primary|BTN_primary|button-primary/.test(cls)) score -= 5;
                return score;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''));
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(700)
        return True

    def _click_global_button(self, page, text: str, timeout: int = 10000) -> None:
        if not self._click_global_button_if_visible(page, text, timeout=timeout):
            raise RuntimeError(f"没有找到或无法点击“{text}”按钮。")

    def _skip_empty_compliance_upload(self, selected_count: int, log_label: str) -> bool:
        if selected_count > 0:
            return False
        self.log(f"当前没有待上传的{log_label}商品，已跳过提交。")
        return True

    def _has_empty_compliance_products(self, page) -> bool:
        return bool(
            page.evaluate(
                """
                ([containerSelector]) => {
                  return [...document.querySelectorAll(containerSelector)]
                    .filter(visible)
                    .some(el => {
                      const text = norm(el.innerText || el.textContent);
                      return text.includes('选择合规信息类型')
                        && text.includes('选择商品')
                        && text.includes('暂无数据');
                    });

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                }
                """,
                [compliance_modal_container_selector()],
            )
        )

    def _set_compliance_modal_page_size(self, page, size: int) -> None:
        rect = page.evaluate(
            """
            ([modalTitle, containerSelector]) => {
              const drawer = findDrawer(modalTitle);
              if (!drawer) return null;
              const body = drawer.querySelector('.rocket-drawer-body') || drawer;
              body.scrollTop = 0;
              const candidates = [...drawer.querySelectorAll('.rocket-pagination-options-size-changer,.rocket-pagination-options .rocket-select,input,button,span,div')]
                .filter(visible)
                .filter(el => /\\d+\\s*条\\/页/.test(norm(el.innerText || el.textContent || el.value)))
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  const aSelect = /rocket-select/.test(String(a.className || '')) ? 0 : 1;
                  const bSelect = /rocket-select/.test(String(b.className || '')) ? 0 : 1;
                  return aSelect - bSelect || br.top - ar.top || br.left - ar.left || area(a) - area(b);
                });
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findDrawer(title) {
                const candidates = [...document.querySelectorAll(containerSelector)]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent).includes(title))
                  .filter(el => el.querySelector('.rocket-pagination-options-size-changer,.rocket-pagination-options'))
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
        )
        if not rect:
            if self._has_empty_compliance_products(page):
                self.log("当前合规筛选结果暂无待上传商品，跳过分页设置。")
                return
            raise RuntimeError("没有找到批量上传弹窗中的分页条数控件。")
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(500)
        selected = False
        for target in page_size_option_labels(size):
            if self._click_compliance_dropdown_option(page, target):
                selected = True
                break
            if self._activate_popup_item(page, target, exact=target.isdigit()):
                selected = True
                break
        if not selected:
            self._fill_compliance_page_size_control(page, size)
        self._wait_for_page_ready(page, timeout=7000)
        page.wait_for_timeout(1200)
        if str(size) in self._read_compliance_page_size_text(page):
            return
        self._fill_compliance_page_size_control(page, size)
        self._wait_for_page_ready(page, timeout=7000)
        page.wait_for_timeout(1200)
        if str(size) not in self._read_compliance_page_size_text(page):
            if self._has_empty_compliance_products(page):
                self.log("当前合规筛选结果暂无待上传商品，跳过分页设置。")
                return
            current_text = self._read_compliance_page_size_text(page) or "未知"
            self.log(f"未能切换到 {size}条/页，当前分页显示为“{current_text}”，继续处理当前页。")

    def _fill_compliance_page_size_control(self, page, size: int) -> None:
        rect = page.evaluate(
            """
            ([modalTitle, containerSelector]) => {
              const drawer = findDrawer(modalTitle);
              if (!drawer) return null;
              const candidates = [...drawer.querySelectorAll('.rocket-pagination-options-size-changer,.rocket-pagination-options .rocket-select,input,button,span,div')]
                .filter(visible)
                .filter(el => {
                  const text = norm((el.innerText || el.textContent || '') + ' ' + (el.value || ''));
                  return /\\d+\\s*条\\/页/.test(text) || /^\\d+$/.test(el.value || '');
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  const aInput = a.tagName === 'INPUT' ? 0 : 1;
                  const bInput = b.tagName === 'INPUT' ? 0 : 1;
                  return aInput - bInput || br.top - ar.top || br.left - ar.left || area(a) - area(b);
                });
              const target = candidates[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findDrawer(title) {
                const candidates = [...document.querySelectorAll(containerSelector)]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent).includes(title))
                  .filter(el => el.querySelector('.rocket-pagination-options-size-changer,.rocket-pagination-options'))
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
        )
        if not rect:
            return
        page.mouse.click(rect["x"], rect["y"])
        page.keyboard.press("Control+A")
        page.keyboard.type(str(size), delay=20)
        page.keyboard.press("Enter")

    def _read_compliance_page_size_text(self, page) -> str:
        return str(
            page.evaluate(
                """
                ([modalTitle, containerSelector]) => {
                  const drawer = [...document.querySelectorAll(containerSelector)]
                    .filter(visible)
                    .filter(el => norm(el.innerText || el.textContent).includes(modalTitle))
                    .filter(el => el.querySelector('.rocket-pagination-options-size-changer,.rocket-pagination-options'))
                    .sort((a, b) => area(a) - area(b))[0];
                  if (!drawer) return '';
                  const target = [...drawer.querySelectorAll('.rocket-pagination-options-size-changer,.rocket-pagination-options')]
                    .filter(visible)
                    .map(el => norm((el.innerText || el.textContent || '') + ' ' + (el.value || '')))
                    .find(text => /\\d+\\s*条\\/页/.test(text));
                  return target || '';

                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
            )
            or ""
        )

    def _select_all_compliance_products(self, page) -> int:
        rect = page.evaluate(
            """
            ([modalTitle, containerSelector]) => {
              const drawer = findDrawer(modalTitle);
              if (!drawer) return null;
              const body = drawer.querySelector('.rocket-drawer-body') || drawer;
              body.scrollTop = 0;
              const productTable = [...drawer.querySelectorAll('.rocket-table,.rocket-table-container,table')]
                .filter(el => area(el) > 1000)
                .filter(el => norm(el.innerText || el.textContent).includes('商品信息'))
                .filter(el => !norm(el.innerText || el.textContent).includes('需填写信息'))
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return Math.abs(ar.top - 360) - Math.abs(br.top - 360) || area(a) - area(b);
                })[0];
              if (productTable) productTable.scrollIntoView({ block: 'center', inline: 'nearest' });
              const scope = productTable || drawer;
              const header = [...scope.querySelectorAll('.rocket-table-thead,thead')]
                .filter(el => norm(el.innerText || el.textContent).includes('商品信息'))[0] || scope;
              const checkboxes = [...scope.querySelectorAll('input[type="checkbox"],[role="checkbox"],.rocket-checkbox')]
                .filter(visible)
                .filter(el => !ancestorText(el).includes('我确认'))
                .filter(el => !ancestorText(el).includes('需填写信息'))
                .filter(el => {
                  if (!header || header === scope) return true;
                  return header.contains(el);
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const target = checkboxes[0];
              if (!target) return null;
              target.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = target.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findDrawer(title) {
                const candidates = [...document.querySelectorAll(containerSelector)]
                  .filter(visible)
                  .filter(el => String(el.innerText || el.textContent || '').includes(title))
                  .filter(el => el.querySelector('.rocket-table,.rocket-table-container,table'))
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function ancestorText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                }
                return chunks.join(' ');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
        )
        if not rect:
            if self._has_empty_compliance_products(page):
                self.log("当前合规筛选结果暂无待上传商品，跳过全选。")
                return 0
            raise RuntimeError("没有找到商品列表全选框。")
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(1200)
        selected_count = self._read_compliance_selected_count(page)
        if selected_count <= 0:
            if self._has_empty_compliance_products(page):
                self.log("当前合规筛选结果暂无待上传商品，跳过全选。")
                return 0
            raise RuntimeError("已点击全选，但没有检测到已选商品数量。")
        self.log(f"已全选合规信息商品：{selected_count} 个。")
        return selected_count

    def _read_compliance_selected_count(self, page) -> int:
        return int(
            page.evaluate(
                """
                ([modalTitle, containerSelector]) => {
                  const drawer = findDrawer(modalTitle);
                  if (!drawer) return 0;
                  const text = String(drawer.innerText || '');
                  const match = text.match(/已选[:：]\\s*(\\d+)/);
                  if (match) return Number(match[1]);
                  const checked = [...drawer.querySelectorAll('input[type="checkbox"]')]
                    .filter(input => input.checked)
                    .filter(input => !ancestorText(input).includes('我确认'));
                  return checked.length;

                  function findDrawer(title) {
                    const candidates = [...document.querySelectorAll(containerSelector)]
                      .filter(visible)
                      .filter(el => String(el.innerText || el.textContent || '').includes(title))
                      .filter(el => el.querySelector('.rocket-table,.rocket-table-container,table'))
                      .sort((a, b) => area(a) - area(b));
                    return candidates[0] || null;
                  }
                  function ancestorText(el) {
                    const chunks = [];
                    let node = el;
                    for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                      chunks.push(node.innerText || node.textContent || '');
                    }
                    return chunks.join(' ');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, compliance_modal_container_selector()],
            )
            or 0
        )

    def _fill_packaging_material_info(self, page) -> None:
        self._select_packaging_table_dropdown(page, "材质分类", "塑料")
        page.wait_for_timeout(800)
        self._select_packaging_table_dropdown(page, "材料名称", "聚丙烯")
        self._select_packaging_table_dropdown(page, "它是否含有一次性塑料", "否")
        self._select_packaging_table_dropdown(page, "包装类型", "软包装")
        self._fill_packaging_table_input(page, "包装材料重量", "8.00")
        self.log("已填写包装材料信息：塑料、聚丙烯、否、软包装、8.00g。")

    def _select_packaging_table_dropdown(self, page, column_text: str, option_text: str) -> None:
        rect = None
        for _attempt in range(8):
            rect = self._packaging_table_control_rect(page, column_text, ".rocket-select", require_enabled=True)
            if rect:
                break
            page.wait_for_timeout(600)
        if not rect:
            raise RuntimeError(f"没有找到包装材料信息中的“{column_text}”下拉框。")

        selected = False
        for search_text in self._compliance_dropdown_search_terms(option_text):
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(400)
            if self._click_compliance_dropdown_option(page, option_text):
                selected = True
                break
            self._search_compliance_dropdown(page, rect, search_text)
            if self._click_compliance_dropdown_option(page, option_text):
                selected = True
                break
        if not selected:
            raise RuntimeError(f"没有找到包装材料信息“{column_text}”的“{option_text}”选项。")
        page.wait_for_timeout(700)
        if not self._packaging_table_cell_has_value(page, column_text, option_text):
            raise RuntimeError(f"已尝试选择“{option_text}”，但包装材料信息“{column_text}”仍未显示为已选择。")

    def _fill_packaging_table_input(self, page, column_text: str, value: str) -> None:
        rect = self._packaging_table_control_rect(page, column_text, "input", require_enabled=True)
        if not rect:
            raise RuntimeError(f"没有找到包装材料信息中的“{column_text}”输入框。")
        page.mouse.click(rect["x"], rect["y"])
        page.keyboard.press("Control+A")
        page.keyboard.type(value, delay=20)
        page.wait_for_timeout(500)
        if not self._packaging_table_cell_has_value(page, column_text, value):
            raise RuntimeError(f"已尝试填写“{value}”，但包装材料信息“{column_text}”仍未显示该值。")

    def _select_compliance_loose_label_dropdown(
        self,
        page,
        label_text: str,
        option_text: str,
        search_text: str | None = None,
    ) -> None:
        rect = None
        for _attempt in range(8):
            rect = self._compliance_loose_label_dropdown_rect(page, label_text)
            if rect:
                break
            page.wait_for_timeout(600)
        if not rect:
            raise RuntimeError(f"没有找到“{label_text}”下拉框。")

        if self._compliance_loose_label_dropdown_has_value(page, label_text, option_text):
            return

        search_terms = [search_text or option_text]
        if option_text not in search_terms:
            search_terms.append(option_text)
        selected = False
        for term in search_terms:
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(300)
            page.keyboard.press("Control+A")
            page.keyboard.press("Backspace")
            page.keyboard.type(term, delay=20)
            page.wait_for_timeout(1000)
            selected = self._click_compliance_dropdown_option_containing(page, option_text)
            if selected:
                break
        if not selected:
            raise RuntimeError(f"没有找到“{option_text}”选项。")
        page.wait_for_timeout(800)
        if not self._compliance_loose_label_dropdown_has_value(page, label_text, option_text):
            raise RuntimeError(f"已尝试选择“{option_text}”，但“{label_text}”仍未显示为已选择。")

    def _compliance_loose_label_dropdown_rect(self, page, label_text: str) -> dict | None:
        return page.evaluate(
            """
            ([modalTitle, labelText, containerSelector]) => {
              const drawer = findDrawer(modalTitle);
              if (!drawer) return null;
              const body = drawer.querySelector('.rocket-drawer-body') || drawer;
              body.scrollTop = body.scrollHeight;
              const target = compact(labelText);
              const labels = [...drawer.querySelectorAll('label,span,div')]
                .filter(visible)
                .filter(el => compact(el.innerText || el.textContent).includes(target))
                .filter(el => !compact(el.innerText || el.textContent).includes('如还未申报'))
                .sort((a, b) => area(a) - area(b));
              for (const label of labels) {
                const labelRect = label.getBoundingClientRect();
                const controls = [...drawer.querySelectorAll('.rocket-select,input,[role="combobox"]')]
                  .filter(visible)
                  .filter(el => {
                    const rect = el.getBoundingClientRect();
                    return rect.width > 100
                      && rect.left > labelRect.left + 20
                      && rect.top < labelRect.bottom + 28
                      && rect.bottom > labelRect.top - 28;
                  })
                  .sort((a, b) => a.getBoundingClientRect().left - b.getBoundingClientRect().left || area(b) - area(a));
                const control = controls[0];
                if (!control) continue;
                control.scrollIntoView({ block: 'center', inline: 'center' });
                const rect = control.getBoundingClientRect();
                return { x: rect.left + 14, y: rect.top + rect.height / 2 };
              }
              return null;

              function findDrawer(title) {
                const candidates = [...document.querySelectorAll(containerSelector)]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent).includes(title))
                  .filter(el => {
                    const target = compact(labelText);
                    return [...el.querySelectorAll('label,span,div')]
                      .filter(visible)
                      .some(label => compact(label.innerText || label.textContent).includes(target));
                  })
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/[\\s:*：＊]/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, label_text, compliance_modal_container_selector()],
        )

    def _click_compliance_dropdown_option_containing(self, page, option_text: str) -> bool:
        rect = page.evaluate(
            """
            ([targetText]) => {
              const target = compact(targetText);
              const dropdowns = [...document.querySelectorAll('.rocket-select-dropdown')]
                .filter(visible)
                .filter(el => !String(el.className || '').includes('rocket-select-dropdown-hidden'));
              const options = dropdowns.flatMap(dropdown =>
                [...dropdown.querySelectorAll('.rocket-select-item-option,[role="option"]')]
                  .filter(visible)
                  .filter(el => compact(el.innerText || el.textContent).includes(target))
              );
              const option = options.sort((a, b) => area(a) - area(b))[0];
              if (!option) return null;
              option.scrollIntoView({ block: 'center', inline: 'center' });
              const rect = option.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [option_text],
        )
        if not rect:
            return False
        page.mouse.click(rect["x"], rect["y"])
        page.wait_for_timeout(500)
        return True

    def _compliance_loose_label_dropdown_has_value(self, page, label_text: str, option_text: str) -> bool:
        return bool(
            page.evaluate(
                """
                ([modalTitle, labelText, optionText, containerSelector]) => {
                  const drawer = findDrawer(modalTitle);
                  if (!drawer) return false;
                  const target = compact(labelText);
                  const expected = compact(optionText);
                  const labels = [...drawer.querySelectorAll('label,span,div')]
                    .filter(visible)
                    .filter(el => compact(el.innerText || el.textContent).includes(target))
                    .filter(el => !compact(el.innerText || el.textContent).includes('如还未申报'))
                    .sort((a, b) => area(a) - area(b));
                  for (const label of labels) {
                    const labelRect = label.getBoundingClientRect();
                    const control = [...drawer.querySelectorAll('.rocket-select,input,[role="combobox"]')]
                      .filter(visible)
                      .filter(el => {
                        const rect = el.getBoundingClientRect();
                        return rect.width > 100
                          && rect.left > labelRect.left + 20
                          && rect.top < labelRect.bottom + 28
                          && rect.bottom > labelRect.top - 28;
                      })[0];
                    if (!control) continue;
                    const text = compact((control.innerText || control.textContent || control.value || '')
                      + ' '
                      + ((control.closest('.rocket-form-field,.rocket-form-item,div') || control).innerText || ''));
                    if (text.includes(expected)) return true;
                  }
                  return false;

                  function findDrawer(title) {
                    const candidates = [...document.querySelectorAll(containerSelector)]
                      .filter(visible)
                      .filter(el => norm(el.innerText || el.textContent).includes(title))
                      .filter(el => {
                        const target = compact(labelText);
                        return [...el.querySelectorAll('label,span,div')]
                          .filter(visible)
                          .some(label => compact(label.innerText || label.textContent).includes(target));
                      })
                      .sort((a, b) => area(a) - area(b));
                    return candidates[0] || null;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/[\\s:*：＊]/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, label_text, option_text, compliance_modal_container_selector()],
            )
        )

    def _packaging_table_control_rect(
        self,
        page,
        column_text: str,
        selector: str,
        require_enabled: bool,
    ) -> dict | None:
        return page.evaluate(
            """
            ([modalTitle, columnText, selector, requireEnabled, containerSelector]) => {
              const drawer = findDrawer(modalTitle);
              if (!drawer) return null;
              const body = drawer.querySelector('.rocket-drawer-body') || drawer;
              body.scrollTop = body.scrollHeight;
              const table = findPackagingTable(drawer);
              if (!table) return null;
              table.scrollIntoView({ block: 'center', inline: 'nearest' });
              const header = findHeader(table, columnText);
              if (!header) return null;
              const headerRect = header.getBoundingClientRect();
              const rows = [...table.querySelectorAll('tbody tr,.rocket-table-row')]
                .filter(visible)
                .sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
              for (const row of rows) {
                const cells = [...row.querySelectorAll('td,.rocket-table-cell')]
                  .filter(visible)
                  .filter(cell => Math.abs(cell.getBoundingClientRect().left - headerRect.left) < 12
                    || overlaps(cell.getBoundingClientRect(), headerRect));
                const cell = cells[0];
                if (!cell) continue;
                const controls = [...cell.querySelectorAll(selector)]
                  .filter(visible)
                  .filter(el => !requireEnabled || !isDisabled(el));
                const control = controls[0];
                if (!control) continue;
                control.scrollIntoView({ block: 'center', inline: 'center' });
                const rect = control.getBoundingClientRect();
                return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };
              }
              return null;

              function findDrawer(title) {
                const candidates = [...document.querySelectorAll(containerSelector)]
                  .filter(visible)
                  .filter(el => norm(el.innerText || el.textContent).includes(title))
                  .filter(el => findPackagingTable(el))
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function findPackagingTable(scope) {
                return [...scope.querySelectorAll('.rocket-table,table')]
                  .filter(visible)
                  .filter(el => {
                    const text = norm(el.innerText || el.textContent);
                    return text.includes('材质分类')
                      && text.includes('材料名称')
                      && text.includes('包装材料重量');
                  })
                  .sort((a, b) => area(a) - area(b))[0] || null;
              }
              function findHeader(table, column) {
                const target = compact(column);
                return [...table.querySelectorAll('th')]
                  .filter(visible)
                  .find(th => compact(th.innerText || th.textContent).includes(target)) || null;
              }
              function overlaps(a, b) {
                return a.left < b.right && a.right > b.left;
              }
              function isDisabled(el) {
                return el.disabled === true
                  || el.getAttribute('disabled') !== null
                  || el.getAttribute('aria-disabled') === 'true'
                  || /disabled/.test(String(el.className || ''))
                  || Boolean(el.closest('[aria-disabled="true"],.rocket-select-disabled'));
              }
              function norm(value) {
                return String(value || '').replace(/\\s+/g, ' ').trim();
              }
              function compact(value) {
                return norm(value).replace(/\\s+/g, '');
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, column_text, selector, require_enabled, compliance_modal_container_selector()],
        )

    def _packaging_table_cell_has_value(self, page, column_text: str, value: str) -> bool:
        return bool(
            page.evaluate(
                """
                ([modalTitle, columnText, expectedValue, containerSelector]) => {
                  const drawer = findDrawer(modalTitle);
                  if (!drawer) return false;
                  const table = findPackagingTable(drawer);
                  if (!table) return false;
                  const header = findHeader(table, columnText);
                  if (!header) return false;
                  const headerRect = header.getBoundingClientRect();
                  const expected = compact(expectedValue);
                  const rows = [...table.querySelectorAll('tbody tr,.rocket-table-row')]
                    .filter(visible)
                    .sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
                  for (const row of rows) {
                    const cell = [...row.querySelectorAll('td,.rocket-table-cell')]
                      .filter(visible)
                      .find(td => Math.abs(td.getBoundingClientRect().left - headerRect.left) < 12
                        || overlaps(td.getBoundingClientRect(), headerRect));
                    if (!cell) continue;
                    const text = compact(cell.innerText || cell.textContent || '');
                    const values = [...cell.querySelectorAll('input,textarea')]
                      .map(input => compact(input.value || input.getAttribute('placeholder') || ''))
                      .join('');
                    if ((text + values).includes(expected)) return true;
                  }
                  return false;

                  function findDrawer(title) {
                    const candidates = [...document.querySelectorAll(containerSelector)]
                      .filter(visible)
                      .filter(el => norm(el.innerText || el.textContent).includes(title))
                      .filter(el => findPackagingTable(el))
                      .sort((a, b) => area(a) - area(b));
                    return candidates[0] || null;
                  }
                  function findPackagingTable(scope) {
                    return [...scope.querySelectorAll('.rocket-table,table')]
                      .filter(visible)
                      .filter(el => {
                        const text = norm(el.innerText || el.textContent);
                        return text.includes('材质分类')
                          && text.includes('材料名称')
                          && text.includes('包装材料重量');
                      })
                      .sort((a, b) => area(a) - area(b))[0] || null;
                  }
                  function findHeader(table, column) {
                    const target = compact(column);
                    return [...table.querySelectorAll('th')]
                      .filter(visible)
                      .find(th => compact(th.innerText || th.textContent).includes(target)) || null;
                  }
                  function overlaps(a, b) {
                    return a.left < b.right && a.right > b.left;
                  }
                  function norm(value) {
                    return String(value || '').replace(/\\s+/g, ' ').trim();
                  }
                  function compact(value) {
                    return norm(value).replace(/\\s+/g, '');
                  }
                  function visible(el) {
                    const rect = el.getBoundingClientRect();
                    const style = getComputedStyle(el);
                    return rect.width > 0 && rect.height > 0
                      && style.display !== 'none'
                      && style.visibility !== 'hidden';
                  }
                  function area(el) {
                    const rect = el.getBoundingClientRect();
                    return rect.width * rect.height;
                  }
                }
                """,
                [TEXT_COMPLIANCE_MODAL_TITLE, column_text, value, compliance_modal_container_selector()],
            )
        )

    def _check_compliance_confirmation(self, page) -> None:
        self._check_compliance_checkbox_by_text(page, "我确认")

    def _check_optional_compliance_checkbox_by_text(self, page, text: str) -> None:
        try:
            self._check_compliance_checkbox_by_text(page, text)
        except RuntimeError as exc:
            if f"没有找到“{text}”勾选框" not in str(exc):
                raise
            self.log(f"未检测到“{text}”勾选框，跳过该项。")

    def _check_compliance_checkbox_by_text(self, page, text: str) -> None:
        rect = page.evaluate(
            """
            ([modalTitle, targetText, containerSelector]) => {
              const drawer = findDrawer(modalTitle);
              if (!drawer) return null;
              const body = drawer.querySelector('.rocket-drawer-body') || drawer;
              body.scrollTop = body.scrollHeight;
              const candidates = [...drawer.querySelectorAll('label,input[type="checkbox"],[role="checkbox"],.rocket-checkbox')]
                .filter(visible)
                .filter(el => ancestorText(el).includes(targetText))
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  const aScore = checkboxScore(a);
                  const bScore = checkboxScore(b);
                  if (aScore !== bScore) return aScore - bScore;
                  return br.top - ar.top || ar.left - br.left;
                });
              const target = candidates[0];
              if (!target) return null;
              if (isChecked(target)) return { already: true };
              const clickable = target.closest('label') || target;
              clickable.scrollIntoView({ block: 'center', inline: 'center' });
              const box = checkboxBox(clickable) || clickable;
              const rect = box.getBoundingClientRect();
              return { x: rect.left + rect.width / 2, y: rect.top + rect.height / 2 };

              function findDrawer(title) {
                const candidates = [...document.querySelectorAll(containerSelector)]
                  .filter(visible)
                  .filter(el => String(el.innerText || el.textContent || '').includes(title))
                  .filter(el => String(el.innerText || el.textContent || '').includes(targetText))
                  .sort((a, b) => area(a) - area(b));
                return candidates[0] || null;
              }
              function ancestorText(el) {
                const chunks = [];
                let node = el;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                  chunks.push(node.innerText || node.textContent || '');
                }
                return chunks.join(' ');
              }
              function isChecked(el) {
                const input = el.matches?.('input[type="checkbox"]')
                  ? el
                  : el.querySelector?.('input[type="checkbox"]');
                if (input?.checked === true) return true;
                if (el.getAttribute('aria-checked') === 'true') return true;
                let node = el;
                for (let depth = 0; node && depth < 4; depth += 1, node = node.parentElement) {
                  const cls = String(node.className || '');
                  if (/checkbox-checked|checked/.test(cls) && !/indeterminate/.test(cls)) return true;
                }
                return false;
              }
              function checkboxBox(el) {
                return el.matches?.('input[type="checkbox"],.rocket-checkbox,[role="checkbox"]')
                  ? el
                  : el.querySelector?.('input[type="checkbox"],.rocket-checkbox,[role="checkbox"]');
              }
              function checkboxScore(el) {
                let value = 10;
                if (el.tagName === 'LABEL') value -= 4;
                if (el.matches?.('input[type="checkbox"]')) value -= 3;
                if (String(el.className || '').includes('rocket-checkbox-wrapper')) value -= 2;
                return value;
              }
              function visible(el) {
                const rect = el.getBoundingClientRect();
                const style = getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              }
              function area(el) {
                const rect = el.getBoundingClientRect();
                return rect.width * rect.height;
              }
            }
            """,
            [TEXT_COMPLIANCE_MODAL_TITLE, text, compliance_modal_container_selector()],
        )
        if not rect:
            raise RuntimeError(f"没有找到“{text}”勾选框。")
        if not rect.get("already"):
            page.mouse.click(rect["x"], rect["y"])
            page.wait_for_timeout(500)
