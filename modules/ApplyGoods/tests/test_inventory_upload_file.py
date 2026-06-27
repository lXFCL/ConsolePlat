from __future__ import annotations

import unittest
import inspect
from pathlib import Path

from openpyxl import load_workbook

from app_paths import runtime_root
import temu_goods
from temu_goods import TemuGoodsList


class InventoryUploadFileTest(unittest.TestCase):
    def test_generate_inventory_setting_upload_file_writes_sku_ids_from_third_row(self) -> None:
        worker = TemuGoodsList(user_data_dir=runtime_root() / ".browser-profile")
        sku_ids = ["11111111111", "22222222222"]

        output_path = worker.generate_inventory_setting_upload_file(sku_ids)

        self.addCleanup(lambda: Path(output_path).unlink(missing_ok=True))
        workbook = load_workbook(output_path)
        worksheet = workbook.worksheets[0]

        self.assertEqual(worksheet.cell(row=1, column=1).value, "SKU ID")
        self.assertEqual(worksheet.cell(row=3, column=1).value, 11111111111)
        self.assertEqual(worksheet.cell(row=4, column=1).value, 22222222222)
        self.assertEqual(worksheet.cell(row=3, column=4).value, 999)
        self.assertEqual(worksheet.cell(row=4, column=4).value, 999)
        self.assertIsNone(worksheet.cell(row=5, column=1).value)
        self.assertIsNone(worksheet.cell(row=5, column=4).value)

    def test_read_only_assets_use_packaged_resource_path(self) -> None:
        source = inspect.getsource(temu_goods)

        self.assertIn("resource_path", source)
        self.assertIn('COMPLIANCE_TEMPLATE_PATH = resource_path("合规相关资料") / "合规.xlsx"', source)
        self.assertIn('INVENTORY_SETTING_TEMPLATE_PATH = resource_path("合规相关资料") / "库存设置.xlsx"', source)
        self.assertIn('COMPLIANT_FRONT_IMAGE_PATH = resource_path("合规相关资料") / "商品正视图.jpg"', source)
        self.assertIn('COMPLIANT_SIDE_IMAGE_PATH = resource_path("合规相关资料") / "商品侧视图.jpg"', source)
        self.assertIn('GENERATED_COMPLIANCE_DIR = runtime_root() / "合规相关资料" / "generated"', source)


if __name__ == "__main__":
    unittest.main()
