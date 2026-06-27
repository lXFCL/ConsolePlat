import os


def find_latest_xlsx_in_directory(folder_path: str) -> str:
    folder_path = (folder_path or "").strip()
    if not folder_path:
        raise RuntimeError("最新Excel文件夹路径为空")
    if not os.path.isdir(folder_path):
        raise RuntimeError(f"最新Excel文件夹不存在：{folder_path}")

    candidates = []
    for name in os.listdir(folder_path):
        if name.startswith("~$"):
            continue
        path = os.path.join(folder_path, name)
        if os.path.isfile(path) and name.lower().endswith(".xlsx"):
            candidates.append(path)
    if not candidates:
        raise RuntimeError(f"未找到Excel文件：{folder_path}")
    return max(candidates, key=lambda p: (os.path.getmtime(p), os.path.basename(p).lower()))


def load_product_rows_from_xlsx(file_path: str):
    file_path = (file_path or "").strip()
    if not file_path:
        raise RuntimeError("Excel路径为空")
    if not os.path.exists(file_path):
        raise RuntimeError(f"Excel文件不存在：{file_path}")

    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise RuntimeError("缺少openpyxl，无法导入Excel") from e

    headers = ["店铺名称", "产品分类", "产品标题", "产品序列号", "颜色"]

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        header_cells = [ws.cell(row=1, column=i + 1).value for i in range(ws.max_column)]
        header_cells = [str(v).strip() if v is not None else "" for v in header_cells]
        idx = {}
        for h in headers:
            try:
                idx[h] = header_cells.index(h) + 1
            except ValueError:
                raise RuntimeError(f"Excel缺少标题列：{h}")

        out = []
        for r in range(2, ws.max_row + 1):
            row = {}
            empty = True
            for h in headers:
                v = ws.cell(row=r, column=idx[h]).value
                s = (str(v).strip() if v is not None else "")
                row_key = {
                    "店铺名称": "shop_name",
                    "产品分类": "category",
                    "产品标题": "title",
                    "产品序列号": "sku",
                    "颜色": "color",
                }[h]
                row[row_key] = s
                if s:
                    empty = False
            if empty:
                continue
            out.append(row)
        return out
    finally:
        try:
            wb.close()
        except Exception:
            pass
