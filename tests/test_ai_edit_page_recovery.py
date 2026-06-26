from pathlib import Path
import site
import base64
import json
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

from PyQt5.QtCore import QProcessEnvironment
from PyQt5.QtGui import QShowEvent
from PyQt5.QtWidgets import QApplication, QDialog, QWidget
from openpyxl import Workbook, load_workbook
from PIL import Image

from consoleplat.adapters.posaiimg_adapter import AIEditJob
from consoleplat.config import AppSettings, SettingsStore
from consoleplat.services.ai_edit_formalize_service import AIEditFormalizeSummary
from consoleplat.services.putaway_sync_service import PutawaySyncSummary
from consoleplat.ui.ai_edit_page import (
    AIEditBackgroundWorker,
    AIEditPage,
    AIEditTaskDetailDialog,
    AIEditTaskRecord,
    BackgroundTaskResult,
)


class _FakeSignal:
    def __init__(self):
        self.callbacks = []

    def connect(self, callback):
        self.callbacks.append(callback)


def _page_with_temp_store(tmp_path, monkeypatch, settings=None):
    path = tmp_path / "settings.json"
    SettingsStore(path).save(settings or AppSettings(ai_edit_api_key="stored-key"))
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(path))
    return AIEditPage(), path


def test_ai_edit_page_builds_job_and_persists_preferences(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image = tmp_path / "input.png"
    image.write_bytes(b"fake")
    page, path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_api_base="https://example.test/v1",
            ai_edit_model="demo-image-edit",
            ai_edit_size="1536x1024",
            default_ai_provider_id="provider-2",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "old-key",
                    "api_base": "https://old.example/v1",
                    "model": "old-model",
                    "size": "1024x1024",
                },
                {
                    "provider_id": "provider-2",
                    "name": "备用接口",
                    "api_key": "stored-key",
                    "api_base": "https://example.test/v1",
                    "model": "demo-image-edit",
                    "size": "1536x1024",
                },
            ],
            posai_gallery_root=str(tmp_path / "gallery-root"),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )

    page._add_image_item(str(image))
    page.prompt_edit.setPlainText("keep subject")
    page.split_collage_check.setChecked(True)
    page.split_count_spin.setValue(12)
    page.total_return_count_spin.setValue(8)
    page.prefix_combo.setCurrentText("SZW")
    page.start_spin.setValue(3113)
    page.test_mode_check.setChecked(False)

    job = page.build_job()
    saved = SettingsStore(path).load()

    assert isinstance(job, AIEditJob)
    assert job.images == [image]
    assert job.api_key == "stored-key"
    assert job.api_base == "https://example.test/v1"
    assert job.model == "demo-image-edit"
    assert job.size == "1536x1024"
    assert job.split_collage is True
    assert job.split_count == 12
    assert job.total_return_count == 8
    assert job.prefix == "SZW"
    assert job.start_number == 3113
    assert job.test_mode is False
    assert str(job.output_dir).startswith(str(tmp_path / "gallery-root" / "SZW"))
    assert saved.ai_edit_split_collage is True
    assert saved.ai_edit_split_count == 12
    assert saved.ai_edit_total_return_count == 8

    page.close()


def test_ai_edit_page_save_persists_prompt_and_reference_dir(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image = tmp_path / "refs" / "input.png"
    image.parent.mkdir(parents=True, exist_ok=True)
    image.write_bytes(b"fake")
    page, path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_prompt="old prompt",
            ai_edit_reference_dir="",
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )

    page._add_image_item(str(image))
    page.prompt_edit.setPlainText("new prompt")
    page._save_preferences()

    saved = SettingsStore(path).load()

    assert saved.ai_edit_prompt == "new prompt"
    assert saved.ai_edit_reference_dir == str(image.parent)

    page.close()


def test_ai_edit_page_can_clear_reference_images(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image = tmp_path / "refs" / "input.png"
    image.parent.mkdir(parents=True, exist_ok=True)
    image.write_bytes(b"fake")
    page, path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_reference_dir=str(image.parent),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )

    page._add_image_item(str(image))
    page.clear_reference_images()

    saved = SettingsStore(path).load()

    assert page.image_list.count() == 0
    assert page.selected_image_path_label.text() == "--"
    assert page._selected_images() == []
    assert saved.ai_edit_reference_dir == ""

    page.close()


def test_ai_edit_page_save_persists_ai_edit_controls(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            publish_prefix="BO",
            publish_start_number=1421,
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )

    page.prefix_combo.setCurrentText("SZW")
    page.start_spin.setValue(3113)
    page.total_return_count_spin.setValue(6)
    page.split_collage_check.setChecked(True)
    page.split_count_spin.setValue(9)
    page.test_mode_check.setChecked(False)
    page._save_preferences()

    saved = SettingsStore(path).load()

    assert saved.publish_prefix == "SZW"
    assert saved.publish_start_number == 3113
    assert saved.ai_edit_total_return_count == 6
    assert saved.ai_edit_split_collage is True
    assert saved.ai_edit_split_count == 9
    assert saved.local_image_test_mode is False

    page.close()


def test_ai_edit_page_restores_tasks_and_deletes_failed(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    (tasks_dir / "ai_edit_tasks.json").write_text(
        '[{"task_id":"20260623120000","title":"failed","status":"失败","stage_text":"失败","progress_percent":0,"logs":[],'
        '"output_dir":"E:/out/a","outputs":[],"failed":[],"warnings":[],"round_sources":[],"active_round_index":0,'
        '"final_transparent_dir":"","final_product_dir":"","xlsx_path":"","split_profile":{},'
        '"job":{"images":[],"prompt":"a","api_key":"","api_base":"https://api.openai.com/v1","model":"gpt-image-2",'
        '"output_dir":"E:/out/a","size":"1024x1024","split_collage":true,"split_count":10,"total_return_count":2,'
        '"prefix":"BO","start_number":1661,"test_mode":true,"gallery_root":"E:/gallery","mockup_root":"E:/mockup","xlsx_root":"E:/xlsx"}},'
        '{"task_id":"20260623120100","title":"success","status":"完成","stage_text":"已完成","progress_percent":100,"logs":[],'
        '"output_dir":"E:/out/b","outputs":[],"failed":[],"warnings":[],"round_sources":[],"active_round_index":0,'
        '"final_transparent_dir":"","final_product_dir":"","xlsx_path":"","split_profile":{},'
        '"job":{"images":[],"prompt":"b","api_key":"","api_base":"https://api.openai.com/v1","model":"gpt-image-2",'
        '"output_dir":"E:/out/b","size":"1024x1024","split_collage":true,"split_count":10,"total_return_count":2,'
        '"prefix":"BO","start_number":1662,"test_mode":true,"gallery_root":"E:/gallery","mockup_root":"E:/mockup","xlsx_root":"E:/xlsx"}}]',
        encoding="utf-8",
    )
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()

    assert page.task_list.count() == 2

    page.delete_failed_tasks()

    assert [record.title for record in page.tasks] == ["success"]
    saved = (tasks_dir / "ai_edit_tasks.json").read_text(encoding="utf-8")
    assert '"title": "failed"' not in saved

    page.close()


def test_ai_edit_page_load_history_does_not_backfill_colors_synchronously(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    (tasks_dir / "ai_edit_tasks.json").write_text(
        '[{"task_id":"20260623120100","title":"success","status":"完成","stage_text":"已完成","progress_percent":100,"logs":[],'
        '"output_dir":"E:/out/b","outputs":[],"failed":[],"warnings":[],"round_sources":[],"active_round_index":0,'
        '"final_transparent_dir":"E:/gallery/final","final_product_dir":"","xlsx_path":"E:/gallery/batch.xlsx","split_profile":{},'
        '"job":{"images":[],"prompt":"b","api_key":"","api_base":"https://api.openai.com/v1","model":"gpt-image-2",'
        '"output_dir":"E:/out/b","size":"1024x1024","split_collage":true,"split_count":10,"total_return_count":2,'
        '"prefix":"BO","start_number":1662,"test_mode":true,"gallery_root":"E:/gallery","mockup_root":"E:/mockup","xlsx_root":"E:/xlsx"}}]',
        encoding="utf-8",
    )
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    called = {"count": 0}

    def fake_backfill(*args, **kwargs):
        called["count"] += 1
        return 1

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.backfill_xlsx_colors_from_transparent_dir", fake_backfill)

    page = AIEditPage()

    assert page.task_list.count() == 1
    assert called["count"] == 0

    page.close()


def test_ai_edit_page_show_task_detail_schedules_backfill_for_selected_task(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_drop_first_per_round=False,
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    record = AIEditTaskRecord(
        task_id="20260623123002",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="prompt", prefix="BO", start_number=1661),
        xlsx_path=str(tmp_path / "batch.xlsx"),
        final_transparent_dir=str(tmp_path / "final-transparent"),
    )
    page.tasks = [record]
    page._rebuild_task_list()

    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["task_id"] = current_record.task_id

    class FakeDialog:
        def __init__(self, record_arg, parent):
            self.record = record_arg
            self.parent = parent

        def exec_(self):
            return 0

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.AIEditTaskDetailDialog", FakeDialog)

    page.show_task_detail(page.task_list.item(0))

    assert scheduled == {"action": "backfill_colors", "task_id": "20260623123002"}

    page.close()


def test_ai_edit_page_copy_split_outputs_and_detail_dialog(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", posai_gallery_root=str(tmp_path / "gallery")),
    )
    record = AIEditTaskRecord(
        task_id="20260623123000",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="prompt",
            split_collage=True,
            split_count=2,
            prefix="BO",
            start_number=1661,
            gallery_root=str(tmp_path / "gallery"),
            mockup_root=str(tmp_path / "mockup"),
            xlsx_root=str(tmp_path / "xlsx"),
        ),
        output_dir=str(tmp_path / "output"),
    )
    split_dir = tmp_path / "output"
    split_dir.mkdir(parents=True)
    split_a = split_dir / "edited_transparent_part_01.png"
    split_b = split_dir / "edited_transparent_part_02.png"
    split_a.write_bytes(b"a")
    split_b.write_bytes(b"b")

    outputs = page._copy_split_outputs_to_final_gallery(record, [split_a, split_b])
    record.outputs = [str(path) for path in outputs]
    record.round_sources = [str(path) for path in outputs]
    dialog = AIEditTaskDetailDialog(record, page)

    assert [path.name for path in outputs] == ["BO-1661.png", "BO-1662.png"]
    assert dialog.start_split_button.isEnabled() is True
    assert dialog.convert_transparent_button.isEnabled() is True
    assert dialog.export_product_button.isEnabled() is False
    assert dialog.export_xlsx_button.isEnabled() is False
    assert dialog.sync_putaway_button.isEnabled() is False
    assert dialog.batch_path_button.text() == str(tmp_path / "output")
    assert dialog.transparent_path_button.text() == str(tmp_path / "output")

    dialog.close()
    page.close()


def test_ai_edit_task_detail_enables_batch_action_buttons_when_outputs_exist(tmp_path):
    app = QApplication.instance() or QApplication([])

    product_dir = tmp_path / "final-product"
    product_dir.mkdir(parents=True, exist_ok=True)
    (product_dir / "BO-1661_title.png").write_bytes(b"image")
    xlsx_path = tmp_path / "batch.xlsx"
    xlsx_path.write_bytes(b"xlsx")
    source = tmp_path / "edited_round_01_transparent.png"
    source.write_bytes(b"image")

    record = AIEditTaskRecord(
        task_id="20260623123001",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="prompt", prefix="BO", start_number=1661),
        round_sources=[str(source)],
        final_product_dir=str(product_dir),
        xlsx_path=str(xlsx_path),
    )

    dialog = AIEditTaskDetailDialog(record)

    assert dialog.export_product_button.isEnabled() is True
    assert dialog.export_xlsx_button.isEnabled() is True
    assert dialog.sync_putaway_button.isEnabled() is True

    dialog.close()


def test_ai_edit_task_detail_uses_two_rows_of_four_action_buttons(tmp_path):
    app = QApplication.instance() or QApplication([])

    source = tmp_path / "edited_round_01_transparent.png"
    source.write_bytes(b"image")
    record = AIEditTaskRecord(
        task_id="20260623223001",
        title="AI 改图 SZW-3338",
        job=AIEditJob(images=[], prompt="prompt", prefix="SZW", start_number=3338),
        round_sources=[str(source)],
    )

    dialog = AIEditTaskDetailDialog(record)

    assert hasattr(dialog, "action_grid")
    assert dialog.action_grid.count() == 8
    assert dialog.action_grid.itemAtPosition(0, 0).widget() is dialog.prev_round_button
    assert dialog.action_grid.itemAtPosition(0, 1).widget() is dialog.next_round_button
    assert dialog.action_grid.itemAtPosition(0, 2).widget() is dialog.convert_transparent_button
    assert dialog.action_grid.itemAtPosition(0, 3).widget() is dialog.start_split_button
    assert dialog.action_grid.itemAtPosition(1, 0).widget() is dialog.edit_split_profile_button
    assert dialog.action_grid.itemAtPosition(1, 1).widget() is dialog.export_product_button
    assert dialog.action_grid.itemAtPosition(1, 2).widget() is dialog.export_xlsx_button
    assert dialog.action_grid.itemAtPosition(1, 3).widget() is dialog.sync_putaway_button

    dialog.close()


def test_ai_edit_task_detail_can_apply_new_start_number(tmp_path):
    app = QApplication.instance() or QApplication([])

    source = tmp_path / "edited_round_01_transparent.png"
    source.write_bytes(b"image")
    record = AIEditTaskRecord(
        task_id="20260623223002",
        title="AI 改图 SZW-3438",
        job=AIEditJob(images=[], prompt="prompt", prefix="SZW", start_number=3438),
        round_sources=[str(source)],
    )

    class _PageStub(QWidget):
        def __init__(self) -> None:
            super().__init__()
            self.calls: list[tuple[AIEditTaskRecord, int]] = []

        def update_task_start_number(self, current_record, start_number) -> None:
            self.calls.append((current_record, start_number))

    page = _PageStub()
    dialog = AIEditTaskDetailDialog(record, parent=page)
    dialog.start_number_spin.setValue(4500)
    dialog.apply_start_number()

    assert len(page.calls) == 1
    assert page.calls[0][0] is record
    assert page.calls[0][1] == 4500

    dialog.close()


def test_ai_edit_page_finalize_task_schedules_post_process_in_background(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_drop_first_per_round=False,
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    page.current_task = AIEditTaskRecord(
        task_id="20260623192000",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=2,
            prefix="BO",
            start_number=1661,
            test_mode=False,
        ),
        output_dir=str(tmp_path),
    )

    scheduled = {}

    def fake_start_post_process(task):
        scheduled["task_id"] = task.task_id

    monkeypatch.setattr(page, "_start_post_process_job", fake_start_post_process)

    class _Summary:
        ok = True
        output_dir = str(tmp_path)
        outputs = [str(tmp_path / "edited_round_01.png")]
        failed = []
        warnings = []
        message = "done"

    page._finalize_task(_Summary(), 0)

    assert scheduled["task_id"] == "20260623192000"

    page.close()


def test_ai_edit_page_split_current_round_schedules_background_job(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    transparent = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(transparent)
    record = AIEditTaskRecord(
        task_id="20260623192100",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=2,
            prefix="BO",
            start_number=1661,
        ),
        output_dir=str(tmp_path),
    )

    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["task_id"] = current_record.task_id

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.split_current_round(record, str(transparent))

    assert scheduled == {"action": "split_current_round", "task_id": "20260623192100"}

    page.close()


def test_ai_edit_page_split_current_round_keeps_full_round_list_for_detail_view(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_drop_first_per_round=False,
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    first_round = tmp_path / "edited_round_01_transparent.png"
    second_round = tmp_path / "edited_round_02_transparent.png"
    first_round.write_bytes(b"round1")
    second_round.write_bytes(b"round2")
    record = AIEditTaskRecord(
        task_id="20260623192101",
        title="AI 改图 SZW-3438",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=25,
            prefix="SZW",
            start_number=3438,
        ),
        output_dir=str(tmp_path),
        round_sources=[str(first_round), str(second_round)],
    )

    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["round_sources"] = list(current_record.round_sources)
        scheduled["start_number"] = current_record.job.start_number

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.split_current_round(record, str(second_round))

    assert record.round_sources == [str(first_round), str(second_round)]
    assert record.job.start_number == 3438
    assert scheduled == {
        "action": "split_current_round",
        "round_sources": [str(second_round)],
        "start_number": 3463,
    }

    page.close()


def test_ai_edit_page_split_current_round_uses_effective_count_when_drop_first_enabled(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_drop_first_per_round=True,
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    first_round = tmp_path / "edited_round_01_transparent.png"
    second_round = tmp_path / "edited_round_02_transparent.png"
    first_round.write_bytes(b"round1")
    second_round.write_bytes(b"round2")
    record = AIEditTaskRecord(
        task_id="20260623192101x",
        title="AI 改图 SZW-3438",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=25,
            prefix="SZW",
            start_number=3438,
        ),
        output_dir=str(tmp_path),
        round_sources=[str(first_round), str(second_round)],
    )

    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["round_sources"] = list(current_record.round_sources)
        scheduled["start_number"] = current_record.job.start_number

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.split_current_round(record, str(second_round))

    assert scheduled == {
        "action": "split_current_round",
        "round_sources": [str(second_round)],
        "start_number": 3462,
    }

    page.close()


def test_ai_edit_background_worker_split_current_round_passes_drop_first_setting(tmp_path, monkeypatch):
    source = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(source)
    final_dir = tmp_path / "final-transparent"
    record = AIEditTaskRecord(
        task_id="20260623192101y",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=2,
            prefix="BO",
            start_number=1661,
        ),
        output_dir=str(tmp_path),
        round_sources=[str(source)],
        final_transparent_dir=str(final_dir),
        split_profile={"x_guides": [500], "y_guides": [500]},
    )
    captured = {}

    def fake_split(source_image, output_dir, split_count, x_guides=None, y_guides=None, original_image=None, drop_first=False):
        captured["drop_first"] = drop_first
        captured["x_guides"] = list(x_guides or [])
        captured["y_guides"] = list(y_guides or [])
        target = Path(output_dir) / "edited_round_01_transparent_part_02.png"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"part2")
        return [str(target)]

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.split_collage_image_with_guides", fake_split)

    worker = AIEditBackgroundWorker(
        "split_current_round",
        record,
        AppSettings(ai_edit_drop_first_per_round=True),
    )

    result = worker._run_split_current_round()

    assert captured == {"drop_first": True, "x_guides": [500], "y_guides": [500]}
    assert result.outputs == [str(final_dir / "BO-1661.png")]


def test_ai_edit_background_worker_renumbers_outputs_and_syncs(tmp_path, monkeypatch):
    final_transparent_dir = tmp_path / "final-transparent"
    final_transparent_dir.mkdir()
    old_a = final_transparent_dir / "SZW-3438.png"
    old_b = final_transparent_dir / "SZW-3439.png"
    old_a.write_bytes(b"old-a")
    old_b.write_bytes(b"old-b")
    final_product_dir = tmp_path / "final-product"
    final_product_dir.mkdir()
    old_product = final_product_dir / "SZW-3438_SZW fixed title.png"
    old_product.write_bytes(b"old-product")
    xlsx_path = tmp_path / "batch.xlsx"
    xlsx_path.write_bytes(b"old-xlsx")
    model_dir = tmp_path / "models"
    model_dir.mkdir()
    putaway_dir = tmp_path / "putaway"
    record = AIEditTaskRecord(
        task_id="20260623192101z",
        title="AI 改图 SZW-3438",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=2,
            prefix="SZW",
            start_number=4500,
            test_mode=False,
        ),
        output_dir=str(tmp_path),
        outputs=[str(final_product_dir / "SZW-3438_title.png")],
        round_sources=[str(old_a), str(old_b)],
        final_transparent_dir=str(final_transparent_dir),
        final_product_dir=str(final_product_dir),
        xlsx_path=str(xlsx_path),
    )
    captured = {}

    def fake_build_product_images(*, print_paths, final_product_dir, product_title, model_dir, **kwargs):
        captured["print_paths"] = [path.name for path in print_paths]
        outputs = []
        for print_path in print_paths:
            output = final_product_dir / f"{print_path.stem}_{product_title}.png"
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_bytes(print_path.read_bytes())
            outputs.append(output)
        return outputs, {path.stem: "黑" for path in print_paths}

    def fake_write_xlsx(**kwargs):
        captured["xlsx_prefix"] = kwargs["prefix"]
        captured["xlsx_start_number"] = kwargs["start_number"]
        captured["xlsx_count"] = kwargs["count"]
        kwargs["xlsx_path"].write_bytes(b"new-xlsx")
        return kwargs["xlsx_path"]

    def fake_sync_putaway_assets(**kwargs):
        captured["sync_source_images_dir"] = kwargs["source_images_dir"]
        captured["sync_source_xlsx_path"] = kwargs["source_xlsx_path"]
        captured["sync_target_data_dir"] = kwargs["target_data_dir"]
        captured["sync_force_replace"] = kwargs["force_replace"]
        return PutawaySyncSummary(ok=True, message="synced")

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.build_product_images", fake_build_product_images)
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.write_xlsx", fake_write_xlsx)
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.sync_putaway_assets", fake_sync_putaway_assets)

    worker = AIEditBackgroundWorker(
        "renumber_outputs",
        record,
        AppSettings(
            szw_product_title="SZW fixed title",
            posai_model_root=str(model_dir),
            putaway_data_dir=str(putaway_dir),
        ),
    )

    result = worker._run_renumber_outputs()

    assert sorted(path.name for path in final_transparent_dir.glob("*.png")) == ["SZW-4500.png", "SZW-4501.png"]
    assert (final_transparent_dir / "SZW-4500.png").read_bytes() == b"old-a"
    assert captured["print_paths"] == ["SZW-4500.png", "SZW-4501.png"]
    assert captured["xlsx_prefix"] == "SZW"
    assert captured["xlsx_start_number"] == 4500
    assert captured["xlsx_count"] == 2
    assert captured["sync_source_images_dir"] == final_product_dir
    assert captured["sync_source_xlsx_path"] == xlsx_path
    assert captured["sync_target_data_dir"] == putaway_dir
    assert captured["sync_force_replace"] is True
    assert result.outputs == [
        str(final_product_dir / "SZW-4500_SZW fixed title.png"),
        str(final_product_dir / "SZW-4501_SZW fixed title.png"),
    ]
    assert result.round_sources == [
        str(final_transparent_dir / "SZW-4500.png"),
        str(final_transparent_dir / "SZW-4501.png"),
    ]
    assert old_product.exists() is False
    assert any("重新编号完成" in line for line in result.logs)


def test_ai_edit_page_update_start_number_schedules_renumber_and_persists_record(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    record = AIEditTaskRecord(
        task_id="20260623192101w",
        title="AI 改图 SZW-3438\n2 轮 · 正式模式",
        job=AIEditJob(images=[], prompt="keep subject", prefix="SZW", start_number=3438),
        output_dir=str(tmp_path),
    )
    page.tasks = [record]
    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["start_number"] = current_record.job.start_number
        scheduled["title"] = current_record.title

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.update_task_start_number(record, 4500)

    assert record.job.start_number == 4500
    assert "SZW-4500" in record.title
    assert "重新编号中" == record.stage_text
    assert scheduled == {
        "action": "renumber_outputs",
        "start_number": 4500,
        "title": record.title,
    }

    page.close()


def test_ai_edit_page_split_current_round_uses_recovered_second_round_start_number(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_drop_first_per_round=False,
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    output_dir = tmp_path / "output"
    output_dir.mkdir(parents=True)
    first_round = output_dir / "edited_round_01_transparent.png"
    second_round = output_dir / "edited_round_02_transparent.png"
    first_round.write_bytes(b"round1")
    second_round.write_bytes(b"round2")
    record = AIEditTaskRecord(
        task_id="20260623192102",
        title="AI edit SZW-3438",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=25,
            total_return_count=2,
            prefix="SZW",
            start_number=3438,
        ),
        output_dir=str(output_dir),
        round_sources=[str(second_round)],
    )

    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["round_sources"] = list(current_record.round_sources)
        scheduled["start_number"] = current_record.job.start_number

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.split_current_round(record, str(second_round))

    assert scheduled == {
        "action": "split_current_round",
        "round_sources": [str(second_round)],
        "start_number": 3463,
    }

    page.close()


def test_ai_edit_page_start_background_job_retains_worker_reference(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    record = AIEditTaskRecord(
        task_id="20260623192200",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="keep subject", prefix="BO", start_number=1661),
        output_dir=str(tmp_path),
    )

    started = {}

    def fake_start(self):
        started["called"] = True

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.QThread.start", fake_start)

    page._start_background_job("post_process", record)

    assert started["called"] is True
    assert len(page._background_threads) == 1
    assert len(page._background_workers) == 1

    page.close()


def test_ai_edit_page_background_post_process_result_updates_formal_outputs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    record = AIEditTaskRecord(
        task_id="20260623192300",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="keep subject", prefix="BO", start_number=1661, test_mode=False),
        output_dir=str(tmp_path),
        final_transparent_dir=str(tmp_path / "final-transparent"),
        final_product_dir=str(tmp_path / "final-product"),
        xlsx_path=str(tmp_path / "batch.xlsx"),
    )
    page.tasks = [record]
    page.current_task = record

    result = BackgroundTaskResult(
        action="post_process",
        task_id="20260623192300",
        outputs=[str(tmp_path / "final-product" / "BO-1661_title.png")],
        round_sources=[str(tmp_path / "final-transparent" / "BO-1661.png")],
        final_transparent_dir=str(tmp_path / "final-transparent"),
        final_product_dir=str(tmp_path / "final-product"),
        xlsx_path=str(tmp_path / "batch.xlsx"),
        warnings=["synced"],
        logs=["formalize done"],
    )

    page._on_background_job_finished(result)

    assert page.current_task.outputs == [str(tmp_path / "final-product" / "BO-1661_title.png")]
    assert page.current_task.round_sources == [str(tmp_path / "final-transparent" / "BO-1661.png")]
    assert page.current_task.final_transparent_dir == str(tmp_path / "final-transparent")
    assert page.current_task.final_product_dir == str(tmp_path / "final-product")
    assert page.current_task.xlsx_path == str(tmp_path / "batch.xlsx")
    assert "synced" in page.current_task.warnings
    assert any("formalize done" in line for line in page.current_task.logs)

    page.close()


def test_ai_edit_page_background_split_result_updates_final_transparent_outputs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    final_transparent_dir = tmp_path / "final-transparent"
    final_transparent_dir.mkdir(parents=True, exist_ok=True)
    split_dir = tmp_path / "output" / "edited_round_01_transparent_split"
    split_dir.mkdir(parents=True, exist_ok=True)
    split_a = split_dir / "edited_round_01_transparent_part_01.png"
    split_b = split_dir / "edited_round_01_transparent_part_02.png"
    split_a.write_bytes(b"a")
    split_b.write_bytes(b"b")

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    record = AIEditTaskRecord(
        task_id="20260623192350",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="keep subject", prefix="BO", start_number=1661, test_mode=False),
        output_dir=str(tmp_path / "output"),
        final_transparent_dir=str(final_transparent_dir),
    )
    page.tasks = [record]
    page.current_task = record

    result = BackgroundTaskResult(
        action="split_current_round",
        task_id="20260623192350",
        outputs=[str(split_a), str(split_b)],
        round_sources=[str(tmp_path / "output" / "edited_round_01_transparent.png")],
        final_transparent_dir=str(final_transparent_dir),
    )

    page._on_background_job_finished(result)

    assert record.outputs == [str(final_transparent_dir / "BO-1661.png"), str(final_transparent_dir / "BO-1662.png")]

    page.close()


def test_ai_edit_page_background_split_result_preserves_rounds_and_updates_second_range(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    final_transparent_dir = tmp_path / "final-transparent"
    final_transparent_dir.mkdir(parents=True, exist_ok=True)
    output_dir = tmp_path / "output"
    split_dir = output_dir / "edited_round_02_transparent_split"
    split_dir.mkdir(parents=True, exist_ok=True)
    split_a = split_dir / "edited_round_02_transparent_part_01.png"
    split_b = split_dir / "edited_round_02_transparent_part_02.png"
    split_a.write_bytes(b"new-second-a")
    split_b.write_bytes(b"new-second-b")
    first_round = output_dir / "edited_round_01_transparent.png"
    second_round = output_dir / "edited_round_02_transparent.png"
    first_round.write_bytes(b"round1")
    second_round.write_bytes(b"round2")

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            ai_edit_drop_first_per_round=True,
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    record = AIEditTaskRecord(
        task_id="20260623192351",
        title="AI 改图 SZW-3438",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            split_collage=True,
            split_count=2,
            prefix="SZW",
            start_number=3438,
            test_mode=False,
        ),
        output_dir=str(output_dir),
        round_sources=[str(first_round), str(second_round)],
        final_transparent_dir=str(final_transparent_dir),
    )
    page.tasks = [record]
    page.current_task = record

    result = BackgroundTaskResult(
        action="split_current_round",
        task_id="20260623192351",
        outputs=[str(split_a), str(split_b)],
        round_sources=[str(second_round)],
        final_transparent_dir=str(final_transparent_dir),
    )

    page._on_background_job_finished(result)

    assert record.round_sources == [str(first_round), str(second_round)]
    assert record.outputs == [str(final_transparent_dir / "SZW-3439.png"), str(final_transparent_dir / "SZW-3440.png")]
    assert (final_transparent_dir / "SZW-3439.png").read_bytes() == b"new-second-a"
    assert (final_transparent_dir / "SZW-3440.png").read_bytes() == b"new-second-b"

    page.close()


def test_ai_edit_task_detail_dialog_shows_failed_and_warning_messages(tmp_path):
    app = QApplication.instance() or QApplication([])

    record = AIEditTaskRecord(
        task_id="20260623123100",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="prompt", prefix="BO", start_number=1661),
        status="失败",
        stage_text="后处理失败",
        progress_percent=0,
        logs=["12:00:00  运行结束"],
        failed=["未找到可正式入库的改图产物"],
        warnings=["已同步 0 张图片"],
    )

    dialog = AIEditTaskDetailDialog(record)
    text = dialog.log.toPlainText()

    assert "未找到可正式入库的改图产物" in text
    assert "已同步 0 张图片" in text
    assert "运行结束" in text

    dialog.close()


def test_ai_edit_page_refreshes_open_task_detail_dialog(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(tmp_path, monkeypatch)
    record = AIEditTaskRecord(
        task_id="20260623123200",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="prompt", prefix="BO", start_number=1661),
        status="运行中",
        stage_text="准备启动",
        progress_percent=5,
        logs=["12:00:00  启动任务"],
    )
    page.current_task = record
    dialog = AIEditTaskDetailDialog(record, page)
    page.task_detail_dialog = dialog

    page._update_current_task(status="失败", stage_text="后处理失败", progress_percent=0)
    page.current_task.failed = ["未找到可正式入库的改图产物"]
    page._append_log("正式模式后处理失败：未找到可正式入库的改图产物")

    assert "后处理失败" in dialog.meta_label.text()
    assert "未找到可正式入库的改图产物" in dialog.log.toPlainText()

    dialog.close()
    page.close()


def test_ai_edit_page_process_env_exports_key_and_user_site(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image = tmp_path / "input.png"
    image.write_bytes(b"fake")
    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            default_ai_provider_id="provider-1",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "stored-key",
                    "api_base": "https://api.openai.com/v1",
                    "model": "gpt-image-2",
                    "size": "1024x1024",
                }
            ],
            posai_gallery_root=str(tmp_path / "gallery-root"),
        ),
    )
    page._add_image_item(str(image))
    page.prompt_edit.setPlainText("go")
    captured = {}

    class ReadySignal:
        def connect(self, _callback):
            return None

    class FakeEnv:
        def __init__(self):
            self.values = {}

        def insert(self, key, value):
            self.values[key] = value

        def value(self, key, default=""):
            return self.values.get(key, default)

    class FakeProcess:
        readyReadStandardOutput = ReadySignal()
        readyReadStandardError = ReadySignal()
        finished = ReadySignal()
        errorOccurred = ReadySignal()

        def __init__(self, _parent=None):
            self.env = FakeEnv()

        def setProgram(self, _program):
            return None

        def setArguments(self, _args):
            return None

        def setWorkingDirectory(self, _cwd):
            return None

        def processEnvironment(self):
            return self.env

        def setProcessEnvironment(self, env):
            captured.update(env.values)

        def start(self):
            return None

        def kill(self):
            return None

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.QProcess", FakeProcess)

    page.start_job()

    assert captured["PYTHONUTF8"] == "1"
    assert captured["PYTHONIOENCODING"] == "utf-8"
    assert captured["CONSOLEPLAT_AI_IMAGE_API_KEY"] == "stored-key"
    assert site.getusersitepackages() in captured["PYTHONPATH"]

    page.close()


def test_ai_edit_page_loads_start_number_from_gallery_on_init(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    gallery_root = tmp_path / "gallery"
    batch_dir = gallery_root / "BO" / "2026" / "06" / "demo" / "final"
    batch_dir.mkdir(parents=True)
    (batch_dir / "BO-1660.png").write_bytes(b"x")

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            posai_gallery_root=str(gallery_root),
        ),
    )

    assert page.start_spin.value() == 1661

    page.close()


def test_ai_edit_page_append_log_updates_progress_from_percent_text(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(tmp_path, monkeypatch)
    page.current_task = AIEditTaskRecord(
        task_id="20260620123456",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体"),
        status="运行中",
    )

    page._append_log("准备处理 20%")

    assert page.current_task.progress_percent == 20

    page.close()


def test_ai_edit_task_detail_prefers_transparent_source_for_split(tmp_path):
    app = QApplication.instance() or QApplication([])

    source_image = tmp_path / "edited.png"
    transparent_image = tmp_path / "edited_transparent.png"
    source_image.write_bytes(b"fake")
    transparent_image.write_bytes(b"fake")
    record = AIEditTaskRecord(
        task_id="20260620123456",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="保留主体",
            split_collage=True,
            split_count=4,
            prefix="BO",
            start_number=1661,
        ),
        output_dir=str(tmp_path),
        outputs=[str(source_image), str(transparent_image)],
    )

    dialog = AIEditTaskDetailDialog(record)

    assert dialog._find_split_source(record) == str(transparent_image)

    dialog.close()


def test_ai_edit_task_detail_round_index_uses_transparent_round_sources_only(tmp_path):
    app = QApplication.instance() or QApplication([])

    transparent_a = tmp_path / "edited_round_01_transparent.png"
    transparent_b = tmp_path / "edited_round_02_transparent.png"
    split_a = tmp_path / "edited_round_01_transparent_part_01.png"
    split_b = tmp_path / "edited_round_02_transparent_part_01.png"
    for path in (transparent_a, transparent_b, split_a, split_b):
        path.write_bytes(b"fake")
    record = AIEditTaskRecord(
        task_id="20260620123456",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="保留主体",
            split_collage=True,
            split_count=25,
            total_return_count=2,
            prefix="BO",
            start_number=1661,
        ),
        output_dir=str(tmp_path),
        round_sources=[str(transparent_a), str(transparent_b)],
        outputs=[str(transparent_a), str(split_a), str(transparent_b), str(split_b)],
    )

    dialog = AIEditTaskDetailDialog(record)

    assert dialog.round_index_label.text() == "1/2"
    assert dialog.source_path_label.text() == str(transparent_a)
    dialog.show_next_round()
    assert dialog.round_index_label.text() == "2/2"
    assert dialog.source_path_label.text() == str(transparent_b)

    dialog.close()


def test_split_profile_editor_dialog_parses_guides(tmp_path):
    app = QApplication.instance() or QApplication([])

    record = AIEditTaskRecord(
        task_id="20260620123456",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体"),
        split_profile={"x_guides": [200], "y_guides": [100]},
    )

    from consoleplat.ui.ai_edit_page import SplitProfileEditorDialog

    dialog = SplitProfileEditorDialog(record, str(tmp_path / "edited_round_01_transparent.png"))
    dialog.x_guides_edit.setText("120, 360, 120")
    dialog.y_guides_edit.setText("80，240")

    assert dialog.parsed_guides() == ([120, 360], [80, 240])

    dialog.close()


def test_split_profile_editor_dialog_uses_fixed_default_guides(tmp_path):
    app = QApplication.instance() or QApplication([])

    image_path = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(image_path)
    record = AIEditTaskRecord(
        task_id="20260623170000",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体", split_collage=True, split_count=25),
        split_profile={},
    )

    from consoleplat.ui.ai_edit_page import SplitProfileEditorDialog

    dialog = SplitProfileEditorDialog(record, str(image_path))

    assert dialog.parsed_guides() == ([428, 805, 1229, 1638], [482, 852, 1229, 1587])

    dialog.close()


def test_split_profile_editor_dialog_can_reset_guides_to_fixed_defaults(tmp_path):
    app = QApplication.instance() or QApplication([])

    image_path = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (500, 1000), (0, 0, 0, 0)).save(image_path)
    record = AIEditTaskRecord(
        task_id="20260623170100",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体", split_collage=True, split_count=25),
        split_profile={"x_guides": [111], "y_guides": [222]},
    )

    from consoleplat.ui.ai_edit_page import SplitProfileEditorDialog

    dialog = SplitProfileEditorDialog(record, str(image_path))
    dialog.reset_guides()

    assert dialog.parsed_guides() == ([428, 805, 1229, 1638], [482, 852, 1229, 1587])

    dialog.close()


def test_split_profile_editor_dialog_updates_text_fields_when_guides_change(tmp_path):
    app = QApplication.instance() or QApplication([])

    image_path = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(image_path)
    record = AIEditTaskRecord(
        task_id="20260623170200",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体", split_collage=True, split_count=25),
        split_profile={},
    )

    from consoleplat.ui.ai_edit_page import SplitProfileEditorDialog

    dialog = SplitProfileEditorDialog(record, str(image_path))
    dialog.preview.set_guides([210, 420, 620, 810], [190, 390, 610, 805])

    assert dialog.x_guides_edit.text() == "210,420,620,810"
    assert dialog.y_guides_edit.text() == "190,390,610,805"

    dialog.close()


def test_split_profile_editor_dialog_defaults_to_current_preview_size(tmp_path):
    app = QApplication.instance() or QApplication([])

    image_path = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(image_path)
    record = AIEditTaskRecord(
        task_id="20260623170201",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体", split_collage=True, split_count=25),
        split_profile={"x_guides": [210, 420], "y_guides": [190, 390]},
    )

    from consoleplat.ui.ai_edit_page import SplitProfileEditorDialog

    dialog = SplitProfileEditorDialog(record, str(image_path))

    assert dialog.zoom_combo.currentText() == "1x"
    assert dialog.preview.minimumWidth() == 480
    assert dialog.preview.minimumHeight() == 480
    assert dialog.size().width() == 980
    assert dialog.size().height() == 720

    dialog.close()


def test_split_profile_editor_dialog_zoom_resizes_preview_and_dialog_without_changing_guides(tmp_path):
    app = QApplication.instance() or QApplication([])

    image_path = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(image_path)
    record = AIEditTaskRecord(
        task_id="20260623170202",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体", split_collage=True, split_count=25),
        split_profile={"x_guides": [210, 420], "y_guides": [190, 390]},
    )

    from consoleplat.ui.ai_edit_page import SplitProfileEditorDialog

    dialog = SplitProfileEditorDialog(record, str(image_path))

    dialog.zoom_combo.setCurrentText("1.5x")
    assert dialog.preview.minimumWidth() == 720
    assert dialog.preview.minimumHeight() == 720
    assert dialog.size().width() > 980
    assert dialog.size().height() > 720
    assert dialog.parsed_guides() == ([210, 420], [190, 390])

    dialog.zoom_combo.setCurrentText("2x")
    assert dialog.preview.minimumWidth() == 960
    assert dialog.preview.minimumHeight() == 960
    assert dialog.parsed_guides() == ([210, 420], [190, 390])

    dialog.close()


def test_ai_edit_page_load_split_profile_uses_fixed_default_guides_without_saved_profile(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(tmp_path, monkeypatch)
    page.split_profile_store = None
    job = AIEditJob(images=[], prompt="保留主体", split_collage=True, split_count=25)

    profile = page._load_split_profile_for_job(job)

    assert profile["x_guides"] == [428, 805, 1229, 1638]
    assert profile["y_guides"] == [482, 852, 1229, 1587]
    assert profile["split_count"] == 25
    assert profile["columns"] == 5
    assert profile["rows"] == 5

    page.close()


def test_split_guide_preview_widget_move_guide_clamps_and_sorts(tmp_path):
    app = QApplication.instance() or QApplication([])

    image_path = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(image_path)

    from consoleplat.ui.ai_edit_page import SplitGuidePreviewWidget

    widget = SplitGuidePreviewWidget(str(image_path), [200, 400, 600, 800], [200, 400, 600, 800])
    widget.move_guide("x", 2, 50)
    widget.move_guide("y", 1, 1200)

    assert widget.guides() == ([50, 200, 400, 800], [200, 600, 800, 999])


def test_ai_edit_task_detail_manual_split_button_uses_current_round_source(tmp_path):
    app = QApplication.instance() or QApplication([])

    transparent_a = tmp_path / "edited_round_01_transparent.png"
    transparent_b = tmp_path / "edited_round_02_transparent.png"
    for path in (transparent_a, transparent_b):
        path.write_bytes(b"fake")

    record = AIEditTaskRecord(
        task_id="20260620123456",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="保留主体",
            split_collage=True,
            split_count=25,
            total_return_count=2,
            prefix="BO",
            start_number=1661,
        ),
        output_dir=str(tmp_path),
        round_sources=[str(transparent_a), str(transparent_b)],
        outputs=[str(transparent_a), str(transparent_b)],
    )

    class _PageStub(QWidget):
        def __init__(self) -> None:
            super().__init__()
            self.calls: list[tuple[AIEditTaskRecord, str]] = []

        def edit_split_profile(self, current_record, source_path) -> None:
            self.calls.append((current_record, source_path))

    page = _PageStub()
    dialog = AIEditTaskDetailDialog(record, parent=page)

    dialog.show_next_round()
    dialog.edit_split_profile()

    assert len(page.calls) == 1
    assert page.calls[0][0] is record
    assert page.calls[0][1] == str(transparent_b)

    dialog.close()


def test_ai_edit_page_edit_split_profile_saves_preview_guides_and_replaces_split_outputs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image_path = tmp_path / "edited_round_01_transparent.png"
    Image.new("RGBA", (1000, 1000), (0, 0, 0, 0)).save(image_path)
    page, _ = _page_with_temp_store(tmp_path, monkeypatch)
    record = AIEditTaskRecord(
        task_id="20260623170300",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体", split_collage=True, split_count=25),
        output_dir=str(tmp_path),
        outputs=[str(image_path)],
    )
    page.tasks = [record]
    page.split_profile_store = None

    class _ProfileStoreStub:
        def __init__(self) -> None:
            self.saved = None

        def save(self, profile) -> None:
            self.saved = profile

    page.split_profile_store = _ProfileStoreStub()

    class _FakeDialog:
        def __init__(self, *_args, **_kwargs):
            pass

        def exec_(self):
            return QDialog.Accepted

        def parsed_guides(self):
            return [210, 420, 630, 840], [205, 405, 615, 825]

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SplitProfileEditorDialog", _FakeDialog)
    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["task_id"] = current_record.task_id
        scheduled["x_guides"] = list(current_record.split_profile.get("x_guides") or [])
        scheduled["y_guides"] = list(current_record.split_profile.get("y_guides") or [])

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.edit_split_profile(record, str(image_path))

    assert scheduled["action"] == "split_current_round"
    assert scheduled["task_id"] == "20260623170300"
    assert scheduled["x_guides"] == [210, 420, 630, 840]
    assert scheduled["y_guides"] == [205, 405, 615, 825]


def test_ai_edit_page_restores_legacy_collage_transparent_sources(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    transparent_image = tmp_path / "edited_transparent.png"
    transparent_image.write_bytes(b"fake")
    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    (tasks_dir / "ai_edit_tasks.json").write_text(
        (
            "[{"
            "\"task_id\":\"20260623120000\","
            "\"title\":\"AI 改图 BO-1661\\n2 轮 · 正式模式\","
            "\"status\":\"完成\","
            "\"stage_text\":\"已完成\","
            "\"progress_percent\":100,"
            "\"logs\":[],"
            f"\"output_dir\":\"{tmp_path.as_posix()}\","
            "\"outputs\":[],"
            "\"failed\":[],"
            "\"warnings\":[],"
            "\"round_sources\":[],"
            f"\"collage_transparent_sources\":[\"{transparent_image.as_posix()}\"],"
            "\"active_round_index\":0,"
            "\"final_transparent_dir\":\"\","
            "\"final_product_dir\":\"\","
            "\"xlsx_path\":\"\","
            "\"split_profile\":{},"
            "\"job\":{"
            "\"images\":[],"
            "\"prompt\":\"保留主体\","
            "\"api_key\":\"\","
            "\"api_base\":\"https://api.openai.com/v1\","
            "\"model\":\"gpt-image-2\","
            f"\"output_dir\":\"{tmp_path.as_posix()}\","
            "\"size\":\"1024x1024\","
            "\"split_collage\":true,"
            "\"split_count\":10,"
            "\"total_return_count\":2,"
            "\"prefix\":\"BO\","
            "\"start_number\":1661,"
            "\"test_mode\":false,"
            "\"gallery_root\":\"\","
            "\"mockup_root\":\"\","
            "\"xlsx_root\":\"\""
            "}"
            "}]"
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()

    assert len(page.tasks[0].round_sources) == 1
    assert Path(page.tasks[0].round_sources[0]) == transparent_image

    page.close()


def test_ai_edit_page_restores_round_sources_from_legacy_publish_outputs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    batch_dir = tmp_path / "AI改图_SZW-3438-SZW-3439_2026.0624.2041.49"
    output_dir = batch_dir / "临时输出"
    final_transparent_dir = batch_dir / "最终透明底"
    output_dir.mkdir(parents=True)
    final_transparent_dir.mkdir(parents=True)

    round_one = output_dir / "edited_round_01.png"
    round_two = output_dir / "edited_round_02.png"
    round_one.write_bytes(b"fake1")
    round_two.write_bytes(b"fake2")
    (final_transparent_dir / "SZW-3438.png").write_bytes(b"a")
    (final_transparent_dir / "SZW-3439.png").write_bytes(b"b")

    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    (tasks_dir / "ai_edit_tasks.json").write_text(
        (
            "[{"
            "\"task_id\":\"20260624204149\","
            "\"title\":\"AI 改图 SZW-3438\\n2 轮 · 正式模式\","
            "\"status\":\"完成\","
            "\"stage_text\":\"已完成\","
            "\"progress_percent\":100,"
            "\"logs\":[],"
            f"\"output_dir\":\"{output_dir.as_posix()}\","
            f"\"outputs\":[\"{(final_transparent_dir / 'SZW-3438.png').as_posix()}\",\"{(final_transparent_dir / 'SZW-3439.png').as_posix()}\"] ,"
            "\"failed\":[],"
            "\"warnings\":[],"
            "\"round_sources\":[],"
            "\"collage_transparent_sources\":[],"
            "\"active_round_index\":0,"
            f"\"final_transparent_dir\":\"{final_transparent_dir.as_posix()}\","
            "\"final_product_dir\":\"\","
            "\"xlsx_path\":\"\","
            "\"split_profile\":{},"
            "\"job\":{"
            "\"images\":[],"
            "\"prompt\":\"保留主体\","
            "\"api_key\":\"\","
            "\"api_base\":\"https://api.openai.com/v1\","
            "\"model\":\"gpt-image-2\","
            f"\"output_dir\":\"{output_dir.as_posix()}\","
            "\"size\":\"1024x1024\","
            "\"split_collage\":true,"
            "\"split_count\":25,"
            "\"total_return_count\":2,"
            "\"prefix\":\"SZW\","
            "\"start_number\":3438,"
            "\"test_mode\":false,"
            "\"gallery_root\":\"\","
            "\"mockup_root\":\"\","
            "\"xlsx_root\":\"\""
            "}"
            "}]"
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()

    assert len(page.tasks[0].round_sources) == 2
    assert Path(page.tasks[0].round_sources[0]) == round_one
    assert Path(page.tasks[0].round_sources[1]) == round_two

    page.close()


def test_ai_edit_page_rewrites_legacy_round_sources_when_saved_as_final_transparents(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    batch_dir = tmp_path / "AI改图_SZW-3438-SZW-3439_2026.0624.2041.49"
    output_dir = batch_dir / "临时输出"
    final_transparent_dir = batch_dir / "最终透明底"
    output_dir.mkdir(parents=True)
    final_transparent_dir.mkdir(parents=True)

    round_one = output_dir / "edited_round_01.png"
    round_two = output_dir / "edited_round_02.png"
    round_one.write_bytes(b"fake1")
    round_two.write_bytes(b"fake2")
    final_one = final_transparent_dir / "SZW-3438.png"
    final_two = final_transparent_dir / "SZW-3439.png"
    final_one.write_bytes(b"a")
    final_two.write_bytes(b"b")

    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    (tasks_dir / "ai_edit_tasks.json").write_text(
        (
            "[{"
            "\"task_id\":\"20260624204149\","
            "\"title\":\"AI 改图 SZW-3438\\n2 轮 · 正式模式\","
            "\"status\":\"完成\","
            "\"stage_text\":\"已完成\","
            "\"progress_percent\":100,"
            "\"logs\":[],"
            f"\"output_dir\":\"{output_dir.as_posix()}\","
            f"\"outputs\":[\"{final_one.as_posix()}\",\"{final_two.as_posix()}\"] ,"
            "\"failed\":[],"
            "\"warnings\":[],"
            f"\"round_sources\":[\"{final_one.as_posix()}\",\"{final_two.as_posix()}\"],"
            "\"collage_transparent_sources\":[],"
            "\"active_round_index\":0,"
            f"\"final_transparent_dir\":\"{final_transparent_dir.as_posix()}\","
            "\"final_product_dir\":\"\","
            "\"xlsx_path\":\"\","
            "\"split_profile\":{},"
            "\"job\":{"
            "\"images\":[],"
            "\"prompt\":\"保留主体\","
            "\"api_key\":\"\","
            "\"api_base\":\"https://api.openai.com/v1\","
            "\"model\":\"gpt-image-2\","
            f"\"output_dir\":\"{output_dir.as_posix()}\","
            "\"size\":\"1024x1024\","
            "\"split_collage\":true,"
            "\"split_count\":25,"
            "\"total_return_count\":2,"
            "\"prefix\":\"SZW\","
            "\"start_number\":3438,"
            "\"test_mode\":false,"
            "\"gallery_root\":\"\","
            "\"mockup_root\":\"\","
            "\"xlsx_root\":\"\""
            "}"
            "}]"
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()

    assert len(page.tasks[0].round_sources) == 2
    assert Path(page.tasks[0].round_sources[0]) == round_one
    assert Path(page.tasks[0].round_sources[1]) == round_two

    page.close()


def test_ai_edit_page_legacy_round_restore_prefers_transparent_rounds_only(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    batch_dir = tmp_path / "AI改图_SZW-3438-SZW-3439_2026.0624.2041.49"
    output_dir = batch_dir / "临时输出"
    final_transparent_dir = batch_dir / "最终透明底"
    output_dir.mkdir(parents=True)
    final_transparent_dir.mkdir(parents=True)

    (output_dir / "edited_round_01.png").write_bytes(b"orig1")
    (output_dir / "edited_round_02.png").write_bytes(b"orig2")
    round_one = output_dir / "edited_round_01_transparent.png"
    round_two = output_dir / "edited_round_02_transparent.png"
    round_one.write_bytes(b"transparent1")
    round_two.write_bytes(b"transparent2")
    final_one = final_transparent_dir / "SZW-3438.png"
    final_two = final_transparent_dir / "SZW-3439.png"
    final_one.write_bytes(b"a")
    final_two.write_bytes(b"b")

    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    (tasks_dir / "ai_edit_tasks.json").write_text(
        (
            "[{"
            "\"task_id\":\"20260624204149\","
            "\"title\":\"AI 改图 SZW-3438\\n2 轮 · 正式模式\","
            "\"status\":\"完成\","
            "\"stage_text\":\"已完成\","
            "\"progress_percent\":100,"
            "\"logs\":[],"
            f"\"output_dir\":\"{output_dir.as_posix()}\","
            f"\"outputs\":[\"{final_one.as_posix()}\",\"{final_two.as_posix()}\"] ,"
            "\"failed\":[],"
            "\"warnings\":[],"
            f"\"round_sources\":[\"{final_one.as_posix()}\",\"{final_two.as_posix()}\"],"
            "\"collage_transparent_sources\":[],"
            "\"active_round_index\":0,"
            f"\"final_transparent_dir\":\"{final_transparent_dir.as_posix()}\","
            "\"final_product_dir\":\"\","
            "\"xlsx_path\":\"\","
            "\"split_profile\":{},"
            "\"job\":{"
            "\"images\":[],"
            "\"prompt\":\"保留主体\","
            "\"api_key\":\"\","
            "\"api_base\":\"https://api.openai.com/v1\","
            "\"model\":\"gpt-image-2\","
            f"\"output_dir\":\"{output_dir.as_posix()}\","
            "\"size\":\"1024x1024\","
            "\"split_collage\":true,"
            "\"split_count\":25,"
            "\"total_return_count\":2,"
            "\"prefix\":\"SZW\","
            "\"start_number\":3438,"
            "\"test_mode\":false,"
            "\"gallery_root\":\"\","
            "\"mockup_root\":\"\","
            "\"xlsx_root\":\"\""
            "}"
            "}]"
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()

    assert page.tasks[0].round_sources == [str(round_one), str(round_two)]

    page.close()


def test_ai_edit_page_repairs_short_saved_round_sources_from_batch_rounds(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    batch_dir = tmp_path / "ai-edit-batch"
    output_dir = batch_dir / "temp-output"
    final_transparent_dir = batch_dir / "final-transparent"
    output_dir.mkdir(parents=True)
    final_transparent_dir.mkdir(parents=True)

    round_one = output_dir / "edited_round_01_transparent.png"
    round_two = output_dir / "edited_round_02_transparent.png"
    round_one.write_bytes(b"round1")
    round_two.write_bytes(b"round2")

    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    payload = [
        {
            "task_id": "20260624204150",
            "title": "AI edit SZW-3438",
            "status": "done",
            "stage_text": "done",
            "progress_percent": 100,
            "logs": [],
            "output_dir": str(output_dir),
            "outputs": [],
            "failed": [],
            "warnings": [],
            "round_sources": [str(round_two)],
            "collage_transparent_sources": [],
            "active_round_index": 1,
            "final_transparent_dir": str(final_transparent_dir),
            "final_product_dir": "",
            "xlsx_path": "",
            "split_profile": {},
            "job": {
                "images": [],
                "prompt": "keep subject",
                "api_key": "",
                "api_base": "https://api.openai.com/v1",
                "model": "gpt-image-2",
                "output_dir": str(output_dir),
                "size": "1024x1024",
                "split_collage": True,
                "split_count": 25,
                "total_return_count": 2,
                "prefix": "SZW",
                "start_number": 3438,
                "test_mode": False,
                "gallery_root": "",
                "mockup_root": "",
                "xlsx_root": "",
            },
        }
    ]
    (tasks_dir / "ai_edit_tasks.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()
    dialog = AIEditTaskDetailDialog(page.tasks[0], page)

    assert page.tasks[0].round_sources == [str(round_one), str(round_two)]
    assert dialog.round_index_label.text() == "1/2"

    dialog.close()
    page.close()


def test_ai_edit_page_backfills_xlsx_colors_when_opening_task_detail(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))

    final_transparent_dir = tmp_path / "final_transparent"
    final_transparent_dir.mkdir(parents=True)
    Image.new("RGBA", (32, 32), (0, 0, 0, 255)).save(final_transparent_dir / "SZW-3113.png")
    Image.new("RGBA", (32, 32), (240, 240, 240, 255)).save(final_transparent_dir / "SZW-3114.png")

    xlsx_path = tmp_path / "batch.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["店铺名称", "产品分类", "产品标题", "产品序列号", "颜色"])
    ws.append(["YUHOOBO", "T恤", "固定标题", "SZW-3113", ""])
    ws.append(["YUHOOBO", "T恤", "固定标题", "SZW-3114", ""])
    wb.save(xlsx_path)
    wb.close()

    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True)
    (tasks_dir / "ai_edit_tasks.json").write_text(
        (
            "[{"
            "\"task_id\":\"20260623120000\","
            "\"title\":\"AI 改图 SZW-3113\\n2 轮 · 正式模式\","
            "\"status\":\"完成\","
            "\"stage_text\":\"已完成\","
            "\"progress_percent\":100,"
            "\"logs\":[],"
            f"\"output_dir\":\"{tmp_path.as_posix()}\","
            "\"outputs\":[],"
            "\"failed\":[],"
            "\"warnings\":[],"
            "\"round_sources\":[],"
            "\"collage_transparent_sources\":[],"
            "\"active_round_index\":0,"
            f"\"final_transparent_dir\":\"{final_transparent_dir.as_posix()}\","
            "\"final_product_dir\":\"\","
            f"\"xlsx_path\":\"{xlsx_path.as_posix()}\","
            "\"split_profile\":{},"
            "\"job\":{"
            "\"images\":[],"
            "\"prompt\":\"保留主体\","
            "\"api_key\":\"\","
            "\"api_base\":\"https://api.openai.com/v1\","
            "\"model\":\"gpt-image-2\","
            f"\"output_dir\":\"{tmp_path.as_posix()}\","
            "\"size\":\"1024x1024\","
            "\"split_collage\":true,"
            "\"split_count\":10,"
            "\"total_return_count\":2,"
            "\"prefix\":\"SZW\","
            "\"start_number\":3113,"
            "\"test_mode\":false,"
            "\"gallery_root\":\"\","
            "\"mockup_root\":\"\","
            "\"xlsx_root\":\"\""
            "}"
            "}]"
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()
    worker = AIEditBackgroundWorker("backfill_colors", page.tasks[0], SettingsStore(settings_path).load())
    worker._run_backfill_colors()

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    assert ws["E2"].value == "白"
    assert ws["E3"].value == "白"
    wb.close()

    page.close()


def test_ai_edit_page_show_event_reloads_mirror_tasks(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    settings_path = tmp_path / "settings.json"
    program_data_dir = tmp_path / "ConsolePlatData"
    SettingsStore(settings_path).save(AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(program_data_dir)))
    tasks_dir = program_data_dir / "tasks"
    tasks_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.SettingsStore", lambda: SettingsStore(settings_path))

    page = AIEditPage()
    (tasks_dir / "ai_edit_tasks.json").write_text(
        (
            "[{"
            "\"task_id\":\"20260624153500\","
            "\"title\":\"AI mirror\","
            "\"status\":\"完成\","
            "\"stage_text\":\"已完成\","
            "\"progress_percent\":100,"
            "\"logs\":[\"13:05 done\"],"
            "\"output_dir\":\"E:/ai/output\","
            "\"outputs\":[\"E:/ai/output/1.png\"],"
            "\"failed\":[],"
            "\"warnings\":[],"
            "\"round_sources\":[\"E:/ai/output/1.png\"],"
            "\"collage_transparent_sources\":[],"
            "\"active_round_index\":0,"
            "\"final_transparent_dir\":\"E:/ai/transparent\","
            "\"final_product_dir\":\"E:/ai/products\","
            "\"xlsx_path\":\"E:/ai.xlsx\","
            "\"split_profile\":{},"
            "\"job\":{"
            "\"images\":[\"E:/ref.png\"],"
            "\"prompt\":\"keep subject\","
            "\"api_key\":\"\","
            "\"api_base\":\"https://api.openai.com/v1\","
            "\"model\":\"gpt-image-2\","
            "\"output_dir\":\"E:/ai/output\","
            "\"size\":\"1024x1024\","
            "\"split_collage\":true,"
            "\"split_count\":25,"
            "\"total_return_count\":1,"
            "\"prefix\":\"BO\","
            "\"start_number\":1661,"
            "\"test_mode\":true,"
            "\"gallery_root\":\"E:/gallery\","
            "\"mockup_root\":\"E:/mockup\","
            "\"xlsx_root\":\"E:/xlsx\""
            "}"
            "}]"
        ),
        encoding="utf-8",
    )

    page.showEvent(QShowEvent())

    assert any(record.task_id == "20260624153500" for record in page.tasks)
    assert page.task_list.count() == 1
    assert "AI mirror" in page.task_list.item(0).text()

    page.close()


def test_ai_edit_page_reads_remaining_stdout_on_finish(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            default_ai_provider_id="provider-1",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "stored-key",
                    "api_base": "https://api.openai.com/v1",
                    "model": "gpt-image-2",
                    "size": "1024x1024",
                }
            ],
            posai_gallery_root=str(tmp_path / "gallery-root"),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    page.current_task = AIEditTaskRecord(
        task_id="20260623125900",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="保留主体"),
        output_dir=str(tmp_path),
        final_transparent_dir=str(tmp_path / "final-transparent"),
    )

    class FakeProcess:
        def readAllStandardOutput(self):
            return (
                bytes(
                    json.dumps(
                        {
                            "output_dir": str(tmp_path / "out").replace("\\", "/"),
                            "outputs": [str((tmp_path / "out" / "a.png")).replace("\\", "/")],
                            "failed": [],
                            "warnings": [],
                            "message": "ok",
                        }
                    ),
                    "utf-8",
                )
            )

        def readAllStandardError(self):
            return b""

    page.process = FakeProcess()
    page._stdout_buffer = ""
    page._stderr_buffer = ""

    page._on_process_finished(0, None)

    assert page.current_task.status == "完成"
    assert page.current_task.output_dir == str(tmp_path / "out").replace("\\", "/")
    assert page.current_task.outputs

    page.close()


def test_ai_edit_page_formal_mode_passes_model_dir_from_settings(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    model_root = tmp_path / "models"
    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            bo_product_title="BO title",
            putaway_data_dir=str(tmp_path / "putaway-data"),
            posai_model_root=str(model_root),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    transparent_dir = tmp_path / "final-transparent"
    transparent_dir.mkdir(parents=True, exist_ok=True)
    split_path = transparent_dir / "BO-1661.png"
    split_path.write_bytes(b"image")
    page.current_task = AIEditTaskRecord(
        task_id="20260623130110",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="keep subject",
            output_dir=tmp_path / "output",
            prefix="BO",
            start_number=1661,
            split_collage=False,
            split_count=1,
            test_mode=False,
        ),
        output_dir=str(tmp_path / "output"),
        outputs=[str(split_path)],
        final_transparent_dir=str(transparent_dir),
        final_product_dir=str(tmp_path / "final-product"),
        xlsx_path=str(tmp_path / "final-product.xlsx"),
    )

    captured = {}

    def fake_formalize(**kwargs):
        captured["model_dir"] = kwargs["model_dir"]
        return AIEditFormalizeSummary(
            ok=True,
            renamed_outputs=[str(split_path)],
            product_outputs=[str(tmp_path / "final-product" / "BO-1661_title.png")],
            xlsx_path=str(tmp_path / "final-product.xlsx"),
            putaway=PutawaySyncSummary(ok=True, message="synced"),
            message="formalize done",
        )

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.formalize_ai_edit_outputs", fake_formalize)

    page._run_formalize_post_process()

    assert captured["model_dir"] == model_root

    page.close()


def test_ai_edit_page_start_job_runs_cli_and_finishes_success(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image = tmp_path / "input.png"
    Image.new("RGBA", (32, 32), (255, 0, 0, 255)).save(image)

    output_image = tmp_path / "edited.png"
    Image.new("RGBA", (32, 32), (0, 0, 0, 255)).save(output_image)
    image_b64 = base64.b64encode(output_image.read_bytes()).decode("ascii")

    class EditHandler(BaseHTTPRequestHandler):
        def do_POST(self):  # noqa: N802
            length = int(self.headers.get("Content-Length", "0"))
            if length:
                self.rfile.read(length)
            body = {
                "data": [
                    {
                        "b64_json": image_b64,
                        "revised_prompt": "done",
                    }
                ]
            }
            payload = json.dumps(body).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)

        def log_message(self, format, *args):  # noqa: A003
            return

    server = HTTPServer(("127.0.0.1", 0), EditHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    api_base = f"http://127.0.0.1:{server.server_port}/v1"

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            default_ai_provider_id="provider-1",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "stored-key",
                    "api_base": api_base,
                    "model": "gpt-image-2",
                    "size": "1024x1024",
                }
            ],
            posai_gallery_root=str(tmp_path / "gallery-root"),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    page._add_image_item(str(image))
    page.prompt_edit.setPlainText("keep subject")
    try:
        page.start_job()

        assert page.process is not None
        finished = page.process.waitForFinished(15000)

        assert finished is True
        assert page.current_task is not None
        assert page.current_task.status == "完成"
        assert page.current_task.outputs
        assert Path(page.current_task.outputs[0]).exists()
    finally:
        page.close()
        server.shutdown()
        server.server_close()


def test_ai_edit_page_start_job_logs_current_api_base(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image = tmp_path / "input.png"
    image.write_bytes(b"fake")
    api_base = "https://api.example.test/v1"

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            default_ai_provider_id="provider-1",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "stored-key",
                    "api_base": api_base,
                    "model": "gpt-image-2",
                    "size": "1024x1024",
                }
            ],
            posai_gallery_root=str(tmp_path / "gallery-root"),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    page._add_image_item(str(image))
    page.prompt_edit.setPlainText("keep subject")

    class FakeProcess:
        def __init__(self, *_args, **_kwargs):
            self.program = ""
            self.arguments = []
            self.cwd = ""
            self.env = None
            self.readyReadStandardOutput = _FakeSignal()
            self.readyReadStandardError = _FakeSignal()
            self.finished = _FakeSignal()
            self.errorOccurred = _FakeSignal()

        def setProgram(self, program):
            self.program = program

        def setArguments(self, arguments):
            self.arguments = list(arguments)

        def setWorkingDirectory(self, cwd):
            self.cwd = cwd

        def processEnvironment(self):
            return QProcessEnvironment.systemEnvironment()

        def setProcessEnvironment(self, env):
            self.env = env

        def start(self):
            return None

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.QProcess", FakeProcess)
    try:
        page.start_job()

        assert page.current_task is not None
        assert any(f"当前 AI 接口：{api_base}" in line for line in page.current_task.logs)
    finally:
        page.close()


def test_ai_edit_page_start_job_validates_required_inputs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    image = tmp_path / "input.png"
    image.write_bytes(b"fake")

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            default_ai_provider_id="provider-1",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "",
                    "api_base": "https://api.openai.com/v1",
                    "model": "gpt-image-2",
                    "size": "1024x1024",
                }
            ],
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )

    page.start_job()
    assert page.process is None
    assert page.status_label.text() == "请先添加参考图"

    page._add_image_item(str(image))
    page.prompt_edit.clear()
    page.start_job()
    assert page.process is None
    assert page.status_label.text() == "请先填写改图要求"

    page.prompt_edit.setPlainText("keep subject")
    page.start_job()
    assert page.process is None
    assert page.status_label.text() == "请先配置 API Key"

    page.close()


def test_ai_edit_page_debounces_log_and_progress_persistence(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    record = AIEditTaskRecord(
        task_id="20260624171000",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="keep subject", prefix="BO", start_number=1661),
        status="运行中",
        stage_text="生图中",
        progress_percent=5,
    )
    page.current_task = record
    page.tasks = [record]
    page._rebuild_task_list()
    save_calls = []
    rebuild_calls = []
    monkeypatch.setattr(page, "_save_task_history", lambda: save_calls.append("save"))
    monkeypatch.setattr(page, "_rebuild_task_list", lambda: rebuild_calls.append("rebuild"))

    page._append_log("progress 20\nprogress 40")
    page._update_current_task(status="运行中", stage_text="生图中", progress_percent=42)

    assert save_calls == []
    assert rebuild_calls == []
    assert "42%" in page.task_list.item(0).text()

    page._flush_persist()

    assert save_calls == ["save"]

    page.close()


def test_ai_edit_page_formal_mode_runs_post_process_and_updates_outputs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            default_ai_provider_id="provider-1",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "stored-key",
                    "api_base": "https://api.openai.com/v1",
                    "model": "gpt-image-2",
                    "size": "1024x1024",
                }
            ],
            bo_product_title="BO固定产品标题",
            putaway_data_dir=str(tmp_path / "putaway-data"),
            posai_gallery_root=str(tmp_path / "gallery-root"),
            posai_mockup_root=str(tmp_path / "mockup-root"),
            posai_xlsx_root=str(tmp_path / "xlsx-root"),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    page.current_task = AIEditTaskRecord(
        task_id="20260623130000",
        title="AI 改图 BO-1661",
        job=AIEditJob(
            images=[],
            prompt="保留主体",
            output_dir=tmp_path / "output",
            prefix="BO",
            start_number=1661,
            total_return_count=2,
            split_collage=True,
            split_count=2,
            test_mode=False,
        ),
        output_dir=str(tmp_path / "output"),
        final_transparent_dir=str(tmp_path / "final-transparent"),
        final_product_dir=str(tmp_path / "final-product"),
        xlsx_path=str(tmp_path / "final-product.xlsx"),
    )

    scheduled: dict[str, object] = {}

    def fake_start_post_process(task):
        scheduled["task_id"] = task.task_id
        scheduled["outputs"] = list(task.outputs)

    monkeypatch.setattr(page, "_start_post_process_job", fake_start_post_process)

    class FakeProcess:
        def readAllStandardOutput(self):
            return (
                b'{"output_dir":"E:/tmp/out","outputs":["E:/tmp/out/a_part_1.png"],"failed":[],"warnings":[],"message":"ok"}'
            )

        def readAllStandardError(self):
            return b""

    page.process = FakeProcess()
    page._stdout_buffer = ""
    page._stderr_buffer = ""

    page._on_process_finished(0, None)

    assert scheduled["task_id"] == "20260623130000"
    assert scheduled["outputs"] == ["E:/tmp/out/a_part_1.png"]
    assert page.current_task.status == "完成"
    assert page.current_task.stage_text == "已完成"
    assert page.current_task.progress_percent == 100
    assert page.current_task.final_transparent_dir == str(tmp_path / "final-transparent")
    assert page.current_task.final_product_dir == str(tmp_path / "final-product")
    assert page.current_task.xlsx_path == str(tmp_path / "final-product.xlsx")
    assert page.current_task.outputs == ["E:/tmp/out/a_part_1.png"]

    page.close()


def test_ai_edit_page_show_event_skips_reload_when_task_file_unchanged(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    program_data_dir = Path(SettingsStore(path).load().program_data_dir)
    tasks_file = program_data_dir / "tasks" / "ai_edit_tasks.json"
    tasks_file.parent.mkdir(parents=True, exist_ok=True)
    tasks_file.write_text("[]", encoding="utf-8")
    page.showEvent(QShowEvent())
    load_calls = []
    rebuild_calls = []
    monkeypatch.setattr(page.task_store, "load", lambda: load_calls.append("load") or [])
    monkeypatch.setattr(page, "_rebuild_task_list", lambda: rebuild_calls.append("rebuild"))

    page.showEvent(QShowEvent())

    assert load_calls == []
    assert rebuild_calls == []

    page.close()


def test_ai_edit_page_formal_mode_accepts_non_split_png_outputs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            default_ai_provider_id="provider-1",
            ai_providers=[
                {
                    "provider_id": "provider-1",
                    "name": "主接口",
                    "api_key": "stored-key",
                    "api_base": "https://api.openai.com/v1",
                    "model": "gpt-image-2",
                    "size": "1024x1024",
                }
            ],
            szw_product_title="SZW固定产品标题",
            putaway_data_dir=str(tmp_path / "putaway-data"),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    page.current_task = AIEditTaskRecord(
        task_id="20260623130100",
        title="AI 改图 SZW-3113",
        job=AIEditJob(
            images=[],
            prompt="保留主体",
            output_dir=tmp_path / "output",
            prefix="SZW",
            start_number=3113,
            total_return_count=1,
            split_collage=False,
            split_count=1,
            test_mode=False,
        ),
        output_dir=str(tmp_path / "output"),
        final_transparent_dir=str(tmp_path / "final-transparent"),
        final_product_dir=str(tmp_path / "final-product"),
        xlsx_path=str(tmp_path / "final-product.xlsx"),
    )

    scheduled: dict[str, object] = {}

    def fake_start_post_process(task):
        scheduled["task_id"] = task.task_id
        scheduled["outputs"] = list(task.outputs)

    monkeypatch.setattr(page, "_start_post_process_job", fake_start_post_process)

    class FakeProcess:
        def readAllStandardOutput(self):
            return (
                b'{"output_dir":"E:/tmp/out","outputs":["E:/tmp/out/transparent_master.png"],"failed":[],"warnings":[],"message":"ok"}'
            )

        def readAllStandardError(self):
            return b""

    page.process = FakeProcess()
    page._stdout_buffer = ""
    page._stderr_buffer = ""

    page._on_process_finished(0, None)

    assert scheduled["task_id"] == "20260623130100"
    assert scheduled["outputs"] == ["E:/tmp/out/transparent_master.png"]
    assert page.current_task.status == "完成"
    assert page.current_task.outputs == ["E:/tmp/out/transparent_master.png"]

    page.close()


def test_ai_edit_page_export_product_images_uses_task_target_and_syncs_putaway(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    final_transparent_dir = tmp_path / "final-transparent"
    final_transparent_dir.mkdir(parents=True, exist_ok=True)
    (final_transparent_dir / "BO-1661.png").write_bytes(b"print")
    final_product_dir = tmp_path / "final-product"
    final_product_dir.mkdir(parents=True, exist_ok=True)
    image = final_product_dir / "BO-1661_title.png"
    image.write_bytes(b"image")
    putaway_data_dir = tmp_path / "putaway-data"

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(
            ai_edit_api_key="stored-key",
            putaway_data_dir=str(putaway_data_dir),
            program_data_dir=str(tmp_path / "ConsolePlatData"),
        ),
    )
    record = AIEditTaskRecord(
        task_id="20260623130200",
        title="AI 改图 BO-1661",
        job=AIEditJob(images=[], prompt="keep subject", prefix="BO", start_number=1661),
        output_dir=str(tmp_path / "output"),
        final_transparent_dir=str(final_transparent_dir),
        final_product_dir=str(final_product_dir),
        xlsx_path=str(tmp_path / "batch.xlsx"),
    )
    Path(record.xlsx_path).write_bytes(b"xlsx")
    page.current_task = record

    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["task_id"] = current_record.task_id

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.export_task_product_images(record)

    assert not (final_product_dir / "导出产品图").exists()
    assert image.read_bytes() == b"image"
    assert scheduled == {"action": "export_product_images", "task_id": "20260623130200"}
    assert any(str(final_product_dir) in line for line in page.current_task.logs)

    page.close()


def test_ai_edit_page_export_product_images_schedules_background_job_and_logs(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])

    page, _path = _page_with_temp_store(
        tmp_path,
        monkeypatch,
        AppSettings(ai_edit_api_key="stored-key", program_data_dir=str(tmp_path / "ConsolePlatData")),
    )
    transparent_dir = tmp_path / "final-transparent"
    product_dir = tmp_path / "final-product"
    transparent_dir.mkdir(parents=True)
    product_dir.mkdir(parents=True)
    (transparent_dir / "BO-1661.png").write_bytes(b"print")
    (product_dir / "BO-1661_title.png").write_bytes(b"old")
    xlsx_path = tmp_path / "batch.xlsx"
    xlsx_path.write_bytes(b"xlsx")
    record = AIEditTaskRecord(
        task_id="20260623130201",
        title="AI edit BO-1661",
        job=AIEditJob(images=[], prompt="keep subject", prefix="BO", start_number=1661),
        final_transparent_dir=str(transparent_dir),
        final_product_dir=str(product_dir),
        xlsx_path=str(xlsx_path),
    )
    page.current_task = record
    page.tasks = [record]
    scheduled = {}

    def fake_start_background_job(action, current_record):
        scheduled["action"] = action
        scheduled["task_id"] = current_record.task_id

    monkeypatch.setattr(page, "_start_background_job", fake_start_background_job)

    page.export_task_product_images(record)

    assert scheduled == {"action": "export_product_images", "task_id": "20260623130201"}
    assert any("开始后台重贴产品图" in line for line in record.logs)

    page.close()


def test_ai_edit_background_export_product_images_rebuilds_from_transparent_and_syncs(tmp_path, monkeypatch):
    transparent_dir = tmp_path / "final-transparent"
    product_dir = tmp_path / "final-product"
    putaway_data_dir = tmp_path / "putaway-data"
    model_dir = tmp_path / "models"
    transparent_dir.mkdir(parents=True)
    product_dir.mkdir(parents=True)
    model_dir.mkdir(parents=True)
    print_one = transparent_dir / "BO-1661.png"
    print_two = transparent_dir / "BO-1662.png"
    print_one.write_bytes(b"print1")
    print_two.write_bytes(b"print2")
    old_product = product_dir / "BO-1661_title.png"
    old_product.write_bytes(b"old-product")
    xlsx_path = tmp_path / "batch.xlsx"
    xlsx_path.write_bytes(b"original-xlsx")
    record = AIEditTaskRecord(
        task_id="20260623130202",
        title="AI edit BO-1661",
        job=AIEditJob(images=[], prompt="keep subject", prefix="BO", start_number=1661),
        final_transparent_dir=str(transparent_dir),
        final_product_dir=str(product_dir),
        xlsx_path=str(xlsx_path),
    )
    settings = AppSettings(
        ai_edit_api_key="stored-key",
        bo_product_title="title",
        putaway_data_dir=str(putaway_data_dir),
        posai_model_root=str(model_dir),
    )
    captured = {}

    def fake_build_product_images(*, print_paths, final_product_dir, product_title, model_dir, **kwargs):
        captured["print_paths"] = [path.name for path in print_paths]
        captured["final_product_dir"] = final_product_dir
        captured["product_title"] = product_title
        captured["model_dir"] = model_dir
        outputs = []
        for print_path in print_paths:
            output = final_product_dir / f"{print_path.stem}_title.png"
            output.write_bytes(f"rebuilt:{print_path.name}".encode("utf-8"))
            outputs.append(output)
        return outputs, {path.stem: "白" for path in print_paths}

    def fake_sync_putaway_assets(**kwargs):
        captured["sync_source_images_dir"] = kwargs["source_images_dir"]
        captured["sync_source_xlsx_path"] = kwargs["source_xlsx_path"]
        captured["sync_target_data_dir"] = kwargs["target_data_dir"]
        captured["sync_force_replace"] = kwargs["force_replace"]
        return PutawaySyncSummary(
            ok=True,
            copied_images=2,
            copied_xlsx=True,
            images_target_dir=str(putaway_data_dir / "pic" / "1"),
            xlsx_target_path=str(putaway_data_dir / xlsx_path.name),
            message="synced",
        )

    monkeypatch.setattr("consoleplat.ui.ai_edit_page.build_product_images", fake_build_product_images)
    monkeypatch.setattr("consoleplat.ui.ai_edit_page.sync_putaway_assets", fake_sync_putaway_assets)
    worker = AIEditBackgroundWorker("export_product_images", record, settings)
    emitted_logs = []
    worker.log_message.connect(lambda task_id, line: emitted_logs.append((task_id, line)))

    result = worker._run_export_product_images()

    assert captured["print_paths"] == ["BO-1661.png", "BO-1662.png"]
    assert captured["final_product_dir"] == product_dir
    assert captured["product_title"] == "title"
    assert captured["model_dir"] == model_dir
    assert old_product.read_bytes() == b"rebuilt:BO-1661.png"
    assert not (product_dir / "导出产品图").exists()
    assert xlsx_path.read_bytes() == b"original-xlsx"
    assert captured["sync_source_images_dir"] == product_dir
    assert captured["sync_source_xlsx_path"] == xlsx_path
    assert captured["sync_target_data_dir"] == putaway_data_dir
    assert captured["sync_force_replace"] is True
    assert result.action == "export_product_images"
    assert result.outputs == [str(product_dir / "BO-1661_title.png"), str(product_dir / "BO-1662_title.png")]
    assert result.final_product_dir == str(product_dir)
    assert result.xlsx_path == str(xlsx_path)
    assert any(line == "开始重贴产品图：2 张透明底" for _task_id, line in emitted_logs)
    assert any("已重贴产品图 2 张" in line for _task_id, line in emitted_logs)
